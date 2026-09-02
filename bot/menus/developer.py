from __future__ import annotations
import asyncio
import sys
import json
import aiosqlite
import html
import pytz
import calendar
from typing import Optional, Union
from pathlib import Path
from datetime import datetime # NEW IMPORT
from database import DB_PATH
from aiogram import Dispatcher, Bot
from aiogram.types import CallbackQuery, InlineKeyboardMarkup, InlineKeyboardButton, Message, InputMediaPhoto, FSInputFile
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import StatesGroup, State
from aiogram.exceptions import TelegramBadRequest
from bot.config import MAIN_DEVELOPER_ID, DAYS_TOTAL, TIMEZONE
from bot.services.positions import get_all_positions, get_position_by_id, add_position, update_position_name, update_position_days, update_position_type, get_position_stats
from bot.services.cities import get_all_cities, get_city_by_id, add_city, update_city_name, delete_city
from bot.services.logger import get_logger
from bot.services.semantic_search import build_and_reset_embeddings
from database.users import (
    get_all_users,
    get_user_details,
    get_user_by_username,
    get_users_by_full_name, # NEW
    get_all_active_users,
    get_all_inactive_users,
    get_all_interns,
    get_all_workers,
    get_users_by_city,
    delete_user,
    register_user,
    get_reminder_history,
    get_user_progress,
    update_user_role,
    get_users_by_managers,
)
from database.hr import (
    is_developer_user,
    is_hr_user,
)
from database.managers import (
    get_all_managers,
    add_manager,
    delete_manager_by_uid,
    get_manager_by_uid,
    get_all_developers as get_all_devs_from_managers_db, # Alias to avoid name conflict if needed
    get_all_kerivnyky,
    is_territorial_user,
    is_observer_user,
    get_all_observers,
    get_observer_by_uid,
    delete_observer_by_uid,
    is_manager_trainee,
    promote_manager_trainee,
    get_managers_by_responsible,
    get_territorials_by_city,
    update_manager_name,
    update_manager_shops,
    fire_manager,
    search_managers,
    transfer_users_to_territorial_or_admin,
    count_users_in_shops,
    transfer_users_by_shops,
    update_manager_responsible,
)

from bot.services.developer_actions import (
    parse_day_input,
    get_user_days_report,
    open_days_for_user,
    close_days_for_user,
    reset_days_for_user,
)
from database.materials import (
    get_unique_roles_from_materials,
    get_material_by_role_day_type,
    update_material_content,
    update_material_resource_url,
    add_or_update_material,
    get_test_by_role_and_day,
    get_material_by_id, # NEW IMPORT
    toggle_test_status,
)
from database.tokens import get_token_stats, cleanup_expired_tokens, generate_token
from database.analytics import (
    get_daily_stats,
    get_dropout_funnel,
    get_training_added_interns,
    clear_training_added_interns,
    get_training_added_cities,
    get_training_left_inactive,
    get_training_promoted,
    get_training_rejected,
)
from bot.services.health import get_health_status
from bot.services.reports import get_report_data, get_report_details, generate_xlsx_report # NEW IMPORT
from bot.services.test_parser import parse_test_input, format_test_display
from bot.services.learning_progress import get_days_overview, DayStatus
from bot.services.access import get_display_role # Import get_display_role
from bot.constants import AVAILABLE_ROLES, AVAILABLE_SHOPS, AVAILABLE_CITIES
from bot.keyboards import get_pagination_keyboard # Import get_pagination_keyboard
from bot.utils.paginator import split_text # NEW IMPORT
import json


class DeveloperStates(StatesGroup):
    waiting_user_search = State()
    waiting_delete_user = State()
    waiting_add_manager_id = State()
    waiting_add_manager_name = State()
    waiting_add_manager_city = State()
    waiting_add_manager_shops = State()
    waiting_delete_manager = State()
    waiting_days_status = State()
    waiting_days_open = State()
    waiting_days_close = State()
    waiting_days_reset = State()
    waiting_add_developer = State()
    waiting_add_developer_name = State()
    waiting_add_hr = State()
    # Стани для редагування матеріалів
    waiting_material_text = State()
    waiting_material_video = State()
    waiting_material_parts = State() # New state for multi-message input
    waiting_test_input = State() # State for editing tests
    waiting_fix_page = State() # New state for fixing a specific page content
    waiting_complement_type = State() # Start/End/Between
    waiting_insert_start_idx = State() # First page index
    waiting_insert_end_idx = State() # Second page index
    waiting_complement_content = State() # New content to add
    waiting_insert_mode = State() # вибір: одна сторінка чи масив
    waiting_insert_array_content = State() # FSM цикл для вставки масиву до команди ГОТОВО
    waiting_replace_range = State() # очікування діапазону сторінок для заміни, напр. "3-5"
    waiting_replace_array_content = State() # FSM цикл для заміни до команди ГОТОВО
    waiting_announce_decision = State() # рішення щодо розсилки повідомлень користувачам
    # Стани для оповіщення про зміни
    waiting_notify_decision = State()
    waiting_notify_message = State()
    # Стани для зміни даних користувача
    waiting_change_manager_id = State()
    waiting_change_role = State()
    waiting_change_city = State()
    # Стани для редагування відео матеріалів
    waiting_video_role = State()
    waiting_video_day = State()
    waiting_video_uploads = State()
    # Стани для управління посадами
    waiting_add_position_name = State()
    waiting_add_position_days = State()
    waiting_add_position_type = State()  # Вибір типу: ТЗ або ВВ
    waiting_edit_position_name = State()
    waiting_edit_position_days = State()
    # Стани для управління містами
    waiting_add_city_name = State()
    waiting_edit_city_name = State()
    waiting_photo_uploads = State() # New state for photo uploads
    waiting_syllabus_text = State() # State for editing syllabus

    # Стани для управління територіалами
    waiting_add_territorial_uid = State()
    waiting_add_territorial_city = State()
    waiting_add_territorial_type = State()   # Вибір типу: ТЗ або ВВ
    waiting_add_territorial_transfer = State() # Запит чи переносити керівників

    # Стани для управління та пошуку керівників
    waiting_search_manager = State()
    waiting_edit_manager_name = State()
    waiting_edit_manager_shops = State()

async def _ensure_developer(callback: CallbackQuery, require_main: bool = False, allow_observer: bool = False) -> bool:
    user_id = callback.from_user.id
    is_dev = await is_developer_user(user_id)
    
    if require_main:
        # is_dev is only true if user is in 'developers' table AND is MAIN_DEVELOPER_ID
        is_dev = is_dev and user_id == MAIN_DEVELOPER_ID
    
    if allow_observer and not is_dev:
        if await is_observer_user(user_id):
            return True
            
    if not is_dev:
        try:
            await callback.answer("⛔️ Доступ заборонено. Ви не маєте відповідних прав.", show_alert=True)
        except Exception:
            pass
        return False
    return True

async def _check_access(callback: CallbackQuery) -> tuple[bool, bool, bool]:
    """
    Перевіряє доступ для панелі розробника/територіала/наглядача.
    Повертає кортеж (has_access, is_admin, is_territorial).
    Якщо доступу немає, показує alert.
    """
    user_id = callback.from_user.id
    is_admin = await is_developer_user(user_id)
    is_territorial = await is_territorial_user(user_id)
    is_observer = await is_observer_user(user_id)
    has_access = is_admin or is_territorial or is_observer
    
    if not has_access:
        try:
            await callback.answer("⛔️ Доступ заборонено. Ви не маєте потрібної ролі.", show_alert=True)
        except Exception:
            pass
            
    return has_access, is_admin, is_territorial


async def _send_in_chunks(message: Message, text: str, reply_markup=None, chunk_size: int = 3500):
    """
    Відправка довгих текстів частинами.
    - Якщо текст поміщається в один блок — оновлюємо поточне повідомлення (через _edit_or_answer).
    - Якщо ні — перший блок оновлює поточне повідомлення, решта надсилаються новими повідомленнями.
    """
    if len(text) <= chunk_size:
        await _edit_or_answer(message, text, reply_markup=reply_markup)
        return

    parts = []
    for line in text.split("\n"):
        if parts and len(parts[-1]) + len(line) + 1 <= chunk_size:
            parts[-1] += "\n" + line
        else:
            parts.append(line)

    # Першу частину показуємо реактивно в поточному повідомленні
    await _edit_or_answer(message, parts[0])

    # Проміжні частини без клавіатури
    for part in parts[1:-1]:
        await message.answer(part)

    # Остання частина з клавіатурою (якщо вона є)
    await message.answer(parts[-1], reply_markup=reply_markup)


def _admin_cho_keyboard(is_main_dev: bool, is_admin: bool, is_territorial: bool, is_observer: bool = False) -> InlineKeyboardMarkup:
    buttons = []
    if is_admin or is_observer:
        buttons.append([InlineKeyboardButton(text="📚 Навчальні матеріали", callback_data="dev_main_study")])
    buttons.append([InlineKeyboardButton(text="📊 Аналітика", callback_data="dev_main_analyt")])
    buttons.append([InlineKeyboardButton(text="👥 Команда Bulka", callback_data="dev_main_team")])
    if is_admin:
        buttons.append([InlineKeyboardButton(text="🛠 Інше", callback_data="dev_main_other")])
    buttons.append([InlineKeyboardButton(text="🏠 В головне меню", callback_data="main_menu")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)


def _admin_study_keyboard() -> InlineKeyboardMarkup:
    buttons = [
        [
            InlineKeyboardButton(text="📝 Матеріали", callback_data="dev_materials_menu"),
            InlineKeyboardButton(text="🎥 Відео", callback_data="dev_videos_menu")
        ],
        [
            InlineKeyboardButton(text="📝 Тести", callback_data="dev_tests_menu"),
            InlineKeyboardButton(text="🖼 Фото", callback_data="dev_photos_menu")
        ],
        [
            InlineKeyboardButton(text="📚 Змінити змісти", callback_data="dev_syllabus_menu"),
            InlineKeyboardButton(text="🧠 Нагадати тему", callback_data="remind_topic_global")
        ],
        [InlineKeyboardButton(text="🔙 Назад", callback_data="developer_menu")]
    ]
    return InlineKeyboardMarkup(inline_keyboard=buttons)


def _admin_analyt_keyboard(is_observer: bool = False) -> InlineKeyboardMarkup:
    row2 = [InlineKeyboardButton(text="📜 Історія нагадувань", callback_data="dev_reminder_history")]
    if not is_observer:
        row2.append(InlineKeyboardButton(text="📊 XLSX звіт", callback_data="dev_xlsx_menu"))
        
    buttons = [
        [
            InlineKeyboardButton(text="📊 Аналітика", callback_data="dev_analytics_menu"),
            InlineKeyboardButton(text="📊 Помилки тестів", callback_data="show_test_errors")
        ],
        [
            InlineKeyboardButton(text="📋 Ознайомлення", callback_data="dev_an_ack_menu")
        ],
        row2,
        [InlineKeyboardButton(text="🔙 Назад", callback_data="developer_menu")]
    ]
    return InlineKeyboardMarkup(inline_keyboard=buttons)


def _admin_team_keyboard(is_admin: bool, is_territorial: bool, is_observer: bool = False) -> InlineKeyboardMarkup:
    buttons = [
        [
            InlineKeyboardButton(text="👥 Користувачі", callback_data="dev_users_menu"),
            InlineKeyboardButton(text="👔 Команда Керівників", callback_data="dev_manage_managers")
        ]
    ]
    if is_admin:
        buttons.append([InlineKeyboardButton(text="👨‍💻 Команда Адміністраторів", callback_data="dev_team_menu")])
        buttons.append([
            InlineKeyboardButton(text="🗺 Територіали", callback_data="dev_territorials_menu"),
            InlineKeyboardButton(text="👁 Наглядачі", callback_data="dev_observers_menu")
        ])
    buttons.append([InlineKeyboardButton(text="🔙 Назад", callback_data="developer_menu")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)


def _admin_other_keyboard() -> InlineKeyboardMarkup:
    buttons = [
        [
            InlineKeyboardButton(text="👔 Посади", callback_data="dev_positions_menu"),
            InlineKeyboardButton(text="🏙 Міста", callback_data="dev_cities_menu")
        ],
        [
            InlineKeyboardButton(text="🎟 Токени", callback_data="dev_tokens_menu"),
            InlineKeyboardButton(text="💚 Health Status", callback_data="dev_health_status")
        ],
        [InlineKeyboardButton(text="🔙 Назад", callback_data="developer_menu")]
    ]
    return InlineKeyboardMarkup(inline_keyboard=buttons)


async def _show_admin_photo_menu(
    callback: CallbackQuery,
    photo_filename: str,
    caption: str,
    reply_markup: InlineKeyboardMarkup
):
    """
    Універсальний хелпер для головних панелей адмінки:
    - Якщо повідомлення вже має фото -> плавно змінює медіа (edit_media) або підпис (edit_caption).
    - Якщо повідомлення текстове або edit_media падає -> видаляє старе повідомлення і надсилає нове з фото.
    - Якщо фото файл не знайдено -> плавний текстовий фолбек.
    """
    photo_path = Path(photo_filename)
    if not photo_path.exists():
        if (Path("img/admin") / photo_filename).exists():
            photo_path = Path("img/admin") / photo_filename
        elif (Path("img") / photo_filename).exists():
            photo_path = Path("img") / photo_filename

    message = callback.message
    has_photo = bool(getattr(message, "photo", None))

    if has_photo and photo_path.exists():
        try:
            media = InputMediaPhoto(media=FSInputFile(str(photo_path)), caption=caption)
            await message.edit_media(media=media, reply_markup=reply_markup)
            try:
                await callback.answer()
            except Exception:
                pass
            return
        except Exception:
            pass
        try:
            await message.edit_caption(caption=caption, reply_markup=reply_markup)
            try:
                await callback.answer()
            except Exception:
                pass
            return
        except Exception:
            pass

    # Якщо повідомлення було текстовим (без фото) або редагування не вдалося:
    if message:
        try:
            await message.delete()
        except Exception:
            pass

    if photo_path.exists():
        try:
            await message.answer_photo(
                photo=FSInputFile(str(photo_path)),
                caption=caption,
                reply_markup=reply_markup
            )
            try:
                await callback.answer()
            except Exception:
                pass
            return
        except Exception:
            pass

    # Фолбек на текст, якщо фото недоступне
    await message.answer(caption, reply_markup=reply_markup)
    try:
        await callback.answer()
    except Exception:
        pass


# Alias for backward compatibility
_send_or_edit_admin_photo = _show_admin_photo_menu


async def developer_menu_callback(callback: CallbackQuery):
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access:
        return
    is_main = callback.from_user.id == MAIN_DEVELOPER_ID
    is_observer = await is_observer_user(callback.from_user.id)
    
    if is_admin:
        caption_text = "🛠 <b>Панель Адміністратора</b>\nОберіть розділ для керування:"
    elif is_observer:
        caption_text = "👁 <b>Панель Наглядача</b>\nОберіть розділ для перегляду:"
    else:
        caption_text = "🛠 <b>Панель Територіала</b>\nОберіть розділ для керування:"
    
    await _show_admin_photo_menu(
        callback,
        "img/admin/admin_cho.jpg",
        caption_text,
        _admin_cho_keyboard(is_main, is_admin, is_territorial, is_observer=is_observer)
    )


async def dev_main_study_handler(callback: CallbackQuery):
    has_access, is_admin, is_territorial = await _check_access(callback)
    is_observer = await is_observer_user(callback.from_user.id)
    if not has_access or (not is_admin and not is_observer):
        await callback.answer("⛔️ Доступ заборонено.", show_alert=True)
        return
    await _show_admin_photo_menu(
        callback,
        "img/admin/admin_study.jpg",
        "📚 <b>Навчальні матеріали</b>",
        _admin_study_keyboard()
    )


async def dev_main_analyt_handler(callback: CallbackQuery):
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access:
        return
    is_observer = await is_observer_user(callback.from_user.id)
    await _show_admin_photo_menu(
        callback,
        "img/admin/admin_analyt.jpg",
        "📊 <b>Аналітика</b>",
        _admin_analyt_keyboard(is_observer=is_observer)
    )


async def dev_main_team_handler(callback: CallbackQuery):
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access:
        return
    is_observer = await is_observer_user(callback.from_user.id)
    await _show_admin_photo_menu(
        callback,
        "img/admin/admin_spus.jpg",
        "👥 <b>Команда Bulka</b>",
        _admin_team_keyboard(is_admin, is_territorial, is_observer=is_observer)
    )


async def dev_main_other_handler(callback: CallbackQuery):
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access or not is_admin:
        await callback.answer("⛔️ Доступ заборонено.", show_alert=True)
        return
    await _show_admin_photo_menu(
        callback,
        "img/admin/admin_tools.jpg",
        "🛠 <b>Інше</b>",
        _admin_other_keyboard()
    )


async def _edit_or_answer(message: Message, text: str, reply_markup=None):
    """
    Helper: намагається оновити поточне повідомлення, а якщо не виходить — надсилає нове.
    Якщо текст занадто довгий для медіа-підпису (>1024), автоматично переходить на текст.
    """
    # Якщо текст довший за ліміт підпису (з невеликим запасом), 
    # або якщо ми хочемо гарантувати доставку великого списку
    if len(text) > 1000:
        try:
            await message.delete()
        except:
            pass
        return await message.answer(text, reply_markup=reply_markup)

    try:
        if message.photo:
            await message.edit_caption(caption=text, reply_markup=reply_markup)
            return message
        await message.edit_text(text, reply_markup=reply_markup)
        return message
    except TelegramBadRequest as e:
        error_message = str(e).lower()
        if "message to edit not found" in error_message or "there is no text in the message to edit" in error_message:
            try:
                await message.delete()
            except:
                pass
            return await message.answer(text, reply_markup=reply_markup)
        elif "message is not modified" in error_message:
            return message
        else:
            # Для будь-яких інших помилок (наприклад, MEDIA_CAPTION_TOO_LONG), 
            # намагаємося надіслати новим текстовим повідомленням
            try:
                await message.delete()
            except:
                pass
            return await message.answer(text, reply_markup=reply_markup)
    except Exception:
        try:
            await message.delete()
        except:
            pass
        return await message.answer(text, reply_markup=reply_markup)


async def _remember_panel(state: FSMContext, key: str, message: Message):
    await state.update_data(
        **{
            key: {
                "chat_id": message.chat.id,
                "message_id": message.message_id,
            }
        }
    )


async def _refresh_panel_view(bot, panel_info: Optional[dict], builder):
    if not panel_info:
        return
    chat_id = panel_info.get("chat_id")
    message_id = panel_info.get("message_id")
    if not chat_id or not message_id:
        return
    text, kb = await builder()
    try:
        await bot.edit_message_text(text, chat_id=chat_id, message_id=message_id, reply_markup=kb)
    except Exception:
        await bot.send_message(chat_id, text, reply_markup=kb)


async def get_user_avatar_input(bot: Optional[Bot], user_id: int) -> Union[str, FSInputFile, None]:
    """
    Tries to fetch the user's Telegram profile photo.
    Returns the file_id (str) if available, or FSInputFile("img/ava.png") as fallback.
    """
    if bot:
        try:
            photos = await bot.get_user_profile_photos(user_id, limit=1)
            if photos and photos.total_count > 0 and photos.photos:
                # Largest size photo is the last element in the list
                return photos.photos[0][-1].file_id
        except Exception:
            pass
    
    fallback_path = Path("img/ava.png")
    if fallback_path.exists():
        return FSInputFile(str(fallback_path))
    return None


async def _send_or_edit_card_photo(
    callback: CallbackQuery,
    photo_input: Union[str, FSInputFile, None],
    caption: str,
    reply_markup: Optional[InlineKeyboardMarkup] = None
):
    """
    Renders or edits a card photo message:
    - If the message already has a photo, calls edit_media or edit_caption.
    - If it's a text message or edit fails, deletes old message and sends answer_photo.
    - If photo_input is None, falls back to answer.
    """
    message = callback.message
    has_photo = bool(getattr(message, "photo", None))
    
    if has_photo and photo_input:
        try:
            media = InputMediaPhoto(media=photo_input, caption=caption)
            await message.edit_media(media=media, reply_markup=reply_markup)
            try:
                await callback.answer()
            except Exception:
                pass
            return
        except Exception:
            pass
        try:
            await message.edit_caption(caption=caption, reply_markup=reply_markup)
            try:
                await callback.answer()
            except Exception:
                pass
            return
        except Exception:
            pass

    # Якщо повідомлення було текстовим або редагування не вдалося:
    if message:
        try:
            await message.delete()
        except Exception:
            pass

    if photo_input:
        try:
            await message.answer_photo(
                photo=photo_input,
                caption=caption,
                reply_markup=reply_markup
            )
            try:
                await callback.answer()
            except Exception:
                pass
            return
        except Exception:
            pass

    # Фолбек на звичайний текст
    await message.answer(caption, reply_markup=reply_markup)
    try:
        await callback.answer()
    except Exception:
        pass


async def _format_identity(user_id: int, fallback_name=None, fallback_username=None, bot: Optional[Bot] = None):
    # Пріоритет віддаємо даним, переданим напряму (з таблиці managers)
    full_name = fallback_name
    username = fallback_username

    # Якщо дані відсутні, пробуємо отримати їх з таблиці users як запасний варіант
    if not full_name or not username:
        profile = await get_user_details(user_id)
        if profile:
            if not full_name:
                full_name = profile.get("full_name")
            if not username:
                username = profile.get("username")

    # Якщо досі немає username або full_name за замовчуванням, і є bot - запитуємо Telegram API get_chat
    if (not username or not full_name or full_name in ("Головний Адміністратор", "Без імені")) and bot:
        try:
            chat = await bot.get_chat(user_id)
            if chat:
                if not username and chat.username:
                    username = chat.username
                t_fullname = f"{chat.first_name or ''} {chat.last_name or ''}".strip()
                if (not full_name or full_name in ("Головний Адміністратор", "Без імені")) and t_fullname:
                    full_name = t_fullname
                # Оновлюємо кеш в managers та users
                try:
                    async with aiosqlite.connect(DB_PATH) as db:
                        if username:
                            await db.execute("UPDATE managers SET username = ? WHERE uid = ?", (username, user_id))
                            await db.execute("UPDATE users SET username = ? WHERE user_id = ?", (username, user_id))
                        if t_fullname and full_name in ("Головний Адміністратор", "Без імені"):
                            await db.execute("UPDATE managers SET name = ? WHERE uid = ?", (t_fullname, user_id))
                        await db.commit()
                except Exception:
                    pass
        except Exception:
            pass

    # Фінальні перевірки, щоб уникнути None
    full_name = full_name or "Без імені"
    username = username or ""
    
    username_display = f"@{username}" if username else "без username"
    return full_name, username_display


async def _build_managers_team_view(is_admin: bool, is_territorial: bool, user_id: int, page: int = 0, bot: Optional[Bot] = None):
    """Список керівників з пагінацією, пошуком та прямим переходом до карток."""
    is_obs = await is_observer_user(user_id)
    if is_admin or is_obs:
        all_hrs = await get_all_kerivnyky()
    elif is_territorial:
        all_hrs = await get_managers_by_responsible(user_id)
    else:
        all_hrs = []
    
    city_abbr = {
        "Хмельницький": "ХМ",
        "Камʼянець-Подільський": "КП"
    }
    
    if not all_hrs:
        buttons = []
        if is_admin:
            buttons.append([InlineKeyboardButton(text="➕ Додати керівника", callback_data="dev_add_manager")])
        buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_main_team")])
        kb = InlineKeyboardMarkup(inline_keyboard=buttons)
        return "👔 <b>Команда керівників</b>\n\nПоки що немає", kb

    items_per_page = 10
    total_count = len(all_hrs)
    pages_total = (total_count + items_per_page - 1) // items_per_page
    
    if page < 0: page = 0
    if page >= pages_total and pages_total > 0: page = pages_total - 1
    
    start = page * items_per_page
    end = start + items_per_page
    paginated_hrs = all_hrs[start:end]
    
    lines = [f"👔 <b>Команда керівників</b> (Всього: {total_count}, Стор. {page + 1}/{pages_total})", "Натисніть на керівника для перегляду картки:", ""]
    
    buttons = [
        [InlineKeyboardButton(text="🔍 Пошук керівника", callback_data="dev_mgr_search_start")]
    ]
    
    for idx, hr in enumerate(paginated_hrs, start=start + 1):
        name, username = await _format_identity(
            hr["uid"],
            hr.get("full_name"),
            hr.get("username"),
            bot=bot,
        )
        
        shop_list = hr.get("shops") or []
        if isinstance(shop_list, str):
            try:
                shop_list = json.loads(shop_list)
            except Exception:
                shop_list = [shop_list] if shop_list else []
        shop_codes = [s.split(" ")[0] for s in shop_list]
        shops_str = ", ".join(shop_codes) if shop_codes else "Не вказано"
        
        city = hr.get("city")
        if not city and shop_list:
            from bot.constants import AVAILABLE_SHOPS
            for city_name, city_shops in AVAILABLE_SHOPS.items():
                if any(s in city_shops for s in shop_list):
                    city = city_name
                    break
        
        abbr = city_abbr.get(city, "??")
        
        btn_text = f"👔 {name} | {username} | {abbr}"
        buttons.append([InlineKeyboardButton(text=btn_text, callback_data=f"dev_mgr_view:{hr['uid']}")])
    
    # Кнопки навігації
    nav_row = []
    if page > 0:
        nav_row.append(InlineKeyboardButton(text="⬅️", callback_data=f"dev_mgr_page:{page - 1}"))
    nav_row.append(InlineKeyboardButton(text=f"📄 {page + 1}/{pages_total}", callback_data="ignore"))
    if page < pages_total - 1:
        nav_row.append(InlineKeyboardButton(text="➡️", callback_data=f"dev_mgr_page:{page + 1}"))
    
    if len(nav_row) > 1:
        buttons.append(nav_row)
        
    if is_admin:
        buttons.append([InlineKeyboardButton(text="➕ Додати керівника", callback_data="dev_add_manager")])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_main_team")])
    
    return "\n".join(lines), InlineKeyboardMarkup(inline_keyboard=buttons)


async def _build_dev_team_view(bot: Optional[Bot] = None) -> tuple[str, InlineKeyboardMarkup]:
    """Builds the view for the Dev Team management panel with nice formatting and interactive inline cards."""
    developers = await get_all_devs_from_managers_db()
    
    # Сортуємо: головний розробник завжди перший
    developers.sort(key=lambda x: x["uid"] != MAIN_DEVELOPER_ID)
    
    lines = [
        f"👨‍💻 <b>Команда Адміністраторів (всього: {len(developers)})</b>",
        "Натисніть на адміністратора для перегляду картки:",
        ""
    ]
    buttons = []
    if not developers:
        lines.append("  Немає адміністраторів у команді.")
    else:
        for idx, dev in enumerate(developers, start=1):
            name, username = await _format_identity(
                dev["uid"], dev.get("full_name"), dev.get("username"), bot=bot
            )
            btn_text = f"👨‍💻 {name} | {username}"
            buttons.append([InlineKeyboardButton(text=btn_text, callback_data=f"dev_admin_view:{dev['uid']}")])
    
    buttons.append([InlineKeyboardButton(text="➕ Додати адміністратора", callback_data="dev_add_dev")])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_main_team")])
    
    return "\n".join(lines), InlineKeyboardMarkup(inline_keyboard=buttons)


async def developer_dev_team_menu(callback: CallbackQuery, state: FSMContext):
    """Handler to show the Dev Team management menu."""
    if not await _ensure_developer(callback):
        return
    text, kb = await _build_dev_team_view(bot=callback.bot)
    await _send_or_edit_admin_photo(callback, "admin_spus.jpg", text, kb)
    await _remember_panel(state, "dev_panel", callback.message)


async def developer_admin_view(callback: CallbackQuery, state: FSMContext):
    """Відображає картку адміністратора з фото та системною аналітикою."""
    if not await _ensure_developer(callback):
        return
    try:
        admin_uid = int(callback.data.split(":")[1])
    except (ValueError, IndexError):
        return await callback.answer("Помилка ідентифікатора.", show_alert=True)

    bot = callback.bot
    devs = await get_all_devs_from_managers_db()
    admin_info = next((d for d in devs if d["uid"] == admin_uid), None)
    
    name, username = await _format_identity(
        admin_uid, 
        admin_info.get("full_name") if admin_info else None,
        admin_info.get("username") if admin_info else None,
        bot=bot
    )
    
    role_label = "👑 Головний адміністратор" if admin_uid == MAIN_DEVELOPER_ID else "👨‍💻 Адміністратор"
    
    # Збираємо глобальну статистику
    async with aiosqlite.connect(DB_PATH) as db:
        c_all = await db.execute("SELECT COUNT(*) FROM users")
        total_users = (await c_all.fetchone())[0]
        
        c_interns = await db.execute("SELECT COUNT(*) FROM users WHERE status IS NULL OR status != 'Працівник'")
        total_interns = (await c_interns.fetchone())[0]
        
        c_workers = await db.execute("SELECT COUNT(*) FROM users WHERE status = 'Працівник'")
        total_workers = (await c_workers.fetchone())[0]
        
    managers = await get_all_kerivnyky()
    total_managers = len(managers)
    territorials = await get_all_territorials()
    total_territorials = len(territorials)
    
    caption = (
        f"<b>{role_label}:</b> {html.escape(name)}\n"
        f"👤 <b>Username:</b> {html.escape(username)}\n"
        f"🆔 <b>Telegram ID:</b> <code>{admin_uid}</code>\n"
        f"👑 <b>Роль:</b> {role_label}\n\n"
        f"📊 <b>Загальна статистика системи:</b>\n"
        f"• 👥 Всього користувачів: {total_users}\n"
        f"• 🎓 Активних стажерів: {total_interns}\n"
        f"• 👷 Випущених працівників: {total_workers}\n"
        f"• 👔 Керівників: {total_managers}\n"
        f"• 🗺 Територіалів: {total_territorials}"
    )
    
    buttons = []
    if admin_uid != MAIN_DEVELOPER_ID:
        buttons.append([InlineKeyboardButton(text="❌ Видалити адміністратора", callback_data=f"dev_remove_confirm:{admin_uid}")])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад до списку", callback_data="dev_team_menu")])
    
    photo_input = await get_user_avatar_input(bot, admin_uid)
    await _send_or_edit_card_photo(callback, photo_input, caption, InlineKeyboardMarkup(inline_keyboard=buttons))

async def developer_remove_dev_menu(callback: CallbackQuery, state: FSMContext):
    """Shows a menu to select a developer to remove."""
    if not await _ensure_developer(callback):
        return
    
    developers = await get_all_devs_from_managers_db()
    if not developers:
        await callback.answer("Немає адміністраторів для видалення.", show_alert=True)
        return
    
    buttons = []
    for dev in developers:
        name, _ = await _format_identity(dev["uid"], dev.get("full_name"))
        buttons.append([
            InlineKeyboardButton(
                text=f"❌ {name}", 
                callback_data=f"dev_remove_confirm:{dev['uid']}"
            )
        ])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_team_menu")])
    
    await _edit_or_answer(
        callback.message,
        "🗑 Оберіть адміністратора для видалення:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    try:
        await callback.answer()
    except TelegramBadRequest as e:
        if "query is too old" not in str(e):
            raise e

async def developer_remove_dev(callback: CallbackQuery, state: FSMContext):
    """Performs the removal of a developer after confirmation."""
    if not await _ensure_developer(callback):
        return
    
    try:
        user_id_to_remove = int(callback.data.split(":")[1])
    except (ValueError, IndexError):
        await callback.answer("Некоректний ID.", show_alert=True)
        return

    # Захист від видалення головного розробника
    if user_id_to_remove == MAIN_DEVELOPER_ID:
        await callback.answer(
            "⛔️ Ви не можете видалити головного адміністратора.",
            show_alert=True
        )
        return
        
    await delete_manager_by_uid(user_id_to_remove)
    await callback.answer(f"Адміністратора {user_id_to_remove} видалено з команди.", show_alert=True)
    
    # Оновлюємо вигляд меню
    await developer_dev_team_menu(callback, state)

# ---------- Users section ----------

def _users_menu_keyboard(is_admin: bool, is_territorial: bool, is_observer: bool = False) -> InlineKeyboardMarkup:
    buttons = [
        [
            InlineKeyboardButton(text="🎓 Стажери", callback_data="dev_users_interns"),
            InlineKeyboardButton(text="👷 Працівники", callback_data="dev_users_workers"),
        ],
        [
            InlineKeyboardButton(text="🚀 Активні стажери", callback_data="dev_users_active"),
            InlineKeyboardButton(text="😴 Неактивні", callback_data="dev_users_inactive"),
        ],
    ]
    if is_admin or is_observer:
        buttons.append([InlineKeyboardButton(text="🏙️ За містом", callback_data="dev_users_by_city")])
        if is_admin:
            buttons.append([InlineKeyboardButton(text="✅ Перевести завершених у працівники", callback_data="dev_users_bulk_promote")])
        buttons.append([InlineKeyboardButton(text="📋 Список користувачів", callback_data="dev_users_list")])
        buttons.append([InlineKeyboardButton(text="🔍 Пошук", callback_data="dev_users_search")])
        if is_admin:
            buttons.append([InlineKeyboardButton(text="❌ Видалити", callback_data="dev_users_delete")])

    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_main_team")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)

async def _get_users_menu_kb(user_id: int) -> InlineKeyboardMarkup:
    from database.hr import is_developer_user
    from database.managers import is_territorial_user, is_observer_user
    is_admin = await is_developer_user(user_id)
    is_territorial = await is_territorial_user(user_id)
    is_observer = await is_observer_user(user_id)
    return _users_menu_keyboard(is_admin, is_territorial, is_observer=is_observer)


def _cancel_keyboard(callback_data: str) -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[[InlineKeyboardButton(text="❌ Скасувати", callback_data=callback_data)]]
    )


async def _remember_panel(state: FSMContext, key: str, message: Message):
    await state.update_data(
        **{
            key: {
                "chat_id": message.chat.id,
                "message_id": message.message_id,
            }
        }
    )


async def _refresh_panel_view(bot, panel_info: Optional[dict], builder):
    if not panel_info:
        return
    chat_id = panel_info.get("chat_id")
    message_id = panel_info.get("message_id")
    if not chat_id or not message_id:
        return
    text, kb = await builder()
    try:
        await bot.edit_message_text(text, chat_id=chat_id, message_id=message_id, reply_markup=kb)
    except Exception:
        await bot.send_message(chat_id, text, reply_markup=kb)


async def _build_developer_team_view():
    developers = await get_all_devs_from_managers_db()
    text = (
        f"👨‍💻 <b>Команда Адміністраторів</b>\n"
        f"Усього учасників: <b>{len(developers)}</b>\n"
        "Оберіть адміністратора, щоб керувати його доступом, або додайте нового."
    )
    rows = []
    if developers:
        for dev in developers:
            name, _ = await _format_identity(
                dev["uid"], dev.get("full_name"), dev.get("username")
            )
            rows.append([
                InlineKeyboardButton(
                    text=f"{name} ({dev['uid']})",
                    callback_data=f"dev_team_member_{dev['uid']}",
                )
            ])
    rows.append([InlineKeyboardButton(text="➕ Додати адміністратора", callback_data="dev_add_developer")])
    rows.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_main_team")])
    return text, InlineKeyboardMarkup(inline_keyboard=rows)


async def _build_hr_team_view():
    hrs = await get_all_kerivnyky()
    text = (
        f"👨‍💼 <b>Команда HR</b>\n"
        f"Усього учасників: <b>{len(hrs)}</b>\n"
        "Оберіть HR, щоб керувати доступом, або додайте нового."
    )
    rows = []
    if hrs:
        for hr in hrs:
            name, _ = await _format_identity(
                hr["uid"], hr.get("full_name"), hr.get("username")
            )
            rows.append([
                InlineKeyboardButton(
                    text=f"{name} ({hr['uid']})",
                    callback_data=f"hr_team_member_{hr['uid']}",
                )
            ])
    rows.append([InlineKeyboardButton(text="➕ Додати HR", callback_data="dev_add_hr")])
    rows.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="developer_menu")])
    return text, InlineKeyboardMarkup(inline_keyboard=rows)


async def _show_prompt(callback: CallbackQuery, text: str, cancel_callback: str):
    kb = _cancel_keyboard(cancel_callback)
    msg = await _edit_or_answer(callback.message, text, reply_markup=kb)
    return {"chat_id": msg.chat.id, "message_id": msg.message_id}


async def _update_prompt_message(bot, prompt_info: Optional[dict], text: str, cancel_callback: str):
    if not prompt_info:
        return
    kb = _cancel_keyboard(cancel_callback)
    try:
        await bot.edit_message_text(
            chat_id=prompt_info["chat_id"],
            message_id=prompt_info["message_id"],
            text=text,
            reply_markup=kb,
        )
    except Exception:
        pass


async def developer_users_menu(callback: CallbackQuery):
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access:
        return
    await _edit_or_answer(
        callback.message,
        "👥 <b>Користувачі</b>\nОберіть дію:",
        reply_markup=await _get_users_menu_kb(callback.from_user.id),
    )
    await callback.answer()

async def _filter_users_for_territorial(user_id: int, users: list) -> list:
    """
    Фільтрує список користувачів для територіала за:
      1) Містом (user.city == territorial.city)
      2) Типом посади (positions.territorial_type == territorial.territorial_type: 'ТЗ' або 'ВВ')
    """
    from database.managers import get_manager_by_uid
    from database.positions import get_all_positions
    
    mgr = await get_manager_by_uid(user_id)
    if not mgr or mgr.get("process") != "Територіал":
        return users
        
    t_city = (mgr.get("city") or "").strip().lower()
    t_type = (mgr.get("territorial_type") or "ТЗ").strip()
    
    positions = await get_all_positions()
    pos_type_map = {p["name"]: p.get("territorial_type", "ТЗ") for p in positions}
    
    filtered = []
    for u in users:
        u_city = (u.get("city") or "").strip().lower()
        u_role = u.get("role") or ""
        u_type = pos_type_map.get(u_role, "ТЗ")
        
        # Перевірка збігу міста та напрямку посади (ТЗ або ВВ)
        if u_city == t_city and u_type == t_type:
            filtered.append(u)
            
    return filtered


async def _developer_show_users_list(callback: CallbackQuery, users: list, title: str, mode: str, page: int = 0, show_day: bool = True):
    user_id = callback.from_user.id
    from database.hr import is_developer_user
    from database.managers import is_territorial_user
    is_admin = await is_developer_user(user_id)
    is_territorial = await is_territorial_user(user_id)
    if not users:
        await _edit_or_answer(callback.message, f"{title}\n\nСписок порожній.", reply_markup=await _get_users_menu_kb(callback.from_user.id))
        try:
             await callback.answer()
        except:
             pass
        return

    # Розрахунок пагінації
    items_per_page = 8
    total_count = len(users)
    pages_total = (total_count + items_per_page - 1) // items_per_page
    
    # Корекція поточної сторінки
    curr_page = page
    if curr_page < 0: curr_page = 0
    if curr_page >= pages_total and pages_total > 0: curr_page = pages_total - 1

    start_offset = curr_page * items_per_page
    end_offset = start_offset + items_per_page
    paginated_users = users[start_offset:end_offset]

    lines = [f"{title} (Всього: {total_count}, Стор. {curr_page + 1}/{pages_total})", ""]
    
    for idx, user in enumerate(paginated_users, start=start_offset + 1):
        # Визначаємо роль та посаду
        display_role = await get_display_role(user['user_id'])
        is_worker = user.get("status") == "Працівник"
        
        if display_role in ["Dev", "Керівник"]:
            job_title = display_role
        elif is_worker:
            job_title = user.get('role') or "Не вказано"
        elif mode == "all":
            job_title = "Стажер"
        else:
            job_title = user.get('role') or "Не вказано"
            
        user_full_name = user.get('full_name', 'Без імені')
        user_username = user.get('username', 'немає')
        current_block = user.get('current_block', 1)
        
        # Витягуємо короткий номер магазину (наприклад, B-19)
        shop_full = user.get('shop', '') or ''
        shop_short = shop_full.split(' ')[0] if shop_full else 'Не вказано'
        
        # Екранування
        e_full_name = html.escape(user_full_name)
        e_username = html.escape(user_username)
        e_job = html.escape(job_title)
        e_shop = html.escape(shop_short)
        
        # Формуємо рядок: день показуємо тільки якщо show_day=True і це не Dev/Керівник/Працівник
        if not show_day or display_role in ["Dev", "Керівник"] or is_worker:
            status_label = " | Статус: <b>Завершено</b>" if is_worker else ""
            info_line = f"   @{e_username} | Посада: <b>{e_job}</b> | {e_shop}{status_label}"
        else:
            info_line = f"   @{e_username} | Посада: <b>{e_job}</b> | {e_shop} | {current_block} день"
            
        lines.append(f"<b>{idx}. {e_full_name}</b> (ID: <code>{user['user_id']}</code>)")
        lines.append(info_line)
        lines.append("───────────────")

    # Формування кнопок
    pagination_buttons = []
    row = []
    if curr_page > 0:
        row.append(InlineKeyboardButton(text="⬅️ Попередня", callback_data=f"dev_users_pag:{mode}:{curr_page - 1}"))
    if curr_page < pages_total - 1:
        row.append(InlineKeyboardButton(text="Наступна ➡️", callback_data=f"dev_users_pag:{mode}:{curr_page + 1}"))
    if row:
        pagination_buttons.append(row)

    if mode == "interns":
        pagination_buttons.append([InlineKeyboardButton(text="📥 Вивантажити список", callback_data="dev_interns_export_xlsx")])
    elif mode == "workers":
        pagination_buttons.append([InlineKeyboardButton(text="📥 Вивантажити список", callback_data="dev_workers_export_xlsx")])

    pagination_buttons.append([InlineKeyboardButton(text="⬅️ До меню користувачів", callback_data="dev_users_menu")])
    
    await _edit_or_answer(
        callback.message,
        "\n".join(lines),
        reply_markup=InlineKeyboardMarkup(inline_keyboard=pagination_buttons),
    )
    try:
        await callback.answer()
    except:
        pass

async def developer_list_users(callback: CallbackQuery, page: int = 0):
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access: return
    users = await get_all_users()
    if is_territorial: users = await _filter_users_for_territorial(callback.from_user.id, users)
    await _developer_show_users_list(callback, users, "👥 <b>Всі користувачі</b>", "all", page)

async def developer_active_users(callback: CallbackQuery, page: int = 0):
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access: return
    users = await get_all_active_users(days=3)
    if is_territorial: users = await _filter_users_for_territorial(callback.from_user.id, users)
    await _developer_show_users_list(callback, users, "🚀 <b>Активні стажери</b>", "active", page)

async def developer_inactive_users(callback: CallbackQuery, page: int = 0):
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access: return
    users = await get_all_inactive_users(days=3)
    if is_territorial: users = await _filter_users_for_territorial(callback.from_user.id, users)
    await _developer_show_users_list(callback, users, "😴 <b>Неактивні стажери</b>", "inactive", page)

async def developer_list_interns(callback: CallbackQuery, page: int = 0):
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access: return
    users = await get_all_interns()
    if is_territorial: users = await _filter_users_for_territorial(callback.from_user.id, users)
    await _developer_show_users_list(callback, users, "🎓 <b>Стажери</b>", "interns", page, show_day=True)

async def developer_list_workers(callback: CallbackQuery, page: int = 0):
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access: return
    users = await get_all_workers()
    if is_territorial: users = await _filter_users_for_territorial(callback.from_user.id, users)
    await _developer_show_users_list(callback, users, "👷 <b>Працівники</b>", "workers", page, show_day=False)

async def developer_bulk_promote_completed_interns(callback: CallbackQuery):
    """Масово переводить завершених стажерів у працівники."""
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access:
        return

    interns = await get_all_interns()
    promoted_count = 0

    for intern in interns:
        progress = await get_user_progress(intern["user_id"])
        completed_days = sum(1 for p in progress if p.get("completed"))
        if completed_days >= DAYS_TOTAL:
            await update_user_role(intern["user_id"], "Працівник", actor_id=callback.from_user.id)
            promoted_count += 1

    if promoted_count == 0:
        text = (
            "✅ <b>Масовий перехід завершено.</b>\n\n"
            "Не знайдено стажерів, яких потрібно перевести у працівники."
        )
    else:
        text = (
            "✅ <b>Масовий перехід завершено.</b>\n\n"
            f"У статус <b>Працівник</b> переведено: <b>{promoted_count}</b>"
        )

    await _edit_or_answer(callback.message, text, reply_markup=await _get_users_menu_kb(callback.from_user.id))
    await callback.answer("Готово ✅")

async def developer_interns_export_xlsx(callback: CallbackQuery):
    """Вивантажує список усіх стажерів у xlsx з колонками Ім'я / День / Магазин / Керівник."""
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access: return

    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from aiogram.types import BufferedInputFile

    users = await get_all_interns()
    if is_territorial: users = await _filter_users_for_territorial(callback.from_user.id, users)

    # Збираємо імена керівників одним проходом (кешуємо щоб не дублювати запити)
    manager_cache: dict = {}
    for user in users:
        mid = user.get("manager_id")
        if mid and mid not in manager_cache:
            mgr = await get_manager_by_uid(mid)
            if mgr:
                manager_cache[mid] = mgr.get("full_name") or mgr.get("username") or str(mid)
            else:
                # Fallback: developer може бути тільки в таблиці users
                fallback = await get_user_details(mid)
                manager_cache[mid] = (fallback.get("full_name") or fallback.get("username") or str(mid)) if fallback else str(mid)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Стажери"

    # Заголовки
    headers = ["Ім'я", "День", "Магазин", "Керівник"]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(fill_type="solid", fgColor="4472C4")
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Дані
    for row_idx, user in enumerate(users, start=2):
        shop_full = user.get("shop") or ""
        shop_short = shop_full.split(" ")[0] if shop_full else "—"
        manager_name = manager_cache.get(user.get("manager_id"), "—")
        ws.cell(row=row_idx, column=1, value=user.get("full_name") or user.get("username") or "—")
        ws.cell(row=row_idx, column=2, value=user.get("current_block") or 1)
        ws.cell(row=row_idx, column=3, value=shop_short)
        ws.cell(row=row_idx, column=4, value=manager_name)

    # Ширина колонок
    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 30

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    await callback.message.answer_document(
        document=BufferedInputFile(buf.getvalue(), filename="Стажери.xlsx"),
        caption=f"📋 <b>Список стажерів</b> — {len(users)} осіб",
        parse_mode="HTML",
    )
    await callback.answer()

async def developer_workers_export_xlsx(callback: CallbackQuery):
    """Вивантажує список усіх працівників у xlsx з колонками Ім'я / Магазин / Керівник."""
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access:
        await callback.answer("Вам недоступна ця функція.", show_alert=True)
        return

    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from aiogram.types import BufferedInputFile

    users = await get_all_workers()
    if is_territorial and not is_admin:
        users = await _filter_users_for_territorial(callback.from_user.id, users)

    manager_cache: dict = {}
    for user in users:
        mid = user.get("manager_id")
        if mid and mid not in manager_cache:
            mgr = await get_manager_by_uid(mid)
            if mgr:
                manager_cache[mid] = mgr.get("full_name") or mgr.get("username") or str(mid)
            else:
                fallback = await get_user_details(mid)
                manager_cache[mid] = (fallback.get("full_name") or fallback.get("username") or str(mid)) if fallback else str(mid)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Працівники"

    headers = ["Ім'я", "Посада", "Магазин", "Керівник"]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(fill_type="solid", fgColor="217346")
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_idx, user in enumerate(users, start=2):
        shop_full = user.get("shop") or ""
        shop_short = shop_full.split(" ")[0] if shop_full else "—"
        manager_name = manager_cache.get(user.get("manager_id"), "—")
        raw_role = user.get("role") or ""
        job_title = raw_role if raw_role and raw_role != "Працівник" else "—"
        ws.cell(row=row_idx, column=1, value=user.get("full_name") or user.get("username") or "—")
        ws.cell(row=row_idx, column=2, value=job_title)
        ws.cell(row=row_idx, column=3, value=shop_short)
        ws.cell(row=row_idx, column=4, value=manager_name)

    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 30

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    await callback.message.answer_document(
        document=BufferedInputFile(buf.getvalue(), filename="Працівники.xlsx"),
        caption=f"👷 <b>Список працівників</b> — {len(users)} осіб",
        parse_mode="HTML",
    )
    await callback.answer()

async def developer_users_by_city_menu(callback: CallbackQuery):
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access or (is_territorial and not is_admin):
        await callback.answer("Вам недоступна ця функція.", show_alert=True)
        return
    
    buttons = []
    for city in AVAILABLE_CITIES:
        buttons.append([InlineKeyboardButton(text=city, callback_data=f"dev_users_filter_city:{city}")])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_users_menu")])
    
    await _edit_or_answer(
        callback.message,
        "🏙️ <b>Оберіть місто для фільтрації:</b>",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    try: await callback.answer()
    except: pass

async def developer_users_filter_city(callback: CallbackQuery):
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access or (is_territorial and not is_admin):
        await callback.answer("Вам недоступна ця функція.", show_alert=True)
        return
    try:
        city = callback.data.split(":")[1]
    except IndexError:
        return
    users = await get_users_by_city(city)
    await _developer_show_users_list(callback, users, f"🏙️ <b>Стажери: {city}</b>", f"city:{city}", 0)

async def developer_users_pagination_handler(callback: CallbackQuery):
    try:
        parts = callback.data.split(":")
        # Formats:
        # dev_users_pag:mode:page
        # dev_users_pag:city:CityName:page
        # dev_users_page:page (legacy)
        
        if len(parts) == 4 and parts[1] == "city":
            city = parts[2]
            page = int(parts[3])
            users = await get_users_by_city(city)
            await _developer_show_users_list(callback, users, f"🏙️ <b>Стажери: {city}</b>", f"city:{city}", page)
            return
        elif len(parts) == 3:
            mode = parts[1]
            page = int(parts[2])
        else:
            mode = "all"
            page = int(parts[1])
            
    except (ValueError, IndexError):
        return

    if mode == "active":
        await developer_active_users(callback, page)
    elif mode == "inactive":
        await developer_inactive_users(callback, page)
    elif mode == "interns":
        await developer_list_interns(callback, page)
    elif mode == "workers":
        await developer_list_workers(callback, page)
    else:
        await developer_list_users(callback, page)


async def developer_request_user_search(callback: CallbackQuery, state: FSMContext):
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access or (is_territorial and not is_admin):
        await callback.answer("Вам недоступна ця функція.", show_alert=True)
        return
    await state.set_state(DeveloperStates.waiting_user_search)
    await _edit_or_answer(
        callback.message,
        "🔍 <b>Пошук користувача</b>\n\n"
        "Введіть один з варіантів:\n"
        "• ID користувача (наприклад, <code>123456</code>)\n"
        "• Username (наприклад, <code>@username</code>)\n"
        "• ПІБ (наприклад, <code>Іванов Іван</code>)\n\n"
        "<i>💡 Пошук не чутливий до регістру</i>",
        reply_markup=await _get_users_menu_kb(callback.from_user.id),
    )
    await callback.answer()


async def developer_process_user_search(message: Message, state: FSMContext):
    query = (message.text or "").strip()
    users = []
    
    if query.startswith("@"):
        user = await get_user_by_username(query)
        if user:
            users.append(user)
    elif query.isdigit():
        user = await get_user_details(int(query))
        if user:
            users.append(user)
    else:
        # Search by full name if it's not a username or ID
        users = await get_users_by_full_name(query)

    if not users:
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🔍 Новий пошук", callback_data="dev_user_search")],
            [InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_users_menu")]
        ])
        
        await message.answer(
            "❌ <b>Користувача не знайдено</b>\n\n"
            f"Ваш запит: <code>{query}</code>\n\n"
            "💡 <b>Поради:</b>\n"
            "• Перевірте правильність написання\n"
            "• Спробуйте ввести частину імені\n"
            "• Використайте ID якщо він відомий\n"
            "• Або натисніть /cancel для скасування",
            reply_markup=kb
        )
        return

    if len(users) == 1:
        # If only one user is found, show their profile directly
        user = users[0]
        intern_id = int(user["user_id"])
        report = await get_user_days_report(intern_id)
        
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [
                InlineKeyboardButton(text="📅 Навчальні дні", callback_data=f"dev_days_manage:{intern_id}"),
                InlineKeyboardButton(text="✍️ Змінити", callback_data=f"dev_user_modify:{intern_id}")
            ],
            [
                InlineKeyboardButton(text="❌ Видалити", callback_data=f"dev_user_delete_confirm:{intern_id}"),
                InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_users_menu")
            ]
        ])
        
        await message.answer(report, reply_markup=kb)
    else:
        # If multiple users are found, show a list
        lines = [f"👥 <b>Знайдено {len(users)} користувачів:</b>\n"]
        
        # Показуємо максимум 20 користувачів
        display_users = users[:20]
        for user in display_users:
            full_name = user.get('full_name', 'Без імені')
            user_id = user['user_id']
            username = user.get('username', '')
            
            user_info = f"• <b>{full_name}</b> (ID: <code>{user_id}</code>)"
            if username:
                user_info += f" | @{username}"
            lines.append(user_info)
        
        if len(users) > 20:
            lines.append(f"\n<i>... та ще {len(users) - 20} користувачів</i>")
        
        lines.append("\n💡 <b>Уточніть пошук:</b>")
        lines.append("• Введіть повне ім'я")
        lines.append("• Або використайте ID конкретного користувача")
        
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🔍 Новий пошук", callback_data="dev_user_search")],
            [InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_users_menu")]
        ])
        
        await message.answer("\n".join(lines), reply_markup=kb)

    await state.clear()


async def developer_request_delete_user(callback: CallbackQuery, state: FSMContext):
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access or (is_territorial and not is_admin):
        if has_access: await callback.answer("Вам недоступна ця функція.", show_alert=True)
        return
    await state.set_state(DeveloperStates.waiting_delete_user)
    await _edit_or_answer(
        callback.message,
        "❌ Введіть ID користувача, якого потрібно видалити:",
        reply_markup=await _get_users_menu_kb(callback.from_user.id),
    )
    await callback.answer()


async def developer_process_delete_user(message: Message, state: FSMContext):
    try:
        user_id = int(message.text.strip())
    except (ValueError, AttributeError):
        await message.answer("❌ ID повинен бути числом. Спробуйте ще раз.")
        return

    user = await get_user_details(user_id)
    if not user:
        await message.answer(f"❌ Користувача з ID {user_id} не знайдено.")
        await state.clear()
        return

    await delete_user(user_id)
    await message.answer(f"✅ Користувача <b>{user.get('full_name', 'Без імені')}</b> (ID: {user_id}) видалено.", reply_markup=await _get_users_menu_kb(message.from_user.id))
    await state.clear()


# ---------- Mentors section ----------

def _managers_menu_keyboard(is_admin: bool = True, is_territorial: bool = False) -> InlineKeyboardMarkup:
    keyboard = []
    keyboard.append([InlineKeyboardButton(text="📋 Список керівників", callback_data="dev_managers_list")])
    if is_admin or not is_territorial:
        keyboard.append([InlineKeyboardButton(text="➕ Додати керівника", callback_data="dev_manager_add")])
        keyboard.append([InlineKeyboardButton(text="🗑 Видалити керівника", callback_data="dev_manager_delete")])
    keyboard.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="developer_menu")])
    return InlineKeyboardMarkup(inline_keyboard=keyboard)


async def developer_managers_menu(callback: CallbackQuery):
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access:
        return
    await _edit_or_answer(
        callback.message,
        "🧑‍🏫 <b>Керівники</b>\nОберіть дію:",
        reply_markup=_managers_menu_keyboard(is_admin, is_territorial),
    )
    await callback.answer()


async def developer_list_managers(callback: CallbackQuery):
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access:
        return
    managers = await get_all_managers()
    
    if is_territorial and not is_admin:
        my_managers = await get_managers_by_responsible(callback.from_user.id)
        my_mgr_ids = [m['uid'] for m in my_managers]
        managers = [m for m in managers if m['uid'] in my_mgr_ids]

    if not managers:
        await _edit_or_answer(
            callback.message,
            "ℹ️ У системі поки немає керівників.",
            reply_markup=_managers_menu_keyboard(is_admin, is_territorial),
        )
        await callback.answer()
        return

    lines = ["🧑‍🏫 <b>Список керівників</b>", ""]
    for idx, manager in enumerate(managers, 1):
        lines.append(
            f"{idx}. <b>{manager.get('full_name', 'Без імені')}</b> (uid: {manager.get('uid')})\n"
            f"   @{manager.get('username', 'немає')} · Посада: {manager.get('process', 'Не вказано')}"
        )
        if idx >= 50:
            lines.append("… Показано перших 50 керівників.")
            break
    await _edit_or_answer(
        callback.message,
        "\n".join(lines),
        reply_markup=_managers_menu_keyboard(is_admin, is_territorial),
    )
    await callback.answer()


async def developer_request_add_manager(callback: CallbackQuery, state: FSMContext):
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access or (is_territorial and not is_admin):
        if has_access: await callback.answer("Вам недоступна ця функція.", show_alert=True)
        return
    await state.set_state(DeveloperStates.waiting_add_manager)
    await _edit_or_answer(
        callback.message,
        "➕ Введіть дані керівника у форматі <code>ID;посада</code> (наприклад: <code>123456;Керівник магазину</code>).",
        reply_markup=_managers_menu_keyboard(is_admin, is_territorial),
    )
    await callback.answer()


async def developer_process_add_manager(message: Message, state: FSMContext):
    payload = (message.text or "")
    if ";" not in payload:
        await message.answer("❌ Невірний формат. Використайте <code>ID;посада</code>.")
        return
    user_id_part, process_part = payload.split(";", 1)
    try:
        uid = int(user_id_part.strip())
    except ValueError:
        await message.answer("❌ ID повинен бути числом.")
        return
    process = process_part.strip()
    if not process:
        await message.answer("❌ Посада не може бути порожньою.")
        return

    user = await get_user_details(uid)
    if not user:
        await message.answer("❌ Користувача з таким ID не знайдено.")
        await state.clear()
        return
    existing = await get_manager_by_uid(uid)
    if existing:
        await message.answer("⚠️ Користувач вже є керівником.")
        await state.clear()
        return
    await add_manager(uid, process)
    await message.answer(f"✅ Користувача {user.get('full_name', 'Без імені')} призначено керівником ({process}).", reply_markup=_managers_menu_keyboard())
    await state.clear()


async def developer_request_delete_manager(callback: CallbackQuery, state: FSMContext):
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access or (is_territorial and not is_admin):
        if has_access: await callback.answer("Вам недоступна ця функція.", show_alert=True)
        return
    await state.set_state(DeveloperStates.waiting_delete_manager)
    await _edit_or_answer(
        callback.message,
        "🗑 Введіть ID керівника для видалення:",
        reply_markup=_managers_menu_keyboard(is_admin, is_territorial),
    )
    await callback.answer()


async def developer_process_delete_manager(message: Message, state: FSMContext):
    try:
        uid = int(message.text.strip())
    except (ValueError, AttributeError):
        await message.answer("❌ ID повинен бути числом.")
        return

    manager = await get_manager_by_uid(uid)
    if not manager:
        await message.answer("❌ Керівника з таким ID не знайдено.")
        await state.clear()
        return

    await delete_manager_by_uid(uid)
    await message.answer(
        f"✅ Керівник {manager.get('full_name', 'Без імені')} (ID: {uid}) видалений.",
        reply_markup=_managers_menu_keyboard()
    )
    await state.clear()


# ---------- Days section ----------

def _days_menu_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="👀 Статус", callback_data="dev_days_status")],
        [InlineKeyboardButton(text="🔓 Відкрити дні", callback_data="dev_days_open")],
        [InlineKeyboardButton(text="🔒 Закрити дні", callback_data="dev_days_close")],
        [InlineKeyboardButton(text="🔄 Скинути прогрес", callback_data="dev_days_reset")],
        [InlineKeyboardButton(text="⬅️ Назад", callback_data="developer_menu")]
    ])


async def developer_days_menu(callback: CallbackQuery):
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access or (is_territorial and not is_admin):
        if has_access: await callback.answer("Вам недоступна ця функція.", show_alert=True)
        return
    await _edit_or_answer(
        callback.message,
        "📅 <b>Навчальні дні</b>\nОберіть дію:",
        reply_markup=_days_menu_keyboard(),
    )
    await callback.answer()


async def developer_request_days_status(callback: CallbackQuery, state: FSMContext):
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access or (is_territorial and not is_admin):
        if has_access: await callback.answer("Вам недоступна ця функція.", show_alert=True)
        return
    await state.set_state(DeveloperStates.waiting_days_status)
    await _edit_or_answer(
        callback.message,
        "👀 Введіть ID стажера для перегляду статусу:",
        reply_markup=_days_menu_keyboard(),
    )
    await callback.answer()


async def developer_process_days_status(message: Message, state: FSMContext):
    try:
        user_id = int(message.text.strip())
    except (ValueError, AttributeError):
        await message.answer("❌ ID повинен бути числом.")
        return
    try:
        report = await get_user_days_report(user_id)
    except ValueError as exc:
        await message.answer(f"❌ {exc}")
        await state.clear()
        return
    await _send_in_chunks(message, report, reply_markup=_days_menu_keyboard())
    await state.clear()


async def developer_request_days_action(callback: CallbackQuery, state: FSMContext, action: str):
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access or (is_territorial and not is_admin):
        if has_access: await callback.answer("Вам недоступна ця функція.", show_alert=True)
        return
    action_map = {
        "open": "🔓 Введіть дані у форматі <code>ID;дні</code> (приклад: <code>123456;3-5</code>) для відкриття.",
        "close": "🔒 Введіть <code>ID;дні</code>, які потрібно закрити (день 1 закрито бути не може).",
        "reset": "🔄 Введіть <code>ID;дні</code>, які потрібно скинути."
    }
    await state.set_state(getattr(DeveloperStates, f"waiting_days_{action}"))
    await _edit_or_answer(callback.message, action_map[action], reply_markup=_days_menu_keyboard())
    await callback.answer()


async def developer_request_days_open(callback: CallbackQuery, state: FSMContext):
    await developer_request_days_action(callback, state, "open")


async def developer_request_days_close(callback: CallbackQuery, state: FSMContext):
    await developer_request_days_action(callback, state, "close")


async def developer_request_days_reset(callback: CallbackQuery, state: FSMContext):
    await developer_request_days_action(callback, state, "reset")


async def _parse_user_days_payload(text: str):
    if ";" not in text:
        raise ValueError("Невірний формат. Використайте <code>ID;дні</code>.")
    user_part, days_part = text.split(";", 1)
    user_id = int(user_part.strip())
    days = parse_day_input(days_part.strip())
    if not days:
        raise ValueError("Невірний формат днів. Приклад: 3 або 3-5.")
    return user_id, days


async def developer_process_days_open(message: Message, state: FSMContext):
    try:
        user_id, days = await _parse_user_days_payload(message.text.strip())
        response = await open_days_for_user(user_id, days)
    except ValueError as exc:
        await message.answer(f"❌ {exc}")
        return
    await message.answer(response, reply_markup=_days_menu_keyboard())
    await state.clear()


async def developer_process_days_close(message: Message, state: FSMContext):
    try:
        user_id, days = await _parse_user_days_payload(message.text.strip())
        response = await close_days_for_user(user_id, days)
    except ValueError as exc:
        await message.answer(f"❌ {exc}")
        return
    await message.answer(response, reply_markup=_days_menu_keyboard())
    await state.clear()


async def developer_process_days_reset(message: Message, state: FSMContext):
    try:
        user_id, days = await _parse_user_days_payload(message.text.strip())
        response = await reset_days_for_user(user_id, days)
    except ValueError as exc:
        await message.answer(f"❌ {exc}")
        return
    await message.answer(response, reply_markup=_days_menu_keyboard())
    await state.clear()


# ---------- Developer management ----------

# ---------- Managers team (unified Dev + HR) ----------

async def developer_manage_managers_menu(callback: CallbackQuery, state: FSMContext, page: int = 0):
    """Об'єднане меню керівників з пагінацією."""
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access:
        return
    user_id = callback.from_user.id
    text, kb = await _build_managers_team_view(is_admin, is_territorial, user_id, page=page, bot=callback.bot)
    await _send_or_edit_admin_photo(callback, "admin_spus.jpg", text, kb)
    await _remember_panel(state, "managers_panel", callback.message)

async def developer_managers_pagination_handler(callback: CallbackQuery, state: FSMContext):
    """Обробник пагінації для списку керівників."""
    try:
        page = int(callback.data.split(":")[1])
    except (ValueError, IndexError):
        page = 0
    await developer_manage_managers_menu(callback, state, page=page)


async def developer_request_add_manager(callback: CallbackQuery, state: FSMContext):
    """Starts the process of adding a manager by requesting their ID."""
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access or (is_territorial and not is_admin):
        if has_access: await callback.answer("Вам недоступна ця функція.", show_alert=True)
        return
    await state.set_state(DeveloperStates.waiting_add_manager_id)
    await _edit_or_answer(
        callback.message,
        "👔 Введіть ID користувача, якого потрібно призначити керівником:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="❌ Скасувати", callback_data="dev_manage_managers")]
        ])
    )
    await callback.answer()

async def developer_process_add_manager_city(callback: CallbackQuery, state: FSMContext):
    """Processes the city selection and shows the shop selection menu."""
    try:
        city = callback.data.split(":", 1)[1]
    except IndexError:
        return await callback.answer("Помилка даних міста.", show_alert=True)
        
    await state.update_data(manager_city=city, selected_shops=[])
    await state.set_state(DeveloperStates.waiting_add_manager_shops)
    
    kb = await _build_manager_shops_keyboard(city, [])
    await _edit_or_answer(
        callback.message,
        f"Місто: {city}.\nТепер оберіть магазини. Можна обрати до 5.",
        reply_markup=kb
    )
    await callback.answer()

async def _build_manager_shops_keyboard(city: str, selected_shops: list) -> InlineKeyboardMarkup:
    """Builds the keyboard for shop multi-selection for a specific city."""
    shops_in_city = AVAILABLE_SHOPS.get(city, [])
    
    buttons = []
    for i in range(0, len(shops_in_city), 2):
        row = []
        shop1 = shops_in_city[i]
        text1 = f"✅ {shop1}" if shop1 in selected_shops else shop1
        row.append(InlineKeyboardButton(text=text1, callback_data=f"dev_mgr_shop_toggle:{i}"))
        
        if i + 1 < len(shops_in_city):
            shop2 = shops_in_city[i+1]
            text2 = f"✅ {shop2}" if shop2 in selected_shops else shop2
            row.append(InlineKeyboardButton(text=text2, callback_data=f"dev_mgr_shop_toggle:{i+1}"))
        buttons.append(row)
    
    buttons.append([InlineKeyboardButton(text="✅ Готово", callback_data="dev_mgr_shop_done")])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад до міст", callback_data="dev_add_manager")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)

async def developer_process_add_manager_id(message: Message, state: FSMContext):
    """Processes the manager ID and asks for their Full Name."""
    if not message.text or not message.text.isdigit():
        await message.answer("❌ ID повинен бути числом. Спробуйте ще раз.")
        return
    
    user_id = int(message.text)
    await state.update_data(new_manager_id=user_id, selected_shops=[])
    await state.set_state(DeveloperStates.waiting_add_manager_name)
    
    await message.answer(
        f"ID: {user_id}. Тепер введіть ПІБ керівника:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="❌ Скасувати", callback_data="dev_manage_managers")]
        ])
    )

async def developer_process_add_manager_name(message: Message, state: FSMContext):
    """Processes the manager Name and shows the city selection menu."""
    full_name = message.text.strip()
    if len(full_name.split()) < 2:
        await message.answer("Будь ласка, введіть повне ім'я та прізвище (мінімум 2 слова).")
        return

    await state.update_data(new_manager_name=full_name)
    await state.set_state(DeveloperStates.waiting_add_manager_city)
    
    buttons = []
    for city in AVAILABLE_CITIES:
        buttons.append([InlineKeyboardButton(text=city, callback_data=f"dev_mgr_city_select:{city}")])
    buttons.append([InlineKeyboardButton(text="❌ Скасувати", callback_data="dev_manage_managers")])
    
    await message.answer(
        f"Керівник: <b>{full_name}</b>.\nТепер оберіть місто, до якого належать його магазини:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )

async def developer_process_shop_selection(callback: CallbackQuery, state: FSMContext):
    """Handles toggling a shop in the selection menu."""
    data = await state.get_data()
    city = data.get("manager_city")
    if not city:
        # Якщо місто не знайдено у стані, це помилка флоу. Повертаємося в головне меню керівників.
        await callback.answer("Помилка: місто не обрано. Поверніться до меню додавання керівника.", show_alert=True)
        await state.clear()
        await developer_manage_managers_menu(callback, state) # Go back to managers menu
        return
        
    try:
        shop_index = int(callback.data.split(":", 1)[1])
        shops_in_city = AVAILABLE_SHOPS.get(city, [])
        shop_name = shops_in_city[shop_index]
    except (IndexError, ValueError):
        return await callback.answer("Помилка даних магазину.", show_alert=True)

    selected_shops = data.get("selected_shops", [])
    
    if shop_name in selected_shops:
        selected_shops.remove(shop_name)
    else:
        if len(selected_shops) >= 5:
            await callback.answer("⚠️ Можна обрати максимум 5 магазинів.", show_alert=True)
            return
        selected_shops.append(shop_name)
        
    await state.update_data(selected_shops=selected_shops)
    
    kb = await _build_manager_shops_keyboard(city, selected_shops)
    await callback.message.edit_reply_markup(reply_markup=kb)
    await callback.answer()

async def developer_finish_shop_selection(callback: CallbackQuery, state: FSMContext):
    """Finalizes adding the manager with the selected shops."""
    data = await state.get_data()
    user_id = data.get("new_manager_id")
    selected_shops = data.get("selected_shops", [])
    full_name = data.get("new_manager_name", "Керівник")
    city = data.get("manager_city")
    
    if not user_id:
        await callback.answer("Помилка: ID керівника не знайдено.", show_alert=True)
        return await developer_manage_managers_menu(callback, state)

    # Get username if possible, but rely on provided full_name
    try:
        chat_info = await callback.bot.get_chat(user_id)
        username = chat_info.username
    except Exception:
        username = ""

    await register_user(user_id, username=username, full_name=full_name)
    await add_manager(
        uid=user_id,
        process="Керівник",
        full_name=full_name,
        username=username,
        shops=selected_shops,
        city=city
    )
    
    await state.clear()
    await callback.answer("✅ Керівника успішно додано/оновлено!", show_alert=True)
    await developer_manage_managers_menu(callback, state)


async def developer_remove_manager_menu(callback: CallbackQuery, state: FSMContext):
    """Меню видалення керівника."""
    if not await _ensure_developer(callback):
        return
    
    hrs = await get_all_kerivnyky()
    if not hrs:
        await callback.answer("Немає керівників для видалення.", show_alert=True)
        return
    
    buttons = []
    for hr in hrs:
        name, _ = await _format_identity(hr["uid"], hr.get("full_name"), hr.get("username"))
        buttons.append([InlineKeyboardButton(
            text=f"❌ {name}",
            callback_data=f"mgr_remove:{hr['uid']}"
        )])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_manage_managers")])
    
    kb = InlineKeyboardMarkup(inline_keyboard=buttons)
    await _edit_or_answer(callback.message, "🗑 Оберіть керівника для видалення:", reply_markup=kb)
    await callback.answer()


async def developer_remove_manager(callback: CallbackQuery, state: FSMContext):
    """Видалення керівника."""
    if not await _ensure_developer(callback):
        return
    
    try:
        user_id = int(callback.data.split(":")[1])
    except (ValueError, IndexError):
        await callback.answer("Некоректний ID.", show_alert=True)
        return
    
    await delete_manager_by_uid(user_id)
    await callback.answer(f"Керівника {user_id} видалено.", show_alert=True)
    await developer_manage_managers_menu(callback, state)


async def manager_cancel_add(callback: CallbackQuery, state: FSMContext):
    """Скасування додавання керівника."""
    if not await _ensure_developer(callback):
        return
    await state.set_state(None)
    await state.update_data(mgr_prompt=None)
    await developer_manage_managers_menu(callback, state)
    await callback.answer("Додавання скасовано.", show_alert=True)


async def manager_team_back(callback: CallbackQuery, state: FSMContext):
    """Повернення до списку керівників."""
    await developer_manage_managers_menu(callback, state)


# ==============================================================================
# УПРАВЛІННЯ ТА ПОШУК КЕРІВНИКІВ (ПУНКТ 5)
# ==============================================================================

async def _build_manager_card(manager_uid: int, is_admin: bool, is_territorial: bool, is_observer: bool = False, bot: Optional[Bot] = None):
    mgr = await get_manager_by_uid(manager_uid)
    if not mgr:
        return None
    
    async with aiosqlite.connect(DB_PATH) as db:
        c_interns = await db.execute(
            "SELECT COUNT(*) FROM users WHERE manager_id = ? AND (status IS NULL OR status != 'Працівник')",
            (manager_uid,)
        )
        interns_count = (await c_interns.fetchone())[0]
        c_workers = await db.execute(
            "SELECT COUNT(*) FROM users WHERE manager_id = ? AND status = 'Працівник'",
            (manager_uid,)
        )
        workers_count = (await c_workers.fetchone())[0]

    name, username = await _format_identity(mgr["uid"], mgr.get("full_name"), mgr.get("username"), bot=bot)
    city = mgr.get("city") or "Не вказано"
    shops_list = mgr.get("shops") or []
    if isinstance(shops_list, str):
        try:
            shops_list = json.loads(shops_list)
        except Exception:
            shops_list = [shops_list] if shops_list else []
    shops_str = ", ".join(shops_list) if shops_list else "Не призначено"
    
    resp_uid = mgr.get("responsible_uid")
    if resp_uid:
        resp_mgr = await get_manager_by_uid(resp_uid)
        if resp_mgr:
            t_type = resp_mgr.get("territorial_type") or "ТЗ"
            resp_label = f"🗺 Територіал: {resp_mgr.get('full_name', '')} ({t_type})"
        else:
            resp_label = f"🗺 Територіал (UID: {resp_uid})"
    else:
        resp_label = "👑 Адміністратор (пряме підпорядкування)"

    status_tag = ""
    if mgr.get("status") == "fired":
        status_tag = "\n⚠️ <b>Статус:</b> 🚫 Звільнений"

    text = (
        f"👔 <b>Керівник:</b> {html.escape(name)}\n"
        f"👤 <b>Username:</b> {html.escape(username)}\n"
        f"🆔 <b>Telegram ID:</b> <code>{mgr['uid']}</code>\n"
        f"🏙 <b>Місто:</b> {html.escape(city)}\n"
        f"🏪 <b>Магазини ({len(shops_list)}/5):</b> {html.escape(shops_str)}\n"
        f"🗺 <b>Відповідальний:</b> {html.escape(resp_label)}"
        f"{status_tag}\n\n"
        f"📊 <b>Підлеглі:</b>\n"
        f"• 🎓 Стажерів: {interns_count}\n"
        f"• 👷 Працівників: {workers_count}\n"
        f"• 👥 Всього: {interns_count + workers_count}"
    )

    buttons = []
    if mgr.get("status") != "fired" and not is_observer:
        buttons.append([InlineKeyboardButton(text="✏️ Змінити", callback_data=f"dev_mgr_edit_menu:{manager_uid}")])
        if is_admin:
            buttons.append([InlineKeyboardButton(text="🔄 Перепривʼязати", callback_data=f"dev_mgr_reassign_menu:{manager_uid}")])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад до списку", callback_data="dev_manage_managers")])
    
    return text, InlineKeyboardMarkup(inline_keyboard=buttons)


async def _build_edit_shops_keyboard(city: str, selected_shops: list, manager_uid: int) -> InlineKeyboardMarkup:
    """Builds keyboard for editing manager shops with multi-select."""
    from bot.constants import AVAILABLE_SHOPS
    shops_in_city = AVAILABLE_SHOPS.get(city, [])
    
    buttons = []
    for i in range(0, len(shops_in_city), 2):
        row = []
        shop1 = shops_in_city[i]
        text1 = f"✅ {shop1}" if shop1 in selected_shops else shop1
        row.append(InlineKeyboardButton(text=text1, callback_data=f"dev_mgr_eshop_toggle:{i}"))
        
        if i + 1 < len(shops_in_city):
            shop2 = shops_in_city[i+1]
            text2 = f"✅ {shop2}" if shop2 in selected_shops else shop2
            row.append(InlineKeyboardButton(text=text2, callback_data=f"dev_mgr_eshop_toggle:{i+1}"))
        buttons.append(row)
    
    buttons.append([InlineKeyboardButton(text="✅ Готово", callback_data=f"dev_mgr_eshop_done:{manager_uid}")])
    buttons.append([InlineKeyboardButton(text="⬅️ Скасувати", callback_data=f"dev_mgr_view:{manager_uid}")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)


async def dev_mgr_view(callback: CallbackQuery, state: FSMContext):
    """Відображає картку керівника з його фото."""
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access:
        return
    try:
        manager_uid = int(callback.data.split(":")[1])
    except (ValueError, IndexError):
        return await callback.answer("Помилка ідентифікатора.", show_alert=True)
        
    is_obs = await is_observer_user(callback.from_user.id)
    card_data = await _build_manager_card(manager_uid, is_admin, is_territorial, is_observer=is_obs, bot=callback.bot)
    if not card_data:
        await callback.answer("Керівника не знайдено!", show_alert=True)
        return
    text, kb = card_data
    photo_input = await get_user_avatar_input(callback.bot, manager_uid)
    await _send_or_edit_card_photo(callback, photo_input, text, reply_markup=kb)


async def dev_mgr_search_start(callback: CallbackQuery, state: FSMContext):
    """Початок FSM-пошуку керівника."""
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access:
        return
    await state.set_state(DeveloperStates.waiting_search_manager)
    await _edit_or_answer(
        callback.message,
        "🔍 <b>Пошук керівника</b>\nВведіть ПІБ, прізвище, @username або Telegram ID:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="❌ Скасувати", callback_data="dev_manage_managers")]
        ])
    )
    await callback.answer()


async def dev_mgr_search_process(message: Message, state: FSMContext):
    """Обробляє введення пошукового запиту для керівників."""
    user_id = message.from_user.id
    is_admin = await is_developer_user(user_id)
    is_territorial = await is_territorial_user(user_id)
    
    if not is_admin and not is_territorial:
        await state.clear()
        return

    query = message.text.strip()
    results = await search_managers(query, is_admin=is_admin, territorial_uid=user_id if is_territorial else None)
    
    if not results:
        await message.answer(
            f"❌ За запитом «<b>{html.escape(query)}</b>» керівників не знайдено.",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="🔍 Шукати знову", callback_data="dev_mgr_search_start")],
                [InlineKeyboardButton(text="⬅️ До списку керівників", callback_data="dev_manage_managers")]
            ])
        )
        return
        
    if len(results) == 1:
        await state.clear()
        card_data = await _build_manager_card(results[0]["uid"], is_admin, is_territorial)
        if card_data:
            text, kb = card_data
            await message.answer(text, reply_markup=kb)
        return
        
    await state.clear()
    buttons = []
    for m in results:
        name, _ = await _format_identity(m["uid"], m.get("full_name"), m.get("username"))
        shops = ", ".join(m.get("shops") or []) or "—"
        city = m.get("city") or ""
        tag = f" ({city})" if city else ""
        buttons.append([
            InlineKeyboardButton(text=f"👔 {name}{tag} | {shops}", callback_data=f"dev_mgr_view:{m['uid']}")
        ])
    buttons.append([InlineKeyboardButton(text="🔍 Новий пошук", callback_data="dev_mgr_search_start")])
    buttons.append([InlineKeyboardButton(text="⬅️ До списку керівників", callback_data="dev_manage_managers")])
    
    await message.answer(
        f"🔍 Знайдено <b>{len(results)}</b> керівників за запитом «<b>{html.escape(query)}</b>»:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )


async def dev_mgr_edit_menu(callback: CallbackQuery, state: FSMContext):
    """Підменю редагування керівника."""
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access:
        return
    manager_uid = int(callback.data.split(":")[1])
    mgr = await get_manager_by_uid(manager_uid)
    if not mgr:
        await callback.answer("Керівника не знайдено!", show_alert=True)
        return
    
    text = f"✏️ <b>Редагування керівника {html.escape(mgr.get('full_name') or str(manager_uid))}</b>\nОберіть дію:"
    buttons = [
        [InlineKeyboardButton(text="🏪 Змінити магазин", callback_data=f"dev_mgr_ch_shops:{manager_uid}")],
        [InlineKeyboardButton(text="👤 Змінити ПІБ", callback_data=f"dev_mgr_ch_name:{manager_uid}")],
        [InlineKeyboardButton(text="🚫 Звільнити", callback_data=f"dev_mgr_fire_ask:{manager_uid}")],
        [InlineKeyboardButton(text="⬅️ Назад до картки", callback_data=f"dev_mgr_view:{manager_uid}")]
    ]
    await _edit_or_answer(callback.message, text, reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons))
    await callback.answer()


async def dev_mgr_change_name_start(callback: CallbackQuery, state: FSMContext):
    """Початок зміни ПІБ керівника."""
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access:
        return
    manager_uid = int(callback.data.split(":")[1])
    mgr = await get_manager_by_uid(manager_uid)
    if not mgr:
        await callback.answer("Керівника не знайдено!", show_alert=True)
        return
    
    await state.update_data(edit_mgr_uid=manager_uid)
    await state.set_state(DeveloperStates.waiting_edit_manager_name)
    await _edit_or_answer(
        callback.message,
        f"Введіть нове ПІБ для керівника <b>{html.escape(mgr.get('full_name') or '')}</b> (мінімум 2 слова):",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="❌ Скасувати", callback_data=f"dev_mgr_view:{manager_uid}")]
        ])
    )
    await callback.answer()


async def dev_mgr_change_name_process(message: Message, state: FSMContext):
    """Обробка нового ПІБ керівника."""
    data = await state.get_data()
    manager_uid = data.get("edit_mgr_uid")
    if not manager_uid:
        await state.clear()
        return
    
    new_name = message.text.strip()
    if len(new_name.split()) < 2:
        await message.answer("Будь ласка, введіть повне ім'я та прізвище (мінімум 2 слова).")
        return
    
    await update_manager_name(manager_uid, new_name)
    await state.clear()
    
    card_data = await _build_manager_card(manager_uid, True, False)
    if card_data:
        text, kb = card_data
        await message.answer(f"✅ ПІБ оновлено на <b>{html.escape(new_name)}</b>!\n\n{text}", reply_markup=kb)


async def dev_mgr_change_shops_start(callback: CallbackQuery, state: FSMContext):
    """Початок зміни магазинів керівника з тихим перенесенням старих стажерів."""
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access:
        return
    manager_uid = int(callback.data.split(":")[1])
    mgr = await get_manager_by_uid(manager_uid)
    if not mgr:
        await callback.answer("Керівника не знайдено!", show_alert=True)
        return
    
    city = mgr.get("city") or "Хмельницький"
    # Тихо переносимо стажерів старого магазину на територіала або адміна
    await transfer_users_to_territorial_or_admin(manager_uid)
    
    await state.update_data(edit_mgr_uid=manager_uid, manager_city=city, selected_shops=[])
    await state.set_state(DeveloperStates.waiting_edit_manager_shops)
    
    kb = await _build_edit_shops_keyboard(city, [], manager_uid)
    await _edit_or_answer(
        callback.message,
        f"🏪 Зміна магазинів для керівника <b>{html.escape(mgr.get('full_name') or '')}</b> (місто {city}):\n"
        f"Старі стажери переведені до Територіала/Адміна.\nОберіть нові магазини (до 5):",
        reply_markup=kb
    )
    await callback.answer()


async def dev_mgr_edit_shop_toggle(callback: CallbackQuery, state: FSMContext):
    """Перемикання магазину під час редагування."""
    data = await state.get_data()
    city = data.get("manager_city")
    manager_uid = data.get("edit_mgr_uid")
    if not city or not manager_uid:
        await state.clear()
        await callback.answer("Помилка стану.", show_alert=True)
        return
        
    try:
        shop_index = int(callback.data.split(":", 1)[1])
        from bot.constants import AVAILABLE_SHOPS
        shops_in_city = AVAILABLE_SHOPS.get(city, [])
        shop_name = shops_in_city[shop_index]
    except (IndexError, ValueError):
        return await callback.answer("Помилка даних магазину.", show_alert=True)

    selected_shops = data.get("selected_shops", [])
    if shop_name in selected_shops:
        selected_shops.remove(shop_name)
    else:
        if len(selected_shops) >= 5:
            await callback.answer("⚠️ Можна обрати максимум 5 магазинів.", show_alert=True)
            return
        selected_shops.append(shop_name)
        
    await state.update_data(selected_shops=selected_shops)
    kb = await _build_edit_shops_keyboard(city, selected_shops, manager_uid)
    await callback.message.edit_reply_markup(reply_markup=kb)
    await callback.answer()


async def dev_mgr_edit_shop_done(callback: CallbackQuery, state: FSMContext):
    """Завершення зміни магазинів керівника з запитом на переведення стажерів нових магазинів."""
    manager_uid = int(callback.data.split(":")[1])
    data = await state.get_data()
    city = data.get("manager_city") or "Хмельницький"
    selected_shops = data.get("selected_shops", [])
    
    await update_manager_shops(manager_uid, city, selected_shops)
    
    # Перевіряємо чи є користувачі в обраних магазинах
    interns_c, workers_c = await count_users_in_shops(city, selected_shops)
    total_new = interns_c + workers_c
    
    if total_new > 0:
        await state.clear()
        shops_joined = ", ".join(selected_shops)
        text = (
            f"✅ Магазини оновлено: <b>{html.escape(shops_joined)}</b>.\n\n"
            f"👥 У цих магазинах знайдено <b>{interns_c}</b> стажерів та <b>{workers_c}</b> працівників.\n"
            f"Перевести їх під керівництво цього керівника?"
        )
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="✅ Так, перевести", callback_data=f"dev_mgr_tr_yes:{manager_uid}")],
            [InlineKeyboardButton(text="❌ Ні, залишити як є", callback_data=f"dev_mgr_view:{manager_uid}")]
        ])
        await _edit_or_answer(callback.message, text, reply_markup=kb)
        await callback.answer()
    else:
        await state.clear()
        await callback.answer("✅ Магазини успішно оновлено!", show_alert=True)
        has_access, is_admin, is_territorial = await _check_access(callback)
        card_data = await _build_manager_card(manager_uid, is_admin, is_territorial)
        if card_data:
            text, kb = card_data
            await _edit_or_answer(callback.message, text, reply_markup=kb)


async def dev_mgr_transfer_shops_yes(callback: CallbackQuery, state: FSMContext):
    """Підтвердження переведення стажерів/працівників нових магазинів до керівника."""
    manager_uid = int(callback.data.split(":")[1])
    mgr = await get_manager_by_uid(manager_uid)
    if mgr:
        city = mgr.get("city") or "Хмельницький"
        shops = mgr.get("shops") or []
        transferred = await transfer_users_by_shops(city, shops, manager_uid)
        await callback.answer(f"✅ Переведено {transferred} користувачів!", show_alert=True)
        
    has_access, is_admin, is_territorial = await _check_access(callback)
    card_data = await _build_manager_card(manager_uid, is_admin, is_territorial)
    if card_data:
        text, kb = card_data
        await _edit_or_answer(callback.message, text, reply_markup=kb)


async def dev_mgr_fire_ask(callback: CallbackQuery, state: FSMContext):
    """Діалог підтвердження звільнення керівника."""
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access:
        return
    manager_uid = int(callback.data.split(":")[1])
    mgr = await get_manager_by_uid(manager_uid)
    if not mgr:
        await callback.answer("Керівника не знайдено!", show_alert=True)
        return
        
    name = mgr.get("full_name") or str(manager_uid)
    async with aiosqlite.connect(DB_PATH) as db:
        c_interns = await db.execute("SELECT COUNT(*) FROM users WHERE manager_id = ? AND (status IS NULL OR status != 'Працівник')", (manager_uid,))
        interns_count = (await c_interns.fetchone())[0]
        c_workers = await db.execute("SELECT COUNT(*) FROM users WHERE manager_id = ? AND status = 'Працівник'", (manager_uid,))
        workers_count = (await c_workers.fetchone())[0]

    target_name = "Адміністратора"
    if mgr.get("responsible_uid"):
        r_mgr = await get_manager_by_uid(mgr["responsible_uid"])
        if r_mgr:
            target_name = f"Територіала {r_mgr.get('full_name', '')}"
    else:
        city_t = await get_territorials_by_city(mgr.get("city", ""))
        if city_t:
            target_name = f"Територіала {city_t[0].get('full_name', '')}"

    text = (
        f"⚠️ <b>Ви дійсно хочете звільнити керівника {html.escape(name)}?</b>\n\n"
        f"• Доступ до функцій керівника буде заблоковано\n"
        f"• {interns_count} стажерів та {workers_count} працівників тихо перейдуть до {target_name}\n"
        f"• Всі дані та історія збережуться в базі для звітів та аналітики"
    )
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🚫 Так, звільнити", callback_data=f"dev_mgr_fire_confirm:{manager_uid}")],
        [InlineKeyboardButton(text="❌ Скасувати", callback_data=f"dev_mgr_view:{manager_uid}")]
    ])
    await _edit_or_answer(callback.message, text, reply_markup=kb)
    await callback.answer()


async def dev_mgr_fire_confirm(callback: CallbackQuery, state: FSMContext):
    """Виконання звільнення керівника."""
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access:
        return
    manager_uid = int(callback.data.split(":")[1])
    res = await fire_manager(manager_uid)
    await callback.answer(f"✅ Керівника звільнено. Підлеглих переведено до {res.get('target_name')}.", show_alert=True)
    await developer_manage_managers_menu(callback, state)


async def dev_mgr_reassign_menu(callback: CallbackQuery, state: FSMContext):
    """Меню переприв'язки відповідального для керівника."""
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access or (is_territorial and not is_admin):
        if has_access: await callback.answer("Ця дія доступна лише Адміністратору.", show_alert=True)
        return
    manager_uid = int(callback.data.split(":")[1])
    mgr = await get_manager_by_uid(manager_uid)
    if not mgr:
        await callback.answer("Керівника не знайдено!", show_alert=True)
        return

    city = mgr.get("city") or ""
    territorials = await get_territorials_by_city(city)
    
    buttons = [
        [InlineKeyboardButton(text="👑 Адміністратор (пряме підпорядкування)", callback_data=f"dev_mgr_reassign_set:{manager_uid}:0")]
    ]
    for t in territorials:
        t_type = t.get("territorial_type") or "ТЗ"
        t_name = t.get("full_name") or str(t["uid"])
        buttons.append([
            InlineKeyboardButton(text=f"🗺 {t_name} ({t_type})", callback_data=f"dev_mgr_reassign_set:{manager_uid}:{t['uid']}")
        ])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад до картки", callback_data=f"dev_mgr_view:{manager_uid}")])

    text = f"🔄 <b>Перепривʼязка керівника {html.escape(mgr.get('full_name') or '')}</b>\nОберіть нового відповідального:"
    await _edit_or_answer(callback.message, text, reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons))
    await callback.answer()


async def dev_mgr_reassign_set(callback: CallbackQuery, state: FSMContext):
    """Збереження нового відповідального для керівника."""
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access or (is_territorial and not is_admin):
        return
    parts = callback.data.split(":")
    manager_uid = int(parts[1])
    target_uid = int(parts[2])
    
    responsible_uid = None if target_uid == 0 else target_uid
    await update_manager_responsible(manager_uid, responsible_uid)
    await callback.answer("✅ Відповідального змінено!", show_alert=True)
    
    card_data = await _build_manager_card(manager_uid, is_admin, is_territorial)
    if card_data:
        text, kb = card_data
        await _edit_or_answer(callback.message, text, reply_markup=kb)


async def developer_cancel_add_role(callback: CallbackQuery, state: FSMContext):
    """Cancels the process of adding a new developer."""
    if not await _ensure_developer(callback):
        return
    await state.clear()
    # Повертаємо користувача до меню команди розробників
    await developer_dev_team_menu(callback, state)
    await callback.answer("Додавання скасовано.")


async def developer_manage_hrs_menu(callback: CallbackQuery, state: FSMContext):
    await developer_manage_managers_menu(callback, state)


async def developer_request_add_dev(callback: CallbackQuery, state: FSMContext):
    if not await _ensure_developer(callback):
        return
    await state.set_state(DeveloperStates.waiting_add_developer)
    prompt_info = await _show_prompt(
        callback,
        "👨‍💻 Введіть ID користувача, якому потрібно надати роль Адміністратора:",
        "dev_cancel_add_role",
    )
    await state.update_data(dev_prompt=prompt_info)
    try:
        await callback.answer()
    except TelegramBadRequest as e:
        if "query is too old" not in str(e):
            raise e


async def developer_process_add_dev(message: Message, state: FSMContext):
    data = await state.get_data()
    prompt_info = data.get("dev_prompt")
    text_value = (message.text or "").strip()
    if not text_value.isdigit():
        await _update_prompt_message(
            message.bot,
            prompt_info,
            "⚠️ ID повинен бути числом.\n\n👨‍💻 Введіть ID користувача, якому потрібно надати роль Адміністратора:",
            "dev_cancel_add_role",
        )
        return

    user_id = int(text_value)
    await state.update_data(new_dev_id=user_id)
    await state.set_state(DeveloperStates.waiting_add_developer_name)
    
    await _update_prompt_message(
        message.bot,
        prompt_info,
        f"ID: {user_id}. Тепер введіть ПІБ адміністратора:",
        "dev_cancel_add_role"
    )

async def developer_process_add_dev_name(message: Message, state: FSMContext):
    full_name = message.text.strip()
    if len(full_name.split()) < 2:
        await message.answer("Будь ласка, введіть повне ім'я та прізвище (мінімум 2 слова).")
        return

    data = await state.get_data()
    user_id = data.get("new_dev_id")
    prompt_info = data.get("dev_prompt")
    panel_info = data.get("dev_panel")

    # Отримуємо username якщо можливо
    try:
        chat_info = await message.bot.get_chat(user_id)
        username = chat_info.username
    except Exception:
        username = ""

    # Реєструємо або оновлюємо користувача в таблиці users
    await register_user(user_id, username=username, full_name=full_name)

    await add_manager(
        uid=user_id,
        process="Developer",
        full_name=full_name,
        username=username,
    )
    await state.set_state(None)
    await state.update_data(dev_prompt=None)
    await _update_prompt_message(
        message.bot,
        prompt_info,
        f"✅ Користувача {full_name} додано до Команди Адміністраторів.\n\nВикористайте кнопку нижче, щоб повернутися.",
        "dev_team_back",
    )
    await _refresh_panel_view(message.bot, panel_info, _build_developer_team_view)


async def developer_manage_hrs_menu(callback: CallbackQuery, state: FSMContext):
    if not await _ensure_developer(callback):
        return
    text, kb = await _build_hr_team_view()
    msg = await _edit_or_answer(callback.message, text, reply_markup=kb)
    await _remember_panel(state, "hr_panel", msg or callback.message)
    await callback.answer()


async def developer_request_add_hr(callback: CallbackQuery, state: FSMContext):
    if not await _ensure_developer(callback):
        return
    await state.set_state(DeveloperStates.waiting_add_hr)
    prompt_info = await _show_prompt(
        callback,
        "👨‍💼 Введіть ID користувача, якому потрібно надати роль HR:",
        "hr_cancel_add_role",
    )
    await state.update_data(hr_prompt=prompt_info)
    await callback.answer()


async def developer_process_add_hr(message: Message, state: FSMContext):
    data = await state.get_data()
    prompt_info = data.get("hr_prompt")
    panel_info = data.get("hr_panel")
    text_value = (message.text or "").strip()
    if not text_value.isdigit():
        await _update_prompt_message(
            message.bot,
            prompt_info,
            "⚠️ Введіть, будь ласка, числовий Telegram ID.\n\n"
            "👨‍💼 Введіть ID користувача, якому потрібно надати роль HR:",
            "hr_cancel_add_role",
        )
        return

    user_id = int(text_value)
    user = await get_user_details(user_id)
    if not user:
        await register_user(user_id)
        user = await get_user_details(user_id) or {"username": None, "full_name": None}

    # Use the name from Telegram API or the one already in the DB
    try:
        chat_info = await message.bot.get_chat(user_id)
        full_name = chat_info.full_name
        username = chat_info.username
    except Exception:
        full_name = user.get("full_name", "")
        username = user.get("username", "")

    await add_manager(
        user_id,
        process="Керівник",
        full_name=full_name,
        username=username
    )
    await state.set_state(None)
    await state.update_data(hr_prompt=None)
    await _update_prompt_message(
        message.bot,
        prompt_info,
        f"✅ Користувача {full_name or user_id} додано до HR-команди.\n\nВикористайте кнопку нижче, щоб повернутися.",
        "hr_team_back",
    )
    await _refresh_panel_view(message.bot, panel_info, _build_hr_team_view)


async def developer_cancel_add_role(callback: CallbackQuery, state: FSMContext):
    if not await _ensure_developer(callback, require_main=True):
        return
    await state.set_state(None)
    await state.update_data(dev_prompt=None)
    await developer_dev_team_menu(callback, state)
    await callback.answer("Додавання скасовано.", show_alert=True)


async def hr_cancel_add_role(callback: CallbackQuery, state: FSMContext):
    if not await _ensure_developer(callback):
        return
    await state.set_state(None)
    await state.update_data(hr_prompt=None)
    await developer_manage_hrs_menu(callback, state)
    await callback.answer("Додавання скасовано.", show_alert=True)


async def developer_team_member_handler(callback: CallbackQuery, state: FSMContext):
    if not await _ensure_developer(callback):
        return
    try:
        user_id = int(callback.data.split("_")[-1])
    except (ValueError, IndexError):
        await callback.answer("Некоректний ID.", show_alert=True)
        return
    name, username = await _format_identity(user_id)
    text = (
        f"👨‍💻 <b>{name}</b> (ID: <code>{user_id}</code>)\n"
        f"Username: {username}\n\n"
        "Цей користувач має роль Адміністратор. Що зробити?"
    )
    kb = InlineKeyboardMarkup(
        inline_keyboard=[
            [InlineKeyboardButton(text="🚫 Прибрати з Команди Адміністраторів", callback_data=f"dev_remove_{user_id}")],
            [InlineKeyboardButton(text="⬅️ Назад до списку", callback_data="dev_team_back")],
        ]
    )
    await _edit_or_answer(callback.message, text, reply_markup=kb)
    await callback.answer()


async def developer_remove_member(callback: CallbackQuery, state: FSMContext):
    if not await _ensure_developer(callback):
        return
    try:
        user_id = int(callback.data.split("_")[-1])
    except (ValueError, IndexError):
        await callback.answer("Некоректний ID.", show_alert=True)
        return
    await delete_manager_by_uid(user_id)
    text, kb = await _build_developer_team_view()
    msg = await _edit_or_answer(callback.message, f"✅ Користувача {user_id} вилучено з Команди Адміністраторів.\n\n{text}", reply_markup=kb)
    await _remember_panel(state, "dev_panel", msg or callback.message)
    await callback.answer("Користувача вилучено.", show_alert=True)


async def developer_team_back(callback: CallbackQuery, state: FSMContext):
    await developer_dev_team_menu(callback, state)


async def hr_team_member_handler(callback: CallbackQuery, state: FSMContext):
    if not await _ensure_developer(callback):
        return
    try:
        user_id = int(callback.data.split("_")[-1])
    except (ValueError, IndexError):
        await callback.answer("Некоректний ID.", show_alert=True)
        return
    name, username = await _format_identity(user_id)
    text = (
        f"👨‍💼 <b>{name}</b> (ID: <code>{user_id}</code>)\n"
        f"Username: {username}\n\n"
        "Цей користувач має доступ HR. Що зробити?"
    )
    kb = InlineKeyboardMarkup(
        inline_keyboard=[
            [InlineKeyboardButton(text="🚫 Прибрати з HR-команди", callback_data=f"hr_remove_{user_id}")],
            [InlineKeyboardButton(text="⬅️ Назад до списку", callback_data="hr_team_back")],
        ]
    )
    await _edit_or_answer(callback.message, text, reply_markup=kb)
    await callback.answer()


async def hr_remove_member(callback: CallbackQuery, state: FSMContext):
    if not await _ensure_developer(callback):
        return
    try:
        user_id = int(callback.data.split("_")[-1])
    except (ValueError, IndexError):
        await callback.answer("Некоректний ID.", show_alert=True)
        return
    await delete_manager_by_uid(user_id)
    text, kb = await _build_hr_team_view()
    msg = await _edit_or_answer(callback.message, f"✅ Користувача {user_id} вилучено з HR-команди.\n\n{text}", reply_markup=kb)
    await _remember_panel(state, "hr_panel", msg or callback.message)
    await callback.answer("Користувача вилучено.", show_alert=True)


async def hr_team_back(callback: CallbackQuery, state: FSMContext):
    await developer_manage_hrs_menu(callback, state)


# ==================== Materials Editor ====================

def _materials_menu_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_main_study")]
    ])


async def developer_materials_menu(callback: CallbackQuery):
    """Головне меню редагування матеріалів — вибір посади."""
    if not await _ensure_developer(callback):
        return
    
    # Використовуємо AVAILABLE_ROLES з constants
    buttons = []
    for i, role in enumerate(AVAILABLE_ROLES):
        # Скорочуємо назву для кнопки
        short_name = role[:25] + "..." if len(role) > 28 else role
        buttons.append([InlineKeyboardButton(
            text=short_name,
            callback_data=f"dev_mat_role|{i}"
        )])
    
    buttons.append([InlineKeyboardButton(text="❌ Скасувати", callback_data="dev_main_study")])
    
    await _edit_or_answer(
        callback.message,
        "📝 <b>Редагування навчального матеріалу</b>\n\n"
        "Оберіть посаду для редагування матеріалів:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    await callback.answer()


async def _build_dev_days_picker_keyboard(role: str, callback_prefix: str, back_callback: str, page: int = 0) -> InlineKeyboardMarkup:
    from database.positions import get_days_count_for_role
    total_days = await get_days_count_for_role(role)
    
    PER_PAGE = 6
    pages_total = max(1, (total_days + PER_PAGE - 1) // PER_PAGE)
    page = max(0, min(page, pages_total - 1))
    
    start_day = page * PER_PAGE + 1
    end_day = min(start_day + PER_PAGE, total_days + 1)
    
    buttons = []
    for day in range(start_day, end_day):
        buttons.append([InlineKeyboardButton(
            text=f"📅 День {day}",
            callback_data=f"{callback_prefix}|{day}"
        )])
        
    if pages_total > 1:
        nav_row = []
        if page > 0:
            nav_row.append(InlineKeyboardButton(text="⬅️", callback_data=f"{callback_prefix}_pg|{page - 1}"))
        nav_row.append(InlineKeyboardButton(text=f"📄 {page + 1}/{pages_total}", callback_data="ignore"))
        if page < pages_total - 1:
            nav_row.append(InlineKeyboardButton(text="➡️", callback_data=f"{callback_prefix}_pg|{page + 1}"))
        buttons.append(nav_row)
        
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data=back_callback)])
    return InlineKeyboardMarkup(inline_keyboard=buttons)


async def developer_materials_select_day(callback: CallbackQuery, state: FSMContext):
    """Вибір дня для редагування."""
    if not await _ensure_developer(callback):
        return
    
    try:
        _, role_part = callback.data.split("|", 1)
        role_index = int(role_part)
        if not (0 <= role_index < len(AVAILABLE_ROLES)):
            await callback.answer("Неправильний індекс ролі.", show_alert=True)
            return
        role = AVAILABLE_ROLES[role_index]
    except (ValueError, IndexError):
        await callback.answer("Помилка формату ролі.", show_alert=True)
        return
    
    await state.update_data(edit_role=role, edit_role_index=role_index)
    
    kb = await _build_dev_days_picker_keyboard(role, "dev_mat_day", "dev_materials_menu", page=0)
    
    await _edit_or_answer(
        callback.message,
        f"📝 <b>Редагування матеріалів</b>\n\n"
        f"Посада: <b>{role}</b>\n\n"
        f"Оберіть день:",
        reply_markup=kb
    )
    await callback.answer()


async def developer_materials_select_type(callback: CallbackQuery, state: FSMContext):
    """Вибір типу контенту (текст/відео)."""
    if not await _ensure_developer(callback):
        return
    
    try:
        _, day_str = callback.data.split("|", 1)
        day = int(day_str)
    except (ValueError, IndexError):
        await callback.answer("Помилка формату.", show_alert=True)
        return
    
    data = await state.get_data()
    role = data.get("edit_role", "")
    role_index = data.get("edit_role_index", 0)
    await state.update_data(edit_day=day)
    
    buttons = [
        [InlineKeyboardButton(text="📄 Текст", callback_data="dev_mat_type|text")],
        [InlineKeyboardButton(text="⬅️ Назад", callback_data=f"dev_mat_role|{role_index}")],
    ]
    
    await _edit_or_answer(
        callback.message,
        f"📝 <b>Редагування матеріалів</b>\n\n"
        f"Посада: <b>{role}</b>\n"
        f"День: <b>{day}</b>\n\n"
        f"Оберіть тип контенту:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    await callback.answer()


def parse_material_content(content: str) -> list[dict]:
    """
    Parses material content which can be:
    1. A JSON string representing a list of pages (each page is a dict with 'text' and optional 'photo').
    2. A plain text string (old format).
    
    Returns a list of dicts: [{'text': '...', 'photo': '...'}, ...]
    """
    if not content:
        return []
    
    content = content.strip()
    
    # Try parsing as JSON list of dicts (new format)
    if content.startswith('[') and content.endswith(']'):
        try:
            parsed = json.loads(content)
            if isinstance(parsed, list):
                # Validate items are dicts or can be converted
                normalized = []
                for item in parsed:
                    if isinstance(item, dict):
                        normalized.append(item)
                    elif isinstance(item, str):
                        normalized.append({"text": item})
                return normalized
        except json.JSONDecodeError:
            pass # Fallback to plain text handling
            
    # Fallback: Treat as plain text (old format)
    # Use split_text to paginate plain text
    from bot.utils.paginator import split_text
    text_pages = split_text(content)
    return [{"text": p} for p in text_pages]


async def developer_materials_view(callback: CallbackQuery, state: FSMContext):
    """Показ поточного матеріалу з можливістю редагування."""
    if not await _ensure_developer(callback):
        return
    
    try:
        _, content_type = callback.data.split("|", 1)
    except (ValueError, IndexError):
        await callback.answer("Помилка формату.", show_alert=True)
        return
    
    data = await state.get_data()
    role = data.get("edit_role", "")
    day = data.get("edit_day", 1)
    await state.update_data(edit_type=content_type)
    
    # Отримуємо матеріал з БД
    material = await get_material_by_role_day_type(role, day, content_type)
    
    if material:
        current_content = material.get("content", "")
        current_url = material.get("resource_url", "")
        material_id = material.get("id")
        await state.update_data(edit_material_id=material_id)
        
        if content_type == "text":
            parsed = parse_material_content(current_content)
            if parsed:
                preview_text = parsed[0].get("text", "")[:500]
                if len(parsed) > 1:
                    preview_text += f"\n\n... (ще сторінок: {len(parsed)-1})"
                elif parsed[0].get("photo"):
                    preview_text += "\n\n[Містить фото]"
            else:
                preview_text = "(порожньо)"

            text = (
                f"📝 <b>Текстовий матеріал</b>\n\n"
                f"Посада: <b>{role}</b>\n"
                f"День: <b>{day}</b>\n\n"
                f"<b>Поточний контент:</b>\n"
                f"<code>{preview_text}</code>\n\n"
                f"Натисніть «Редагувати», щоб змінити текст."
            )
        else:
            preview = current_url if current_url else "(не встановлено)"
            text = (
                f"🎥 <b>Відео матеріал</b>\n\n"
                f"Посада: <b>{role}</b>\n"
                f"День: <b>{day}</b>\n\n"
                f"<b>Поточне посилання:</b>\n"
                f"<code>{preview}</code>\n\n"
                f"Натисніть «Редагувати», щоб змінити посилання."
            )
    else:
        await state.update_data(edit_material_id=None)
        text = (
            f"📝 <b>Матеріал не знайдено</b>\n\n"
            f"Посада: <b>{role}</b>\n"
            f"День: <b>{day}</b>\n"
            f"Тип: <b>{content_type}</b>\n\n"
            f"Натисніть «Редагувати», щоб створити новий матеріал."
        )
    
    buttons = []
        
    buttons.append([InlineKeyboardButton(text="➕ Додати навчання", callback_data=f"dev_mat_edit|{content_type}")])
    if material:
        buttons.append([InlineKeyboardButton(text="➕ Доповнити", callback_data=f"dev_mat_complement|{content_type}")])
        buttons.append([InlineKeyboardButton(text="🔄 Змінити діапазон сторінок", callback_data=f"dev_mat_replace|{content_type}")])
    buttons.append([InlineKeyboardButton(text="🔍 Повний перегляд", callback_data=f"dev_mat_full_view|{content_type}")])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data=f"dev_mat_day|{day}")])
    
    await _edit_or_answer(
        callback.message,
        text,
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    await callback.answer()

async def developer_material_complement_menu(callback: CallbackQuery, state: FSMContext):
    """Shows the menu to choose how to complement the material."""
    if not await _ensure_developer(callback): return
    
    try:
        _, content_type = callback.data.split("|", 1)
    except:
        return await callback.answer("Помилка формату.")

    if content_type != "text":
        return await callback.answer("Доповнення поки доступне лише для текстових матеріалів.", show_alert=True)

    data = await state.get_data()
    material_id = data.get("edit_material_id")
    if not material_id:
        return await callback.answer("Матеріал не знайдено.")

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🔝 На початок", callback_data="dev_comp_type:start")],
        [InlineKeyboardButton(text="↔️ Вставити між сторінками", callback_data="dev_comp_type:between")],
        [InlineKeyboardButton(text="🔚 В кінець", callback_data="dev_comp_type:end")],
        [InlineKeyboardButton(text="❌ Скасувати", callback_data=f"dev_mat_type|{content_type}")]
    ])

    await _edit_or_answer(
        callback.message,
        "➕ <b>Доповнення матеріалу</b>\n\nОберіть, куди саме ви хочете додати нову сторінку:",
        reply_markup=kb
    )
    await callback.answer()

async def developer_material_complement_start(callback: CallbackQuery, state: FSMContext):
    """Processes the chosen complement type."""
    comp_type = callback.data.split(":")[1]
    await state.update_data(complement_type=comp_type)
    
    if comp_type == "between":
        await _edit_or_answer(
            callback.message,
            "🔢 <b>Вставка між сторінками</b>\n\nВведіть номер <b>першої</b> сторінки (наприклад, 19):",
            reply_markup=_cancel_keyboard("dev_mat_menu_back") # Simplified back
        )
        await state.set_state(DeveloperStates.waiting_insert_start_idx)
    else:
        # For start/end, go directly to mode choice
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="📄 Одна сторінка", callback_data="dev_ins_mode:single")],
            [InlineKeyboardButton(text="📚 Масив сторінок", callback_data="dev_ins_mode:array")],
            [InlineKeyboardButton(text="❌ Скасувати", callback_data="dev_mat_menu_back")]
        ])
        await _edit_or_answer(callback.message, "Оберіть режим додавання:", reply_markup=kb)
        await state.set_state(DeveloperStates.waiting_insert_mode)
    
    await callback.answer()

async def developer_process_insert_start_idx(message: Message, state: FSMContext):
    if not message.text or not message.text.isdigit():
        return await message.answer("❌ Будь ласка, введіть число (номер сторінки).")
    
    await state.update_data(insert_start_val=int(message.text))
    await message.answer("🔢 Тепер введіть номер <b>другої</b> сторінки (наприклад, 20):")
    await state.set_state(DeveloperStates.waiting_insert_end_idx)

async def developer_process_insert_end_idx(message: Message, state: FSMContext):
    if not message.text or not message.text.isdigit():
        return await message.answer("❌ Будь ласка, введіть число.")
    
    data = await state.get_data()
    start_idx = data.get("insert_start_val")
    end_idx = int(message.text)
    
    if end_idx != start_idx + 1:
        await message.answer(f"⚠️ Ви вказали {start_idx} та {end_idx}. Зазвичай вставляють між сусідніми сторінками (наприклад, {start_idx} та {start_idx+1}).\n\nАле я продовжу. Нова сторінка стане номером {end_idx}, а стара {end_idx} та наступні посунуться.")
    
    await state.update_data(target_insert_pos=end_idx - 1) # 0-based index
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📄 Одна сторінка", callback_data="dev_ins_mode:single")],
        [InlineKeyboardButton(text="📚 Масив сторінок", callback_data="dev_ins_mode:array")],
        [InlineKeyboardButton(text="❌ Скасувати", callback_data="dev_mat_menu_back")]
    ])
    await message.answer("Оберіть режим додавання:", reply_markup=kb)
    await state.set_state(DeveloperStates.waiting_insert_mode)

async def developer_process_complement_content(message: Message, state: FSMContext):
    data = await state.get_data()
    comp_type = data.get("complement_type")
    material_id = data.get("edit_material_id")
    day = data.get("edit_day")
    role = data.get("edit_role")
    
    material = await get_material_by_id(material_id)
    if not material:
        await message.answer("❌ Помилка: матеріал не знайдено.")
        return await state.clear()

    pages = parse_material_content(material.get("content", "[]"))
    
    # Prepare new page
    new_page = {"text": message.caption or message.text or ""}
    if message.photo:
        new_page["photo"] = message.photo[-1].file_id

    # Logic for insertion
    if comp_type == "start":
        pages.insert(0, new_page)
        target_page = 0
    elif comp_type == "end":
        pages.append(new_page)
        target_page = len(pages) - 1
    elif comp_type == "between":
        pos = data.get("target_insert_pos", 0)
        if pos < 0: pos = 0
        if pos > len(pages): pos = len(pages)
        pages.insert(pos, new_page)
        target_page = pos
    
    # Save to DB
    new_content_json = json.dumps(pages, ensure_ascii=False)
    await update_material_content(material_id, new_content_json)
    
    await message.answer(f"✅ Матеріал успішно доповнено! Нова сторінка додана на позицію {target_page + 1}.")
    
    # Move to notification decision
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📢 Оголосити зараз", callback_data="dev_notify:now")],
        [InlineKeyboardButton(text="📥 Обʼєднати з наступним", callback_data="dev_notify:merge")],
        [InlineKeyboardButton(text="🔇 Не оголошувати", callback_data="dev_notify:none")]
    ])
    await message.answer("Сповістити користувачів про зміни в матеріалі?", reply_markup=kb)
    await state.set_state(DeveloperStates.waiting_announce_decision)

async def developer_process_insert_mode(callback: CallbackQuery, state: FSMContext):
    mode = callback.data.split(":")[1]
    data = await state.get_data()
    comp_type = data.get("complement_type")
    
    if mode == "single":
        if comp_type == "between":
            start_idx = data.get("insert_start_val")
            end_idx = data.get("target_insert_pos") + 1
            prompt = f"📝 Надішліть контент (текст/фото), який потрібно вставити між {start_idx} та {end_idx} сторінками:"
        else:
            prompt = "🔝 Надішліть контент, який стане <b>першою</b> сторінкою:" if comp_type == "start" else "🔚 Надішліть контент, який буде додано в <b>кінець</b>:"
        await _edit_or_answer(callback.message, prompt, reply_markup=_cancel_keyboard("dev_mat_menu_back"))
        await state.set_state(DeveloperStates.waiting_complement_content)
    else:
        await state.update_data(insert_array=[])
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="✅ ГОТОВО (Зберегти)", callback_data="dev_ins_array_done")],
            [InlineKeyboardButton(text="❌ Скасувати", callback_data="dev_mat_menu_back")]
        ])
        await _edit_or_answer(callback.message, "📚 <b>Режим: Масив сторінок</b>\n\nНадсилайте сторінки по черзі (текст або фото). Коли закінчите, натисніть <b>ГОТОВО</b>.", reply_markup=kb)
        await state.set_state(DeveloperStates.waiting_insert_array_content)
    await callback.answer()

async def developer_process_insert_array_content(message: Message, state: FSMContext):
    data = await state.get_data()
    insert_array = data.get("insert_array", [])
    
    new_page = {"text": message.caption or message.text or ""}
    if message.photo:
        new_page["photo"] = message.photo[-1].file_id
        
    insert_array.append(new_page)
    await state.update_data(insert_array=insert_array)
    
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="✅ ГОТОВО (Зберегти)", callback_data="dev_ins_array_done")],
        [InlineKeyboardButton(text="❌ Скасувати", callback_data="dev_mat_menu_back")]
    ])
    await message.answer(f"✅ Додано {len(insert_array)}-ю сторінку до масиву.\nНадсилайте наступну або натисніть ГОТОВО.", reply_markup=kb)

async def developer_process_insert_array_done(callback: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    insert_array = data.get("insert_array", [])
    if not insert_array:
        return await callback.answer("Ви не додали жодної сторінки!", show_alert=True)
        
    comp_type = data.get("complement_type")
    material_id = data.get("edit_material_id")
    
    material = await get_material_by_id(material_id)
    if not material:
        await callback.message.answer("❌ Помилка: матеріал не знайдено.")
        return await state.clear()

    pages = parse_material_content(material.get("content", "[]"))
    
    # Logic for insertion
    if comp_type == "start":
        pages = insert_array + pages
    elif comp_type == "end":
        pages.extend(insert_array)
    elif comp_type == "between":
        pos = data.get("target_insert_pos", 0)
        if pos < 0: pos = 0
        if pos > len(pages): pos = len(pages)
        pages = pages[:pos] + insert_array + pages[pos:]
    
    # Save to DB
    new_content_json = json.dumps(pages, ensure_ascii=False)
    await update_material_content(material_id, new_content_json)
    
    await callback.message.answer(f"✅ Успішно додано {len(insert_array)} сторінок!")
    
    # Move to notification decision
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📢 Оголосити зараз", callback_data="dev_notify:now")],
        [InlineKeyboardButton(text="📥 Обʼєднати з наступним", callback_data="dev_notify:merge")],
        [InlineKeyboardButton(text="🔇 Не оголошувати", callback_data="dev_notify:none")]
    ])
    await _edit_or_answer(callback.message, "Сповістити користувачів про зміни в матеріалі?", reply_markup=kb)
    await state.set_state(DeveloperStates.waiting_announce_decision)
    await callback.answer()

async def developer_material_replace_range_start(callback: CallbackQuery, state: FSMContext):
    """Start replacing a range of pages."""
    if not await _ensure_developer(callback): return
    
    await _edit_or_answer(
        callback.message, 
        "🔄 <b>Заміна діапазону сторінок</b>\n\nВведіть діапазон сторінок для заміни (наприклад, '3-5' або '3 5'):", 
        reply_markup=_cancel_keyboard("dev_mat_menu_back")
    )
    await state.set_state(DeveloperStates.waiting_replace_range)
    await callback.answer()

async def developer_process_replace_range(message: Message, state: FSMContext):
    text = message.text.replace("-", " ").strip()
    parts = text.split()
    if len(parts) != 2 or not all(p.isdigit() for p in parts):
        return await message.answer("❌ Невірний формат. Введіть два числа, наприклад '3-5' або '3 5'.")
        
    start_idx, end_idx = int(parts[0]), int(parts[1])
    if start_idx > end_idx or start_idx < 1:
        return await message.answer("❌ Некоректний діапазон. Перше число має бути менше або дорівнювати другому і більше нуля.")
        
    await state.update_data(replace_start=start_idx - 1, replace_end=end_idx, replace_array=[]) # 0-based for start, end is exclusive for slicing
    
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="✅ ГОТОВО (Зберегти)", callback_data="dev_rep_array_done")],
        [InlineKeyboardButton(text="❌ Скасувати", callback_data="dev_mat_menu_back")]
    ])
    await message.answer(f"📚 <b>Режим заміни: Масив сторінок</b>\n\nНадсилайте нові сторінки по черзі, щоб замінити сторінки з {start_idx} по {end_idx}. Коли закінчите, натисніть <b>ГОТОВО</b>.", reply_markup=kb)
    await state.set_state(DeveloperStates.waiting_replace_array_content)

async def developer_process_replace_array_content(message: Message, state: FSMContext):
    data = await state.get_data()
    replace_array = data.get("replace_array", [])
    
    new_page = {"text": message.caption or message.text or ""}
    if message.photo:
        new_page["photo"] = message.photo[-1].file_id
        
    replace_array.append(new_page)
    await state.update_data(replace_array=replace_array)
    
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="✅ ГОТОВО (Зберегти)", callback_data="dev_rep_array_done")],
        [InlineKeyboardButton(text="❌ Скасувати", callback_data="dev_mat_menu_back")]
    ])
    await message.answer(f"✅ Додано {len(replace_array)}-ю сторінку для заміни.\nНадсилайте наступну або натисніть ГОТОВО.", reply_markup=kb)

async def developer_process_replace_array_done(callback: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    replace_array = data.get("replace_array", [])
    
    material_id = data.get("edit_material_id")
    material = await get_material_by_id(material_id)
    if not material:
        await callback.message.answer("❌ Помилка: матеріал не знайдено.")
        return await state.clear()

    pages = parse_material_content(material.get("content", "[]"))
    start_idx = data.get("replace_start", 0)
    end_idx = data.get("replace_end", 0)
    
    if start_idx > len(pages): start_idx = len(pages)
    if end_idx > len(pages): end_idx = len(pages)
    
    # Replace slice
    pages[start_idx:end_idx] = replace_array
    
    # Save to DB
    new_content_json = json.dumps(pages, ensure_ascii=False)
    await update_material_content(material_id, new_content_json)
    
    await callback.message.answer(f"✅ Успішно замінено діапазон на {len(replace_array)} нових сторінок!")
    
    # Move to notification decision
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📢 Оголосити зараз", callback_data="dev_notify:now")],
        [InlineKeyboardButton(text="📥 Обʼєднати з наступним", callback_data="dev_notify:merge")],
        [InlineKeyboardButton(text="🔇 Не оголошувати", callback_data="dev_notify:none")]
    ])
    await _edit_or_answer(callback.message, "Сповістити користувачів про зміни в матеріалі?", reply_markup=kb)
    await state.set_state(DeveloperStates.waiting_announce_decision)
    await callback.answer()

async def developer_process_notify_decision(callback: CallbackQuery, state: FSMContext):
    decision = callback.data.split(":")[1]
    
    if decision == "now":
        from database.material_notifications import create_material_change_event, create_recipients_for_event
        from bot.services.wave_broadcaster import start_event_wave_broadcast
        
        data = await state.get_data()
        role = data.get("edit_role") or data.get("edit_video_role") or data.get("edit_photo_role") or data.get("test_role") or "ALL"
        day = data.get("edit_day") or data.get("edit_video_day") or data.get("edit_photo_day") or data.get("test_day") or 1
        content_type = data.get("content_type") or "text"
        pages = data.get("insert_array") or data.get("replace_array") or data.get("pages")
        if not pages:
            material_id = data.get("material_id")
            if material_id:
                mat = await get_material_by_id(material_id)
                if mat and mat.get("content"):
                    try:
                        pages = json.loads(mat["content"])
                    except Exception:
                        pages = [mat["content"]]
        if not pages:
            pages = [f"Оновлено навчальні матеріали ({content_type})"]

        description = f"Оновлення матеріалів {content_type} (День {day})"
        event_id = await create_material_change_event(role, day, content_type, description, pages)
        recipients_count = await create_recipients_for_event(event_id, role)
        start_event_wave_broadcast(callback.bot, event_id)

        waves_total = max(1, (recipients_count + 39) // 40)
        await callback.message.answer(
            f"📢 <b>Хвильову розсилку запущено!</b>\n\n"
            f"📚 Посада: <b>{role}</b> | 📅 День: <b>{day}</b>\n"
            f"👥 Всього отримувачів: <b>{recipients_count}</b>\n"
            f"🌊 Кількість хвиль: <b>{waves_total}</b> (по 40 осіб на годину)\n"
            f"⏳ Першу хвилю надіслано негайно."
        )
    elif decision == "merge":
        await callback.message.answer("📥 Зміни збережено для наступного оголошення.")
    elif decision == "none":
        await callback.message.answer("🔇 Зміни збережено без оголошення.")
        
    await state.clear()
    await developer_material_menu_back(callback, state)

async def developer_material_delete_request(callback: CallbackQuery):
    """Asks for confirmation before deleting a page."""
    if not await _ensure_developer(callback): return
    
    try:
        parts = callback.data.split(":")
        mid_str, day_str, page_str, ctype = parts[1], parts[2], parts[3], parts[4]
        page = int(page_str)
    except:
        return await callback.answer("Помилка даних.")

    suffix = f"{mid_str}:{day_str}:{page_str}:{ctype}"
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="✅ Так, видалити", callback_data=f"dev_mat_del_conf:{suffix}")],
        [InlineKeyboardButton(text="❌ Ні, назад", callback_data=f"dev_pag:{suffix}")]
    ])

    await _edit_or_answer(
        callback.message,
        f"❓ <b>Ви дійсно бажаєте видалити сторінку №{page + 1}?</b>\n\nЦю дію неможливо буде скасувати.",
        reply_markup=kb
    )
    await callback.answer()

async def developer_material_delete_confirm(callback: CallbackQuery):
    """Performs the actual deletion of a page."""
    if not await _ensure_developer(callback): return
    
    try:
        # callback.data format: dev_mat_del_conf:mid:day:page:ctype
        _, mid_str, day_str, page_str, ctype = callback.data.split(":")
        material_id = int(mid_str)
        day = int(day_str)
        page = int(page_str)
    except:
        return await callback.answer("Помилка видалення.")

    material = await get_material_by_id(material_id)
    if not material:
        return await callback.answer("Матеріал не знайдено.")

    pages = parse_material_content(material.get("content", "[]"))
    
    if 0 <= page < len(pages):
        pages.pop(page)
        
        # Save updated content
        new_content_json = json.dumps(pages, ensure_ascii=False)
        await update_material_content(material_id, new_content_json)
        
        await callback.answer("✅ Сторінку видалено!", show_alert=True)
        
        if not pages:
            # If no pages left, go back to material view
            await developer_materials_view(callback, None)
        else:
            # Show the new page at the same index (or the last one if we deleted the end)
            new_page_idx = page if page < len(pages) else len(pages) - 1
            # Mock callback to trigger pagination view
            callback.data = f"dev_pag:{material_id}:{day}:{new_page_idx}:{ctype}"
            await developer_pagination_handler(callback)
    else:
        await callback.answer("Помилка: сторінка вже не існує.")

async def developer_material_fix_start(callback: CallbackQuery, state: FSMContext):
    """Starts the process of fixing a specific page of material."""
    if not await _ensure_developer(callback): return
    
    try:
        _, mid_str, day_str, page_str = callback.data.split(":")
        material_id = int(mid_str)
        day = int(day_str)
        page = int(page_str)
    except (ValueError, IndexError):
        await callback.answer("Помилка даних.")
        return

    material = await get_material_by_id(material_id)
    if not material:
        await callback.answer("Матеріал не знайдено.")
        return
        
    content_type = material.get("content_type", "text")
    if content_type != "text" and content_type != "photo_files":
        await callback.answer("Виправлення поки доступне тільки для тексту та фото.", show_alert=True)
        return

    if content_type == "photo_files":
        try:
            file_ids = json.loads(material.get("content", "[]"))
            pages = [{"photo": fid} for fid in file_ids]
        except:
            pages = []
        current_text = ""
        prompt_text = "Надішліть нове фото для заміни поточного."
    else:
        pages = parse_material_content(material.get("content", ""))
        current_page_data = pages[page]
        current_text = current_page_data.get("text", "")
        prompt_text = "Надішліть виправлений текст для цієї сторінки.\nЯкщо сторінка містить фото, ви можете надіслати нове фото з підписом для заміни, або просто текст, щоб залишити старе фото."
    
    if not (0 <= page < len(pages)):
        await callback.answer("Сторінка не знайдена.")
        return
        
    await state.update_data(
        fix_material_id=material_id,
        fix_day=day,
        fix_page=page,
        fix_full_content=pages,
        fix_content_type=content_type
    )
    await state.set_state(DeveloperStates.waiting_fix_page)
    
    kb = InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="❌ Скасувати", callback_data=f"dev_pag:{material_id}:{day}:{page}:{content_type}")
    ]])
    
    # 1. Прибираємо кнопки з поточного повідомлення (перегляду), щоб не клікали під час редагування
    try:
        await callback.message.edit_reply_markup(reply_markup=None)
    except:
        pass

    # 2. Надсилаємо інструкцію окремим повідомленням
    instruction_text = (
        f"🛠 <b>Виправлення сторінки {page + 1}</b>\n\n"
        f"Поточний текст:\n<pre>{current_text}</pre>\n\n"
        f"{prompt_text}"
    )
    
    await callback.message.answer(instruction_text, reply_markup=kb)
    await callback.answer()

async def developer_process_fix_page(message: Message, state: FSMContext):
    """Processes the fixed content for a page."""
    data = await state.get_data()
    material_id = data.get("fix_material_id")
    page = data.get("fix_page")
    pages = data.get("fix_full_content")
    day = data.get("fix_day")
    ctype = data.get("fix_content_type")
    
    if not pages or page is None:
        await message.answer("Помилка стану.")
        await state.clear()
        return
    
    if ctype == "photo_files":
        if not message.photo:
            await message.answer("Будь ласка, надішліть фото для заміни.")
            return
        
        if 0 <= page < len(pages):
            pages[page]["photo"] = message.photo[-1].file_id
            
        # For photo_files, content is just list of IDs
        import json
        file_ids = [p["photo"] for p in pages if p.get("photo")]
        new_content_json = json.dumps(file_ids)
        
        await update_material_content(material_id, new_content_json)
        await message.answer(f"✅ Фото {page + 1} оновлено!")
        
    else:
        # Text content logic
        input_text = message.caption if message.caption is not None else message.text
        
        if 0 <= page < len(pages):
            if input_text is not None:
                 pages[page]["text"] = input_text
            if message.photo:
                 pages[page]["photo"] = message.photo[-1].file_id
        
        import json
        new_content_json = json.dumps(pages, ensure_ascii=False)
        await update_material_content(material_id, new_content_json)
        await message.answer(f"✅ Сторінку {page + 1} оновлено!")
    
    # Render view
    kb = _get_dev_pagination_keyboard(page, len(pages), material_id, day, ctype)
    
    page_data = pages[page]
    text_display = page_data.get("text", "")
    photo_display = page_data.get("photo")
    
    if photo_display:
        if len(text_display or "") <= 1024:
            await message.answer_photo(photo_display, caption=text_display, reply_markup=kb)
        else:
            # Long text: send photo then text
            await message.answer_photo(photo_display)
            await message.answer(text_display, reply_markup=kb)
    else:
        await message.answer(text_display, reply_markup=kb)
        
    await state.clear()


async def developer_material_menu_back(callback: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    ctype = data.get("edit_type", "text")
    await state.set_state(None)
    await developer_materials_view(callback, state) # Go back to material view

def _get_dev_pagination_keyboard(current_page, total_pages, material_id, day, content_type):
    buttons = []
    
    # Navigation row (arrows)
    nav_row = []
    if current_page > 0:
        nav_row.append(InlineKeyboardButton(text="⬅️", callback_data=f"dev_pag:{material_id}:{day}:{current_page-1}:{content_type}"))
    
    if total_pages > 1:
        nav_row.append(InlineKeyboardButton(text=f"{current_page+1}/{total_pages}", callback_data="noop"))
    
    if current_page < total_pages - 1:
        nav_row.append(InlineKeyboardButton(text="➡️", callback_data=f"dev_pag:{material_id}:{day}:{current_page+1}:{content_type}"))
    
    if nav_row:
        buttons.append(nav_row)
    
    # "Fix" button on every page
    buttons.append([InlineKeyboardButton(text="🛠 Виправити", callback_data=f"dev_mat_fix:{material_id}:{day}:{current_page}")])
    # "Delete" button on every page
    buttons.append([InlineKeyboardButton(text="🗑 Видалити сторінку", callback_data=f"dev_mat_del_req:{material_id}:{day}:{current_page}:{content_type}")])

    # Action buttons: on the last page OR if there is only 1 page
    if current_page == total_pages - 1:
        buttons.append([InlineKeyboardButton(text="➕ Додати навчання", callback_data=f"dev_mat_edit|{content_type}")])
        buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data=f"dev_mat_type|{content_type}")])
    
    return InlineKeyboardMarkup(inline_keyboard=buttons)

async def developer_pagination_handler(callback: CallbackQuery):
    try:
        _, mid_str, day_str, page_str, ctype = callback.data.split(":")
        material_id = int(mid_str)
        day = int(day_str)
        page = int(page_str)
    except (ValueError, IndexError):
        await callback.answer("Помилка навігації.")
        return

    material = await get_material_by_id(material_id)
    if not material:
        await callback.answer("Матеріал не знайдено.")
        return

    if ctype == "text":
        pages = parse_material_content(material.get("content", ""))
    elif ctype == "video":
        pages = [{"text": material.get("resource_url", "")}]
    elif ctype == "photo_files":
        try:
            file_ids = json.loads(material.get("content", "[]"))
            pages = [{"photo": fid, "text": f"Фото {i+1}"} for i, fid in enumerate(file_ids)]
        except json.JSONDecodeError:
            pages = []
    else:
        # Fallback for other types
        from bot.utils.paginator import split_text
        pages = [{"text": p} for p in split_text(material.get("content", ""))]

    if 0 <= page < len(pages):
        kb = _get_dev_pagination_keyboard(page, len(pages), material_id, day, ctype)
        
        page_data = pages[page]
        text = page_data.get("text", "")
        photo = page_data.get("photo")
        
        is_photo_message = bool(callback.message.photo)
        is_text_message = bool(callback.message.text)
        
        if photo:
            if is_photo_message:
                 media = InputMediaPhoto(media=photo, caption=text)
                 await callback.message.edit_media(media=media, reply_markup=kb)
            else:
                 await callback.message.delete()
                 await callback.message.answer_photo(photo, caption=text, reply_markup=kb)
        else:
            if is_text_message:
                 if callback.message.text != text:
                     await callback.message.edit_text(text, reply_markup=kb)
                 else:
                     await callback.message.edit_reply_markup(reply_markup=kb)
            else:
                 await callback.message.delete()
                 await callback.message.answer(text, reply_markup=kb)
    
    await callback.answer()

async def developer_materials_full_view(callback: CallbackQuery, state: FSMContext):
    """Надсилає повний текст матеріалу розробнику з пагінацією."""
    if not await _ensure_developer(callback):
        return

    try:
        _, content_type = callback.data.split("|", 1)
    except (ValueError, IndexError):
        await callback.answer("Помилка формату.", show_alert=True)
        return

    data = await state.get_data()
    role = data.get("edit_role", "")
    day = data.get("edit_day", 1)

    material = await get_material_by_role_day_type(role, day, content_type)

    if material:
        if content_type == "text":
            pages = parse_material_content(material.get("content", ""))
        elif content_type == "video": 
            pages = [{"text": material.get("resource_url", "(Не встановлено)")}]
        elif content_type == "photo_files":
            try:
                file_ids = json.loads(material.get("content", "[]"))
                pages = [{"photo": fid, "text": f"Фото {i+1}"} for i, fid in enumerate(file_ids)]
            except json.JSONDecodeError:
                pages = []
        else:
            from bot.utils.paginator import split_text
            pages = [{"text": p} for p in split_text(material.get("content", ""))]
        material_id = material.get("id")
    else:
        pages = [{"text": "(Матеріал ще не створено)"}]
        material_id = 0
    
    if not pages:
        pages = [{"text": "(Матеріал порожній)"}]
    
    # Show page 0
    kb = _get_dev_pagination_keyboard(0, len(pages), material_id, day, content_type)
    
    # Always delete the old message (which might be a menu) and send a new message
    try:
        await callback.message.delete()
    except Exception:
        pass 

    page_0 = pages[0]
    text = page_0.get("text", "")
    photo = page_0.get("photo")
    
    if photo:
        await callback.message.answer_photo(photo, caption=text, reply_markup=kb)
    else:
        await callback.message.answer(text, reply_markup=kb)

    await callback.answer()




async def developer_materials_edit_start(callback: CallbackQuery, state: FSMContext):
    """Початок редагування — запит нового контенту."""
    if not await _ensure_developer(callback):
        return
    
    try:
        _, content_type = callback.data.split("|", 1)
    except (ValueError, IndexError):
        await callback.answer("Помилка формату.", show_alert=True)
        return
    
    data = await state.get_data()
    role = data.get("edit_role", "")
    day = data.get("edit_day", 1)
    
    if content_type == "text":
        prompt = (
            f"✏️ <b>Редагування тексту</b>\n\n"
            f"Посада: <b>{role}</b>\n"
            f"День: <b>{day}</b>\n\n"
            f"Надішліть новий текстовий контент для цього матеріалу. Можна відправити кількома повідомленнями. Коли закінчите, напишіть <b>Готово</b>."
        )
        await state.set_state(DeveloperStates.waiting_material_parts)
        await state.update_data(material_parts=[])
    else:
        prompt = (
            f"🎥 <b>Редагування відео</b>\n\n"
            f"Посада: <b>{role}</b>\n"
            f"День: <b>{day}</b>\n\n"
            f"Надішліть нове посилання на відео (YouTube, Vimeo тощо):"
        )
        await state.set_state(DeveloperStates.waiting_material_video)
    
    buttons = [
        [InlineKeyboardButton(text="❌ Скасувати", callback_data=f"dev_mat_type|{content_type}")],
    ]
    
    await _edit_or_answer(
        callback.message,
        prompt,
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    await callback.answer()


async def developer_process_material_parts(message: Message, state: FSMContext):
    """Collects multiple messages for material content until 'Готово' is received."""
    raw_text = (message.text or message.caption or "").strip()
    # Cleaner check for 'готово' (handles dots, spaces, case)
    is_done = raw_text.lower().rstrip('.! ') == 'готово'
    
    data = await state.get_data()
    material_parts = data.get("material_parts", [])
    confirmation_msg_id = data.get("text_confirmation_msg_id")

    if is_done:
        if not material_parts:
            await message.answer("❌ Ви нічого не надіслали. Ввеведіть контент або скасуйте редагування.")
            return

        if confirmation_msg_id:
            try:
                await message.bot.delete_message(chat_id=message.chat.id, message_id=confirmation_msg_id)
            except Exception:
                pass
            
        # Serialize to JSON for new format
        import json
        full_content = json.dumps(material_parts, ensure_ascii=False)
        await _finalize_material_update(message, state, full_content)
    else:
        # Store as object with text and optional photo
        page = {"text": raw_text}
        if message.photo:
            page["photo"] = message.photo[-1].file_id
            
        material_parts.append(page)
        await state.update_data(material_parts=material_parts)

        new_text = f"✅ Отримано частин: {len(material_parts)}. Надішліть наступну частину (текст/фото) або напишіть 'Готово', щоб завершити."

        if confirmation_msg_id:
            try:
                await message.bot.edit_message_text(
                    chat_id=message.chat.id,
                    message_id=confirmation_msg_id,
                    text=new_text,
                )
            except TelegramBadRequest as e:
                if "message is not modified" not in str(e):
                    sent_msg = await message.answer(new_text)
                    await state.update_data(text_confirmation_msg_id=sent_msg.message_id)
        else:
            sent_msg = await message.answer(new_text)
            await state.update_data(text_confirmation_msg_id=sent_msg.message_id)


async def _finalize_material_update(message: Message, state: FSMContext, new_content: str):
    """Обробка нового текстового контенту і збереження."""
    if not new_content:
        await message.answer("❌ Текст не може бути порожнім.")
        return
    
    data = await state.get_data()
    material_id = data.get("edit_material_id") # Get the original material ID

    # Зберігаємо новий контент у стейті для подальшого збереження
    await state.update_data(
        pending_content=new_content,
        pending_content_type="text",
        pending_material_id=material_id # Pass it to the save function
    )
    
    # Одразу зберігаємо матеріал
    success, role, content_type, day = await _save_pending_material(state)
    
    if success:
        await message.answer(
            f"✅ <b>Матеріал оновлено!</b>\n\n"
            f"Посада: {role}\n"
            f"День: {day}\n"
            f"Тип: {content_type}",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="⬅️ До Панелі Адміністратора", callback_data="developer_menu")]
            ])
        )
    else:
        await message.answer(
            "❌ Помилка збереження матеріалу.",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="⬅️ До Панелі Адміністратора", callback_data="developer_menu")]
            ])
        )
    
    await state.clear()


async def developer_process_material_video(message: Message, state: FSMContext):
    """Обробка нового посилання на відео і збереження."""
    new_url = message.text.strip()
    if not new_url:
        await message.answer("❌ Посилання не може бути порожнім.")
        return

    data = await state.get_data()
    material_id = data.get("edit_material_id")
    await state.update_data(
        pending_content=new_url,
        pending_content_type="video",
        pending_material_id=material_id
    )
    
    success, role, content_type, day = await _save_pending_material(state)
    
    if success:
        await message.answer(
            f"✅ <b>Матеріал оновлено!</b>\n\nПосада: {role}\nДень: {day}\nТип: {content_type}",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="⬅️ До Панелі Адміністратора", callback_data="developer_menu")]])
        )
    else:
        await message.answer("❌ Помилка збереження матеріалу.", reply_markup=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="⬅️ До Панелі Адміністратора", callback_data="developer_menu")]]))
    
    await state.clear()


async def _finalize_video_update(message: Message, state: FSMContext, video_file_ids: list[str]):
    """Обробка завантажених відеофайлів і збереження."""
    if not video_file_ids:
        await message.answer("❌ Немає відеофайлів для збереження.")
        return

    data = await state.get_data()
    material_id = data.get("edit_video_material_id")
    await state.update_data(
        pending_content=json.dumps(video_file_ids),
        pending_content_type="video_files",
        pending_material_id=material_id
    )

    success, role, content_type, day = await _save_pending_material(state)
    
    if success:
        await message.answer(
            f"✅ <b>Матеріал оновлено!</b>\n\nПосада: {role}\nДень: {day}\nТип: {content_type}\nЗавантажено файлів: {len(video_file_ids)}",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="⬅️ До Панелі Адміністратора", callback_data="developer_menu")]])
        )
    else:
        await message.answer("❌ Помилка збереження матеріалу.", reply_markup=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="⬅️ До Панелі Адміністратора", callback_data="developer_menu")]]))
    
    await state.clear()


# ==================== Tests Editor ====================

async def developer_tests_menu(callback: CallbackQuery):
    """Головне меню редагування тестів — вибір посади."""
    if not await _ensure_developer(callback):
        return
    
    buttons = []
    for i, role in enumerate(AVAILABLE_ROLES):
        short_name = role[:25] + "..." if len(role) > 28 else role
        buttons.append([InlineKeyboardButton(
            text=short_name,
            callback_data=f"dev_test_role|{i}"
        )])
    
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_main_study")])
    
    await _edit_or_answer(
        callback.message,
        "📝 <b>Редагування тестів</b>\n\n"
        "Оберіть посаду для редагування тестів:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    await callback.answer()


async def developer_tests_select_day(callback: CallbackQuery, state: FSMContext):
    """Вибір дня для тесту."""
    if not await _ensure_developer(callback):
        return
    
    try:
        _, role_part = callback.data.split("|", 1)
        role_index = int(role_part)
        role = AVAILABLE_ROLES[role_index]
    except (ValueError, IndexError):
        await callback.answer("Помилка ролі.", show_alert=True)
        return
    
    await state.update_data(test_role=role, test_role_index=role_index)
    
    kb = await _build_dev_days_picker_keyboard(role, "dev_test_day", "dev_tests_menu", page=0)
    
    await _edit_or_answer(
        callback.message,
        f"📝 <b>Редагування тестів</b>\n\n"
        f"Посада: <b>{role}</b>\n\n"
        f"Оберіть день:",
        reply_markup=kb
    )
    await callback.answer()


async def developer_videos_select_day(callback: CallbackQuery, state: FSMContext):
    """Вибір дня для редагування відео."""
    if not await _ensure_developer(callback):
        return
    
    try:
        _, role_part = callback.data.split("|", 1)
        role_index = int(role_part)
        if not (0 <= role_index < len(AVAILABLE_ROLES)):
            await callback.answer("Неправильний індекс ролі.", show_alert=True)
            return
        role = AVAILABLE_ROLES[role_index]
    except (ValueError, IndexError):
        await callback.answer("Помилка формату ролі.", show_alert=True)
        return
    
    await state.update_data(edit_video_role=role, edit_video_role_index=role_index)
    
    kb = await _build_dev_days_picker_keyboard(role, "dev_video_day", "dev_videos_menu", page=0)
    
    await _edit_or_answer(
        callback.message,
        f"🎥 <b>Редагування відео матеріалів</b>\n\n"
        f"Посада: <b>{role}</b>\n\n"
        f"Оберіть день:",
        reply_markup=kb
    )
    await callback.answer()


async def developer_tests_view(callback: CallbackQuery, state: FSMContext, day: Optional[int] = None, role_index: Optional[int] = None):
    """Перегляд поточного тесту."""
    if not await _ensure_developer(callback):
        return
    
    data = await state.get_data()
    
    try:
        if day is None or role_index is None:
            parts = callback.data.split("|")
            if len(parts) == 3: # dev_test_day|{day}|{role_index}
                day = int(parts[1])
                role_index = int(parts[2])
            else: # old format: dev_test_day|{day}
                day = int(parts[1])
                role_index = data.get("test_role_index", 0) # Fallback to stored role index
        
        role = AVAILABLE_ROLES[role_index]
    except (ValueError, IndexError, TypeError):
        await callback.answer("Помилка дня або ролі.", show_alert=True)
        return
    
    # role is already determined above from AVAILABLE_ROLES[role_index]
    # update state with current day/role if needed, though role is derived
    await state.update_data(test_day=day, test_role=role, test_role_index=role_index)
    
    test_material = await get_test_by_role_and_day(role, day)
    
    display_text = "Тест ще не створено."
    is_enabled = True # Default enabled
    
    if test_material:
        is_enabled = bool(test_material.get("is_enabled", 1))
        if test_material.get("content"):
            try:
                current_questions = json.loads(test_material["content"])
                formatted = format_test_display(current_questions)
                display_text = formatted if formatted else "Тест порожній."
            except json.JSONDecodeError:
                display_text = "⚠️ Помилка формату даних тесту."
    
    status_icon = "✅" if is_enabled else "zzz"
    status_text_display = "АКТИВНИЙ" if is_enabled else "ВИМКНЕНИЙ (не враховується)"
    
    text = (
        f"📝 <b>Тест: {role} — День {day}</b>\n"
        f"Статус: {status_icon} <b>{status_text_display}</b>\n\n"
        f"{display_text}\n\n"
        f"Натисніть «Редагувати», щоб змінити або створити тест."
    )
    
    toggle_btn_text = "🔴 Вимкнути тест" if is_enabled else "🟢 Увімкнути тест"
    toggle_action = "disable" if is_enabled else "enable"
    
    buttons = [
        [InlineKeyboardButton(text="✏️ Редагувати", callback_data="dev_test_edit")],
        [InlineKeyboardButton(text=toggle_btn_text, callback_data=f"dev_test_toggle:{toggle_action}")],
        [InlineKeyboardButton(text="⬅️ Назад", callback_data=f"dev_test_role|{role_index}")]
    ]
    
    await _edit_or_answer(
        callback.message,
        text,
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    await callback.answer()

async def developer_test_toggle(callback: CallbackQuery, state: FSMContext):
    if not await _ensure_developer(callback): return
    
    try:
        action = callback.data.split(":")[1]
        is_enabled = (action == "enable")
    except IndexError:
        return
    
    data = await state.get_data()
    role = data.get("test_role")
    day = data.get("test_day")
    role_index = data.get("test_role_index")
    
    if not role or day is None:
        await callback.answer("Помилка стану. Будь ласка, почніть спочатку.", show_alert=True)
        return

    # Якщо тесту ще немає, створимо порожній, щоб можна було змінити статус
    test_material = await get_test_by_role_and_day(role, day)
    if not test_material:
        await add_or_update_material(
            role=role, day=day, content_type="test", 
            title="Тест", content="[]", resource_url=None
        )

    await toggle_test_status(role, day, is_enabled)
    
    status_msg = "Тест увімкнено! Тепер він обов'язковий." if is_enabled else "Тест вимкнено! Студенти зможуть завершити день без нього."
    await callback.answer(status_msg, show_alert=True)
    
    # Refresh view without modifying callback.data
    await developer_tests_view(callback, state, day=day, role_index=role_index)

async def developer_videos_view(callback: CallbackQuery, state: FSMContext):
    """Показ поточного відео матеріалу з можливістю редагування."""
    if not await _ensure_developer(callback):
        return
    
    try:
        _, day_str = callback.data.split("|", 1)
        day = int(day_str)
    except (ValueError, IndexError):
        await callback.answer("Помилка формату.", show_alert=True)
        return
    
    data = await state.get_data()
    role = data.get("edit_video_role", "")
    await state.update_data(edit_video_day=day)
    
    # Отримуємо матеріал з БД, використовуючи новий тип для завантажених відео
    video_material = await get_material_by_role_day_type(role, day, "video_files")
    
    current_file_ids = []
    video_material_id = None
    if video_material and video_material.get("content"):
        try:
            current_file_ids = json.loads(video_material["content"])
            video_material_id = video_material.get("id")
        except json.JSONDecodeError:
            pass # Will be treated as no content
    
    await state.update_data(edit_video_material_id=video_material_id)
    
    if current_file_ids:
        preview = f"Завантажено відеофайлів: {len(current_file_ids)}"
        text = (
            f"🎥 <b>Відео матеріал</b>\n\n"
            f"Посада: <b>{role}</b>\n"
            f"День: <b>{day}</b>\n\n"
            f"<b>Поточний контент:</b>\n"
            f"{preview}\n\n"
            f"Натисніть «Редагувати», щоб змінити відеофайли."
        )
    else:
        text = (
            f"🎥 <b>Відео матеріал не знайдено</b>\n\n"
            f"Посада: <b>{role}</b>\n"
            f"День: <b>{day}</b>\n\n"
            f"Натисніть «Редагувати», щоб завантажити новий відео матеріал."
        )
    
    buttons = [
        [InlineKeyboardButton(text="✏️ Редагувати", callback_data="dev_video_edit")],
        [InlineKeyboardButton(text="⬅️ Назад", callback_data=f"dev_video_role|{data.get('edit_video_role_index')}")],
    ]
    
    await _edit_or_answer(
        callback.message,
        text,
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    await callback.answer()


async def developer_tests_edit_start(callback: CallbackQuery, state: FSMContext):
    """Початок редагування тесту."""
    if not await _ensure_developer(callback):
        return
        
    data = await state.get_data()
    role = data.get("test_role", "")
    day = data.get("test_day", 1)
    
    template = (
        "1. Питання\n"
        "а. Варіант 1\n"
        "б. Варіант 2 (х)\n"
        "в. Варіант 3\n\n"
        "2. Наступне питання\n"
        "а. Варіант А (х)\n"
        "б. Варіант Б"
    )
    
    text = (
        f"✏️ <b>Редагування тесту: {role} — День {day}</b>\n\n"
        f"Надішліть новий тест у наступному форматі (можна копіювати):\n\n"
        f"<code>{template}</code>\n\n"
        f"⚠️ <b>Важливо:</b>\n"
        f"— Правильну відповідь позначайте як <b>(х)</b> або <b>(x)</b>\n"
        f"— Кожне питання має мати мінімум 2 варіанти."
    )
    
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="❌ Скасувати", callback_data=f"dev_test_day|{day}")]
    ])
    
    await _edit_or_answer(callback.message, text, reply_markup=kb)
    await state.set_state(DeveloperStates.waiting_test_input)
    await callback.answer()


async def developer_video_edit_start(callback: CallbackQuery, state: FSMContext):
    """Початок редагування відео — запит нових відеофайлів."""
    if not await _ensure_developer(callback):
        return
    
    data = await state.get_data()
    role = data.get("edit_video_role", "")
    day = data.get("edit_video_day", 1)
    
    prompt = (
        f"🎥 <b>Завантаження відео матеріалів</b>\n\n"
        f"Посада: <b>{role}</b>\n"
        f"День: <b>{day}</b>\n\n"
        f"Надішліть один або кілька відеофайлів. Коли закінчите, будь ласка, напишіть <b>Готово</b>.\n"
        f"(Максимальний розмір відео файлу - 50МБ, тривалість до 1 хвилини)"
    )
    
    buttons = [
        [InlineKeyboardButton(text="❌ Скасувати", callback_data=f"dev_video_day|{day}")],
    ]
    
    await _edit_or_answer(
        callback.message,
        prompt,
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    await state.set_state(DeveloperStates.waiting_video_uploads)
    await state.update_data(video_file_ids=[])
    await callback.answer()


async def developer_tests_process_input(message: Message, state: FSMContext):
    """Обробка введеного тексту тесту."""
    text_input = message.text
    
    questions, errors = parse_test_input(text_input)
    
    if errors:
        error_msg = "\n".join([f"• {e}" for e in errors])
        await message.answer(
            f"❌ <b>Знайдено помилки у форматі:</b>\n\n{error_msg}\n\n"
            "Виправте помилки та надішліть текст знову, або натисніть /cancel для скасування."
        )
        return
        
    if not questions:
        await message.answer("❌ Не вдалося розпізнати жодного питання. Перевірте формат.")
        return
    
    data = await state.get_data()
    role = data.get("test_role", "")
    day = data.get("test_day", 1)
    
    # Save to DB
    json_content = json.dumps(questions, ensure_ascii=False)
    
    await add_or_update_material(
        role=role,
        day=day,
        content_type="test",
        title=f"Тест: {role} — День {day}",
        content=json_content,
        resource_url=None,
        order_index=999 # Tests usually go last
    )
    
    await message.answer(
        f"✅ <b>Тест успішно збережено!</b>\n"
        f"Розпізнано питань: {len(questions)}",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="⬅️ До перегляду тесту", callback_data=f"dev_test_day|{day}")]
        ])
    )
    await state.clear()


async def developer_process_video_uploads(message: Message, state: FSMContext):
    """Collects multiple video messages for video content until 'Готово' is received."""
    data = await state.get_data()
    video_file_ids = data.get("video_file_ids", [])
    confirmation_msg_id = data.get("video_confirmation_msg_id")

    if message.video:
        file_id = message.video.file_id
        video_file_ids.append(file_id)
        await state.update_data(video_file_ids=video_file_ids)

        new_text = f"✅ Отримано {len(video_file_ids)} відео. Надішліть наступне відео або напишіть 'Готово', щоб завершити."

        if confirmation_msg_id:
            try:
                await message.bot.edit_message_text(
                    chat_id=message.chat.id,
                    message_id=confirmation_msg_id,
                    text=new_text,
                )
            except TelegramBadRequest as e:
                if "message is not modified" not in str(e):
                    sent_msg = await message.answer(new_text)
                    await state.update_data(video_confirmation_msg_id=sent_msg.message_id)
        else:
            sent_msg = await message.answer(new_text)
            await state.update_data(video_confirmation_msg_id=sent_msg.message_id)

    elif message.text and message.text.lower().rstrip('.! ') == 'готово':
        if not video_file_ids:
            await message.answer("❌ Ви не надіслали жодного відео. Надішліть відео або скасуйте редагування.")
            return
            
        if confirmation_msg_id:
            try:
                await message.bot.delete_message(chat_id=message.chat.id, message_id=confirmation_msg_id)
            except Exception:
                pass

        await _finalize_video_update(message, state, video_file_ids)
    else:
        await message.answer("Будь ласка, надішліть відеофайл або напишіть 'Готово' для завершення.")


async def developer_videos_menu(callback: CallbackQuery):
    """Головне меню редагування відео матеріалів — вибір посади."""
    if not await _ensure_developer(callback):
        return
    
    buttons = []
    for i, role in enumerate(AVAILABLE_ROLES):
        short_name = role[:25] + "..." if len(role) > 28 else role
        buttons.append([InlineKeyboardButton(
            text=short_name,
            callback_data=f"dev_video_role|{i}"
        )])
    
    buttons.append([InlineKeyboardButton(text="❌ Скасувати", callback_data="dev_main_study")])
    
    await _edit_or_answer(
        callback.message,
        "🎥 <b>Редагування відео матеріалів</b>\n\n"
        "Оберіть посаду для редагування відео:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    await callback.answer()


async def developer_photos_menu(callback: CallbackQuery):
    """Головне меню редагування фото матеріалів — вибір посади."""
    if not await _ensure_developer(callback):
        return
    
    buttons = []
    for i, role in enumerate(AVAILABLE_ROLES):
        short_name = role[:25] + "..." if len(role) > 28 else role
        buttons.append([InlineKeyboardButton(
            text=short_name,
            callback_data=f"dev_photo_role|{i}"
        )])
    
    buttons.append([InlineKeyboardButton(text="❌ Скасувати", callback_data="dev_main_study")])
    
    await _edit_or_answer(
        callback.message,
        "🖼 <b>Редагування фото матеріалів</b>\n\n"
        "Оберіть посаду для редагування фото:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    await callback.answer()

async def developer_photos_select_day(callback: CallbackQuery, state: FSMContext):
    """Вибір дня для редагування фото."""
    if not await _ensure_developer(callback):
        return
    
    try:
        _, role_part = callback.data.split("|", 1)
        role_index = int(role_part)
        if not (0 <= role_index < len(AVAILABLE_ROLES)):
            await callback.answer("Неправильний індекс ролі.", show_alert=True)
            return
        role = AVAILABLE_ROLES[role_index]
    except (ValueError, IndexError):
        await callback.answer("Помилка формату ролі.", show_alert=True)
        return
    
    await state.update_data(edit_photo_role=role, edit_photo_role_index=role_index)
    
    kb = await _build_dev_days_picker_keyboard(role, "dev_photo_day", "dev_photos_menu", page=0)
    
    await _edit_or_answer(
        callback.message,
        f"🖼 <b>Редагування фото матеріалів</b>\n\n"
        f"Посада: <b>{role}</b>\n\n"
        f"Оберіть день:",
        reply_markup=kb
    )
    await callback.answer()


async def developer_days_page_callback(callback: CallbackQuery, state: FSMContext):
    """Обробляє пагінацію вибору днів у панелі розробника."""
    if not await _ensure_developer(callback):
        return
    try:
        prefix, page_str = callback.data.split("|", 1)
        page = int(page_str)
    except (ValueError, IndexError):
        await callback.answer()
        return

    data = await state.get_data()
    
    if prefix == "dev_mat_day_pg":
        role = data.get("edit_role", AVAILABLE_ROLES[0])
        kb = await _build_dev_days_picker_keyboard(role, "dev_mat_day", "dev_materials_menu", page=page)
        title = f"📝 <b>Редагування матеріалів</b>\n\nПосада: <b>{role}</b>\n\nОберіть день:"
    elif prefix == "dev_test_day_pg":
        role = data.get("test_role", AVAILABLE_ROLES[0])
        kb = await _build_dev_days_picker_keyboard(role, "dev_test_day", "dev_tests_menu", page=page)
        title = f"📝 <b>Редагування тестів</b>\n\nПосада: <b>{role}</b>\n\nОберіть день:"
    elif prefix == "dev_video_day_pg":
        role = data.get("edit_video_role", AVAILABLE_ROLES[0])
        kb = await _build_dev_days_picker_keyboard(role, "dev_video_day", "dev_videos_menu", page=page)
        title = f"🎥 <b>Редагування відео матеріалів</b>\n\nПосада: <b>{role}</b>\n\nОберіть день:"
    elif prefix == "dev_photo_day_pg":
        role = data.get("edit_photo_role", AVAILABLE_ROLES[0])
        kb = await _build_dev_days_picker_keyboard(role, "dev_photo_day", "dev_photos_menu", page=page)
        title = f"🖼 <b>Редагування фото матеріалів</b>\n\nПосада: <b>{role}</b>\n\nОберіть день:"
    else:
        await callback.answer()
        return

    await _edit_or_answer(callback.message, title, reply_markup=kb)
    await callback.answer()

async def developer_photos_view(callback: CallbackQuery, state: FSMContext):
    """Показ поточного фото матеріалу з можливістю редагування."""
    if not await _ensure_developer(callback):
        return
    
    try:
        _, day_str = callback.data.split("|", 1)
        day = int(day_str)
    except (ValueError, IndexError):
        await callback.answer("Помилка формату.", show_alert=True)
        return
    
    data = await state.get_data()
    role = data.get("edit_photo_role", "")
    await state.update_data(edit_photo_day=day)
    
    # Отримуємо матеріал з БД, використовуючи новий тип для завантажених фото
    photo_material = await get_material_by_role_day_type(role, day, "photo_files")
    
    current_file_ids = []
    photo_material_id = None
    if photo_material and photo_material.get("content"):
        try:
            current_file_ids = json.loads(photo_material["content"])
            photo_material_id = photo_material.get("id")
        except json.JSONDecodeError:
            pass # Will be treated as no content
    
    await state.update_data(edit_photo_material_id=photo_material_id)
    # Also set generic edit state for full view handler
    await state.update_data(edit_role=role, edit_day=day)
    
    if current_file_ids:
        preview = f"Завантажено фотофайлів: {len(current_file_ids)}"
        text = (
            f"🖼 <b>Фото матеріал</b>\n\n"
            f"Посада: <b>{role}</b>\n"
            f"День: <b>{day}</b>\n\n"
            f"<b>Поточний контент:</b>\n"
            f"{preview}\n\n"
            f"Натисніть «Повний перегляд», щоб побачити або виправити фото."
        )
        buttons = [
            [InlineKeyboardButton(text="➕ Додати навчання", callback_data="dev_photo_edit")],
            [InlineKeyboardButton(text="🔍 Повний перегляд", callback_data="dev_mat_full_view|photo_files")],
            [InlineKeyboardButton(text="⬅️ Назад", callback_data=f"dev_photo_role|{data.get('edit_photo_role_index')}")],
        ]
    else:
        text = (
            f"🖼 <b>Фото матеріал не знайдено</b>\n\n"
            f"Посада: <b>{role}</b>\n"
            f"День: <b>{day}</b>\n\n"
            f"Натисніть «Додати навчання», щоб завантажити новий фото матеріал."
        )
        buttons = [
            [InlineKeyboardButton(text="➕ Додати навчання", callback_data="dev_photo_edit")],
            [InlineKeyboardButton(text="⬅️ Назад", callback_data=f"dev_photo_role|{data.get('edit_photo_role_index')}")],
        ]
    
    await _edit_or_answer(
        callback.message,
        text,
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    await callback.answer()

async def developer_photo_edit_start(callback: CallbackQuery, state: FSMContext):
    """Початок редагування фото — запит нових фотофайлів."""
    if not await _ensure_developer(callback):
        return
    
    data = await state.get_data()
    role = data.get("edit_photo_role", "")
    day = data.get("edit_photo_day", 1)
    
    prompt = (
        f"🖼 <b>Завантаження фото матеріалів</b>\n\n"
        f"Посада: <b>{role}</b>\n"
        f"День: <b>{day}</b>\n\n"
        f"Надішліть один або кілька фотофайлів (jpg, jpeg, png). Коли закінчите, будь ласка, напишіть <b>Готово</b>."
    )
    
    buttons = [
        [InlineKeyboardButton(text="❌ Скасувати", callback_data=f"dev_photo_day|{day}")],
    ]
    
    await _edit_or_answer(
        callback.message,
        prompt,
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    await state.set_state(DeveloperStates.waiting_photo_uploads)
    await state.update_data(photo_file_ids=[])
    await callback.answer()

async def _finalize_photo_update(message: Message, state: FSMContext, photo_file_ids: list[str]):
    """Обробка завантажених фотофайлів і збереження."""
    if not photo_file_ids:
        await message.answer("❌ Немає фотофайлів для збереження.")
        return

    # edit_role and edit_day are needed by _save_pending_material
    data = await state.get_data()
    role = data.get("edit_photo_role", "")
    day = data.get("edit_photo_day", 1)
    material_id = data.get("edit_photo_material_id")

    await state.update_data(
        pending_content=json.dumps(photo_file_ids),
        pending_content_type="photo_files",
        edit_role=role,
        edit_day=day,
        pending_material_id=material_id
    )

    success, role, content_type, day = await _save_pending_material(state)
    
    if success:
        await message.answer(
            f"✅ <b>Матеріал оновлено!</b>\n\nПосада: {role}\nДень: {day}\nТип: {content_type}\nЗавантажено файлів: {len(photo_file_ids)}",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="⬅️ До Панелі Адміністратора", callback_data="developer_menu")]])
        )
    else:
        await message.answer("❌ Помилка збереження матеріалу.", reply_markup=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="⬅️ До Панелі Адміністратора", callback_data="developer_menu")]]))
    
    await state.clear()


async def developer_process_photo_uploads(message: Message, state: FSMContext):
    """Collects multiple photo messages for photo content until 'Готово' is received."""
    data = await state.get_data()
    photo_file_ids = data.get("photo_file_ids", [])
    confirmation_msg_id = data.get("photo_confirmation_msg_id")

    if message.photo:
        file_id = message.photo[-1].file_id
        photo_file_ids.append(file_id)
        await state.update_data(photo_file_ids=photo_file_ids)

        new_text = f"✅ Отримано {len(photo_file_ids)} фото. Надішліть наступне фото або напишіть 'Готово', щоб завершити."

        if confirmation_msg_id:
            try:
                await message.bot.edit_message_text(
                    chat_id=message.chat.id,
                    message_id=confirmation_msg_id,
                    text=new_text,
                )
            except TelegramBadRequest as e:
                if "message is not modified" not in str(e):
                    sent_msg = await message.answer(new_text)
                    await state.update_data(photo_confirmation_msg_id=sent_msg.message_id)
        else:
            sent_msg = await message.answer(new_text)
            await state.update_data(photo_confirmation_msg_id=sent_msg.message_id)

    elif message.text and message.text.lower().rstrip('.! ') == 'готово':
        if not photo_file_ids:
            await message.answer("❌ Ви не надіслали жодного фото. Надішліть фото або скасуйте редагування.")
            return
            
        if confirmation_msg_id:
            try:
                await message.bot.delete_message(chat_id=message.chat.id, message_id=confirmation_msg_id)
            except Exception:
                pass

        await _finalize_photo_update(message, state, photo_file_ids)
    else:
        if message.text:
             await message.answer("Будь ласка, надішліть фотофайл або напишіть 'Готово' для завершення.")

# ==================== Syllabus Editor ====================

async def developer_syllabus_menu(callback: CallbackQuery):
    """Головне меню редагування змістів — вибір посади."""
    if not await _ensure_developer(callback):
        return
    
    buttons = []
    for i, role in enumerate(AVAILABLE_ROLES):
        short_name = role[:25] + "..." if len(role) > 28 else role
        buttons.append([InlineKeyboardButton(
            text=short_name,
            callback_data=f"dev_syl_role|{i}"
        )])
    
    buttons.append([InlineKeyboardButton(text="❌ Скасувати", callback_data="dev_main_study")])
    
    await _edit_or_answer(
        callback.message,
        "📚 <b>Редагування змістів (Syllabus)</b>\n\n"
        "Оберіть посаду для редагування змісту:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    await callback.answer()

async def developer_syllabus_view(callback: CallbackQuery, state: FSMContext, role_index: Optional[int] = None):
    """Перегляд поточного змісту для обраної посади з можливістю ввімкнення."""
    if not await _ensure_developer(callback):
        return
    
    try:
        if role_index is None:
            _, role_part = callback.data.split("|", 1)
            role_index = int(role_part)
        
        role = AVAILABLE_ROLES[role_index]
    except (ValueError, IndexError, TypeError):
        await callback.answer("Помилка ролі.", show_alert=True)
        return
    
    await state.update_data(syllabus_role=role, syllabus_role_index=role_index)
    
    # Syllabus is stored with day=0 and content_type='syllabus'
    syllabus_material = await get_material_by_role_day_type(role, 0, "syllabus")
    
    is_enabled = True
    if syllabus_material:
        is_enabled = bool(syllabus_material.get("is_enabled", 1))
        content = syllabus_material.get("content", "")
    else:
        content = ""

    display_content = content[:500] + "..." if len(content) > 500 else (content or "(порожньо)")
    
    status_icon = "✅" if is_enabled else "zzz"
    status_text = "АКТИВНИЙ" if is_enabled else "ВИМКНЕНИЙ (прихований від стажерів)"

    text = (
        f"📚 <b>Зміст: {role}</b>\n"
        f"Статус: {status_icon} <b>{status_text}</b>\n\n"
        f"{display_content}\n\n"
        f"Натисніть «Редагувати», щоб змінити текст змісту."
    )
    
    toggle_btn_text = "🔴 Вимкнути зміст" if is_enabled else "🟢 Увімкнути зміст"
    toggle_action = "disable" if is_enabled else "enable"

    buttons = [
        [InlineKeyboardButton(text="✏️ Редагувати", callback_data="dev_syl_edit")],
        [InlineKeyboardButton(text=toggle_btn_text, callback_data=f"dev_syl_toggle:{toggle_action}")],
        [InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_syllabus_menu")]
    ]
    
    await _edit_or_answer(
        callback.message,
        text,
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    await callback.answer()

async def developer_syllabus_toggle(callback: CallbackQuery, state: FSMContext):
    """Вмикає або вимикає зміст."""
    if not await _ensure_developer(callback): return
    
    try:
        action = callback.data.split(":")[1]
        is_enabled = (action == "enable")
    except IndexError:
        return
    
    data = await state.get_data()
    role = data.get("syllabus_role")
    role_index = data.get("syllabus_role_index")
    
    if not role or role_index is None:
        await callback.answer("Помилка стану.", show_alert=True)
        return

    syllabus_material = await get_material_by_role_day_type(role, 0, "syllabus")
    if not syllabus_material:
        await add_or_update_material(
            role=role,
            day=0,
            content_type="syllabus",
            title=f"Зміст: {role}",
            content="",
            resource_url=None,
            order_index=0
        )

    async with aiosqlite.connect(DB_PATH) as db:
        await db.execute(
            "UPDATE materials SET is_enabled = ? WHERE role = ? AND day = 0 AND content_type = 'syllabus'",
            (is_enabled, role)
        )
        await db.commit()
    
    status_msg = "Зміст увімкнено!" if is_enabled else "Зміст вимкнено!"
    await callback.answer(status_msg, show_alert=True)
    
    # Refresh view
    await developer_syllabus_view(callback, state, role_index=role_index)

async def developer_syllabus_edit(callback: CallbackQuery, state: FSMContext):
    """Початок редагування змісту."""
    if not await _ensure_developer(callback):
        return
        
    data = await state.get_data()
    role = data.get("syllabus_role", "")
    role_index = data.get("syllabus_role_index")
    
    if not role or role_index is None:
        await callback.answer("Помилка стану (роль не обрана).", show_alert=True)
        return
    
    text = (
        f"✏️ <b>Редагування змісту: {role}</b>\n\n"
        f"Надішліть новий текст змісту. Ви можете використовувати HTML-розмітку.\n\n"
        f"<b>Приклад оформлення:</b>\n"
        f"<code>&lt;b&gt;День 1&lt;/b&gt;: Знайомство з пекарнею\n"
        f"— Стандарти обслуговування\n"
        f"— &lt;i&gt;Техніка безпеки&lt;/i&gt;\n\n"
        f"&lt;b&gt;День 2&lt;/b&gt;: Робота з касою\n"
        f"— Розрахунок клієнтів</code>"
    )
    
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="❌ Скасувати", callback_data=f"dev_syl_role|{role_index}")]
    ])
    
    await _edit_or_answer(callback.message, text, reply_markup=kb)
    await state.set_state(DeveloperStates.waiting_syllabus_text)
    await callback.answer()

async def developer_process_syllabus(message: Message, state: FSMContext):
    """Збереження змісту."""
    text_input = message.text
    if not text_input:
        await message.answer("❌ Текст не може бути порожнім.")
        return
        
    data = await state.get_data()
    role = data.get("syllabus_role", "")
    role_index = data.get("syllabus_role_index")
    
    if not role:
        await message.answer("❌ Помилка стану (роль не визначена). Спробуйте знову з меню.")
        await state.clear()
        return

    # Save to DB with day=0 and type='syllabus'
    await add_or_update_material(
        role=role,
        day=0,
        content_type="syllabus",
        title=f"Зміст: {role}",
        content=text_input,
        resource_url=None,
        order_index=0
    )
    
    back_cb = f"dev_syl_role|{role_index}" if role_index is not None else "dev_syllabus_menu"
    await message.answer(
        "✅ <b>Зміст успішно збережено!</b>",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="⬅️ До перегляду", callback_data=back_cb)]
        ])
    )
    await state.clear()


# ==================== Notification handlers ====================

async def _save_pending_material(state: FSMContext) -> tuple[bool, str, str, int]:
    """Зберігає pending матеріал. Повертає (success, role, content_type, day)."""
    logger = get_logger()
    data = await state.get_data()
    role = data.get("edit_role") or data.get("edit_video_role") or data.get("edit_photo_role", "")
    day = data.get("edit_day") or data.get("edit_video_day") or data.get("edit_photo_day", 1)
    content_type = data.get("pending_content_type", "text")
    content = data.get("pending_content", "")
    material_id = data.get("pending_material_id") or data.get("edit_photo_material_id") or data.get("edit_video_material_id")
    
    try:
        if content_type == "text":
            if material_id:
                await update_material_content(material_id, content)
            else:
                await add_or_update_material(
                    role=role,
                    day=day,
                    content_type="text",
                    title=f"День {day} — Текст",
                    content=content,
                    resource_url=None,
                    order_index=0,
                )
        elif content_type == "video": # External video URL
            if material_id:
                await update_material_resource_url(material_id, content)
            else:
                await add_or_update_material(
                    role=role,
                    day=day,
                    content_type="video",
                    title=f"День {day} — Відео (посилання)",
                    content=None,
                    resource_url=content,
                    order_index=0,
                )
        elif content_type == "video_files": # Uploaded video files
            if material_id:
                # For video files, we store JSON list of file_ids in content
                await update_material_content(material_id, content)
                await update_material_resource_url(material_id, None) # Clear URL for uploaded files
            else:
                await add_or_update_material(
                    role=role,
                    day=day,
                    content_type="video_files",
                    title=f"День {day} — Відео",
                    content=content, # JSON string of file_ids
                    resource_url=None,
                    order_index=0,
                )
        elif content_type == "photo_files": # Uploaded photo files
            if material_id:
                await update_material_content(material_id, content)
                await update_material_resource_url(material_id, None)
            else:
                await add_or_update_material(
                    role=role,
                    day=day,
                    content_type="photo_files",
                    title=f"День {day} — Фото",
                    content=content, # JSON string of file_ids
                    resource_url=None,
                    order_index=0,
                )
        
        # Schedule the embeddings build to run in the background.
        logger.info("Material saved, scheduling embeddings build.")
        try:
            asyncio.create_task(build_and_reset_embeddings())
        except Exception as emb_err:
            logger.warning(f"Failed to schedule embeddings build: {emb_err}")
        
        return True, role, content_type, day
    except Exception as e:
        logger.error(f"Failed to save material: {e}", exc_info=True)
        return False, role, content_type, day


async def notify_decision_no(callback: CallbackQuery, state: FSMContext):
    """Користувач обрав НІ — зберігаємо без оповіщення."""
    success, role, content_type, day = await _save_pending_material(state)
    
    if success:
        await callback.message.edit_text(
            f"✅ <b>Матеріал оновлено!</b>\n\n"
            f"Посада: {role}\n"
            f"День: {day}\n"
            f"Тип: {content_type}\n\n"
            f"<i>Оповіщення не надсилалось.</i>",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="⬅️ До Панелі Адміністратора", callback_data="developer_menu")]
            ])
        )
    else:
        await callback.message.edit_text(
            "❌ Помилка збереження матеріалу.",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="⬅️ До Панелі Адміністратора", callback_data="developer_menu")]
            ])
        )
    
    await state.clear()
    await callback.answer()


async def notify_decision_yes(callback: CallbackQuery, state: FSMContext):
    """Користувач обрав ТАК — просимо текст оповіщення."""
    data = await state.get_data()
    role = data.get("edit_role") or data.get("edit_video_role") or data.get("edit_photo_role", "")
    day = data.get("edit_day") or data.get("edit_video_day") or data.get("edit_photo_day", 1)
    
    await state.set_state(DeveloperStates.waiting_notify_message)
    
    await callback.message.edit_text(
        f"📢 <b>Оповіщення про зміни</b>\n\n"
        f"Посада: {role}\n"
        f"День: {day}\n\n"
        f"Напишіть опис змін, який отримають користувачі.\n\n"
        f"<b>Приклад:</b>\n"
        f"<i>🔄 Оновлено навчальний матеріал:\n"
        f"• Додано нову інформацію про стандарти обслуговування\n"
        f"• Виправлено помилки в попередній версії</i>",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="❌ Скасувати оповіщення", callback_data="notify_cancel")]
        ])
    )
    await callback.answer()


async def notify_cancel(callback: CallbackQuery, state: FSMContext):
    """Скасування оповіщення — зберігаємо без нього."""
    await notify_decision_no(callback, state)


async def process_notify_message(message: Message, state: FSMContext):
    """Обробка тексту оповіщення та хвильова розсилка."""
    notify_text = message.text.strip() if message.text else ""
    if not notify_text:
        await message.answer("❌ Текст оповіщення не може бути порожнім.")
        return
    
    # Спочатку зберігаємо матеріал
    success, role, content_type, day = await _save_pending_material(state)
    
    if not success:
        await message.answer(
            "❌ Помилка збереження матеріалу.",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="⬅️ До Панелі Адміністратора", callback_data="developer_menu")]
            ])
        )
        await state.clear()
        return

    from database.material_notifications import create_material_change_event, create_recipients_for_event
    from bot.services.wave_broadcaster import start_event_wave_broadcast

    event_id = await create_material_change_event(role, day, content_type, notify_text, [notify_text])
    recipients_count = await create_recipients_for_event(event_id, role)
    start_event_wave_broadcast(message.bot, event_id)
    waves_total = max(1, (recipients_count + 39) // 40)
    
    await message.answer(
        f"✅ <b>Матеріал оновлено та хвильову розсилку запущено!</b>\n\n"
        f"📚 Посада: <b>{role}</b>\n"
        f"📅 День: <b>{day}</b>\n"
        f"📝 Тип: <b>{content_type}</b>\n\n"
        f"👥 Всього отримувачів: <b>{recipients_count}</b>\n"
        f"🌊 Кількість хвиль: <b>{waves_total}</b> (по 40 осіб на годину)\n"
        f"⏳ Першу хвилю надіслано негайно.",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="⬅️ До Панелі Адміністратора", callback_data="developer_menu")]
        ])
    )
    
    await state.clear()


# ==================== Material Acknowledgment Analytics ====================

async def developer_ack_analytics_menu(callback: CallbackQuery):
    """Головне меню аналітики ознайомлення — вибір категорії."""
    if not await _ensure_developer(callback, allow_observer=True):
        return

    text = (
        "📋 <b>Аналітика ознайомлення з матеріалами</b>\n"
        "───────────────────\n"
        "Оберіть категорію співробітників для перегляду звітів про ознайомлення зі змінами:"
    )
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="👔 Керівники", callback_data="dev_ack_cat:керівник")],
        [InlineKeyboardButton(text="👷 Працівники", callback_data="dev_ack_cat:працівник")],
        [InlineKeyboardButton(text="🎓 Стажери", callback_data="dev_ack_cat:стажер")],
        [InlineKeyboardButton(text="⬅️ Назад до аналітики", callback_data="dev_main_analyt")]
    ])
    await _edit_or_answer(callback.message, text, reply_markup=kb)
    await callback.answer()


async def developer_ack_category_events_view(callback: CallbackQuery):
    """Список дат змін навчальних матеріалів для обраної категорії з пагінацією."""
    if not await _ensure_developer(callback, allow_observer=True):
        return

    # Callback formats: dev_ack_cat:{cat} or dev_ack_cat_pg:{cat}:{page}
    parts = callback.data.split(":")
    category = parts[1]
    page = int(parts[2]) if len(parts) > 2 else 0

    from database.material_notifications import get_events_for_category, prune_old_material_events
    await prune_old_material_events(days=90)
    events = await get_events_for_category(category, days_limit=90)

    cat_titles = {
        "керівник": "👔 Керівники",
        "працівник": "👷 Працівники",
        "стажер": "🎓 Стажери"
    }
    cat_title = cat_titles.get(category.lower(), category.title())

    if not events:
        text = (
            f"📋 <b>Аналітика ознайомлення — {cat_title}</b>\n"
            f"───────────────────\n\n"
            f"<i>За останні 90 днів оновлень матеріалів для цієї категорії не знайдено.</i>"
        )
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_an_ack_menu")]
        ])
        await _edit_or_answer(callback.message, text, reply_markup=kb)
        await callback.answer()
        return

    PER_PAGE = 6
    pages_total = max(1, (len(events) + PER_PAGE - 1) // PER_PAGE)
    page = max(0, min(page, pages_total - 1))
    
    start_idx = page * PER_PAGE
    paginated_events = events[start_idx:start_idx + PER_PAGE]

    buttons = []
    for ev in paginated_events:
        created_str = ev.get('created_at', '')
        try:
            dt = datetime.fromisoformat(created_str.replace("Z", "+00:00"))
            dt_formatted = dt.strftime("%d.%m %H:%M")
        except Exception:
            dt_formatted = created_str[:16]

        btn_text = f"📅 {dt_formatted} — {ev['role']} (День {ev['day']})"
        buttons.append([InlineKeyboardButton(
            text=btn_text,
            callback_data=f"dev_ack_ev:{ev['id']}:{category}"
        )])

    # Рядок пагінації якщо подій більше ніж 6
    if pages_total > 1:
        nav_row = []
        if page > 0:
            nav_row.append(InlineKeyboardButton(text="⬅️", callback_data=f"dev_ack_cat_pg:{category}:{page - 1}"))
        nav_row.append(InlineKeyboardButton(text=f"📄 {page + 1}/{pages_total}", callback_data="ignore"))
        if page < pages_total - 1:
            nav_row.append(InlineKeyboardButton(text="➡️", callback_data=f"dev_ack_cat_pg:{category}:{page + 1}"))
        buttons.append(nav_row)

    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_an_ack_menu")])

    text = (
        f"📋 <b>Аналітика ознайомлення — {cat_title}</b>\n"
        f"───────────────────\n"
        f"Оберіть дату та подію зміни матеріалу для перегляду звіту:"
    )
    await _edit_or_answer(callback.message, text, reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons))
    await callback.answer()


async def developer_ack_event_card_view(callback: CallbackQuery):
    """Картка події: перевірка 75h порогу (таймер або повний звіт)."""
    if not await _ensure_developer(callback, allow_observer=True):
        return

    # Format: dev_ack_ev:{event_id}:{category}
    parts = callback.data.split(":")
    event_id = int(parts[1])
    category = parts[2]

    from database.material_notifications import get_event_by_id, get_event_analytics_for_category
    event = await get_event_by_id(event_id)
    if not event:
        await callback.answer("Подію не знайдено.", show_alert=True)
        return

    stats = await get_event_analytics_for_category(event_id, category)

    cat_titles = {
        "керівник": "👔 Керівники",
        "працівник": "👷 Працівники",
        "стажер": "🎓 Стажери"
    }
    cat_title = cat_titles.get(category.lower(), category.title())

    created_str = event.get('created_at', '')
    try:
        dt = datetime.fromisoformat(created_str.replace("Z", "+00:00"))
        dt_formatted = dt.strftime("%d.%m.%Y %H:%M")
    except Exception:
        try:
            dt = datetime.strptime(created_str, "%Y-%m-%d %H:%M:%S")
            dt_formatted = dt.strftime("%d.%m.%Y %H:%M")
        except Exception:
            dt = datetime.utcnow()
            dt_formatted = created_str

    time_passed = datetime.utcnow() - dt
    threshold_seconds = 75 * 3600

    # Перевірка 75-годинного порогу
    if time_passed.total_seconds() < threshold_seconds:
        rem_sec = int(threshold_seconds - time_passed.total_seconds())
        rem_hours = rem_sec // 3600
        rem_mins = (rem_sec % 3600) // 60

        text = (
            f"⏳ <b>Аналітика ознайомлення (Формування звіту)</b>\n"
            f"───────────────────\n"
            f"📚 Посада: <b>{event['role']}</b> | 📅 День: <b>{event['day']}</b>\n"
            f"📝 Зміна: <b>{event.get('description', '')}</b>\n"
            f"🕒 Створено: <b>{dt_formatted}</b>\n\n"
            f"👥 Категорія: <b>{cat_title}</b>\n"
            f"• Всього отримувачів: <b>{stats['total']}</b>\n"
            f"• Вже ознайомились: <b>{stats['acknowledged_count']}</b> ({stats['ack_percent']}%)\n\n"
            f"⚠️ <i>Аналітика формується лише через 75 годин після повідомлення про зміни.</i>\n\n"
            f"⏳ <b>До повного формування звіту залишилось:</b> {rem_hours} год {rem_mins} хв."
        )
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🔄 Оновити статус", callback_data=f"dev_ack_ev:{event_id}:{category}")],
            [InlineKeyboardButton(text="⬅️ Назад до дат", callback_data=f"dev_ack_cat:{category}")]
        ])
    else:
        text = (
            f"📊 <b>Аналітика ознайомлення</b>\n"
            f"───────────────────\n"
            f"📚 Посада: <b>{event['role']}</b> | 📅 День: <b>{event['day']}</b>\n"
            f"📝 Зміна: <b>{event.get('description', '')}</b>\n"
            f"🕒 Дата розсилки: <b>{dt_formatted}</b>\n\n"
            f"👥 Категорія: <b>{cat_title}</b>\n"
            f"• Всього в категорії: <b>{stats['total']}</b> осіб\n"
            f"• ✅ <b>Ознайомились:</b> {stats['acknowledged_count']} ({stats['ack_percent']}%)\n"
            f"  └ ⏳ з них із запізненням (>72 год): <b>{stats['acknowledged_late_count']}</b>\n"
            f"• ❌ <b>Не ознайомились:</b> {stats['unacknowledged_count']} ({stats['unack_percent']}%)\n"
            f"───────────────────\n"
            f"Оберіть групу для перегляду списку по магазинах:"
        )
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text=f"✅ Ознайомились ({stats['acknowledged_count']})", callback_data=f"dev_ack_shops:{event_id}:{category}:ack:0")],
            [InlineKeyboardButton(text=f"❌ Не ознайомились ({stats['unacknowledged_count']})", callback_data=f"dev_ack_shops:{event_id}:{category}:noack:0")],
            [InlineKeyboardButton(text="⬅️ Назад до дат", callback_data=f"dev_ack_cat:{category}")]
        ])

    await _edit_or_answer(callback.message, text, reply_markup=kb)
    await callback.answer()


async def developer_ack_shops_list_view(callback: CallbackQuery):
    """Список магазинів для обраної групи (ознайомлені/неознайомлені) з пагінацією."""
    if not await _ensure_developer(callback, allow_observer=True):
        return

    # Format: dev_ack_shops:{event_id}:{category}:{status_key}:{page}
    parts = callback.data.split(":")
    event_id = int(parts[1])
    category = parts[2]
    status_key = parts[3]
    page = int(parts[4]) if len(parts) > 4 else 0
    is_ack = (status_key == "ack")

    from database.material_notifications import get_event_shop_breakdown
    shops = await get_event_shop_breakdown(event_id, category, is_ack)

    status_title = "✅ Ознайомлені" if is_ack else "❌ Не ознайомлені"

    if not shops:
        text = f"📋 <b>{status_title} по магазинах</b>\n\n<i>Користувачів у цій групі не знайдено.</i>"
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="⬅️ Назад", callback_data=f"dev_ack_ev:{event_id}:{category}")]
        ])
        await _edit_or_answer(callback.message, text, reply_markup=kb)
        await callback.answer()
        return

    PER_PAGE = 6
    pages_total = max(1, (len(shops) + PER_PAGE - 1) // PER_PAGE)
    page = max(0, min(page, pages_total - 1))
    
    start_idx = page * PER_PAGE
    paginated_shops = shops[start_idx:start_idx + PER_PAGE]

    buttons = []
    for s in paginated_shops:
        shop_name = s['shop']
        btn_text = f"🏪 {shop_name} ({s['user_count']} чол.)"
        buttons.append([InlineKeyboardButton(
            text=btn_text,
            callback_data=f"dev_ack_shop_u:{event_id}:{category}:{status_key}:{shop_name[:30]}"
        )])

    if pages_total > 1:
        nav_row = []
        if page > 0:
            nav_row.append(InlineKeyboardButton(text="⬅️", callback_data=f"dev_ack_shops:{event_id}:{category}:{status_key}:{page - 1}"))
        nav_row.append(InlineKeyboardButton(text=f"📄 {page + 1}/{pages_total}", callback_data="ignore"))
        if page < pages_total - 1:
            nav_row.append(InlineKeyboardButton(text="➡️", callback_data=f"dev_ack_shops:{event_id}:{category}:{status_key}:{page + 1}"))
        buttons.append(nav_row)

    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data=f"dev_ack_ev:{event_id}:{category}")])

    text = (
        f"📋 <b>{status_title} — Список магазинів</b>\n"
        f"───────────────────\n"
        f"Оберіть магазин для перегляду співробітників:"
    )
    await _edit_or_answer(callback.message, text, reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons))
    await callback.answer()


async def developer_ack_shop_users_view(callback: CallbackQuery):
    """Детальний список співробітників конкретного магазину."""
    if not await _ensure_developer(callback, allow_observer=True):
        return

    # Format: dev_ack_shop_u:{event_id}:{category}:{status_key}:{shop}
    parts = callback.data.split(":", 4)
    event_id = int(parts[1])
    category = parts[2]
    status_key = parts[3]
    shop = parts[4]
    is_ack = (status_key == "ack")

    from database.material_notifications import get_event_shop_users
    users = await get_event_shop_users(event_id, category, is_ack, shop)

    status_title = "✅ Ознайомлені" if is_ack else "❌ Не ознайомлені"

    lines = [
        f"🏪 <b>Магазин:</b> {shop}",
        f"📋 <b>Статус:</b> {status_title}",
        f"👥 <b>Кількість:</b> {len(users)} чол.",
        "───────────────────"
    ]

    if not users:
        lines.append("<i>Список порожній.</i>")
    else:
        for idx, u in enumerate(users, 1):
            name = u['full_name']
            username_str = f" (@{u['username']})" if u.get('username') else ""
            if is_ack:
                ack_str = u.get('acknowledged_at', '')
                try:
                    dt = datetime.fromisoformat(ack_str.replace("Z", "+00:00"))
                    dt_fmt = dt.strftime("%d.%m %H:%M")
                except Exception:
                    dt_fmt = ack_str[:16]
                late_badge = " ⏳ <i>(після 72 год)</i>" if u.get('is_late') else ""
                lines.append(f"{idx}. <b>{name}</b>{username_str} — {dt_fmt}{late_badge}")
            else:
                lines.append(f"{idx}. <b>{name}</b>{username_str} — ❌ <i>Не ознайомився (>72 год)</i>")

    text = "\n".join(lines)
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="⬅️ Назад до магазинів", callback_data=f"dev_ack_shops:{event_id}:{category}:{status_key}:0")]
    ])
    await _edit_or_answer(callback.message, text, reply_markup=kb)
    await callback.answer()


# ==================== Tokens Management ====================

async def developer_tokens_menu(callback: CallbackQuery):
    """Меню управління токенами з поясненнями."""
    if not await _ensure_developer(callback):
        return
    
    stats = await get_token_stats()
    
    text = (
        "🎟 <b>УПРАВЛІННЯ ТОКЕНАМИ</b>\n"
        "───────────────────\n"
        f"📊 <b>Статистика:</b>\n"
        f"• Активні (діючі): <b>{stats.get('active', 0)}</b>\n"
        f"• Використані: <b>{stats.get('used', 0)}</b>\n"
        f"• Прострочені: <b>{stats.get('expired', 0)}</b>\n"
        f"• Всього в базі: <b>{stats.get('total', 0)}</b>\n\n"
        
        f"⚙️ <b>Функції:</b>\n"
        f"🧹 <b>Очищення прострочених:</b> видаляє з бази даних усі посилання-запрошення, термін дії яких (24 години) вже закінчився і які так і не були використані.\n"
        "───────────────────\n"
        "<i>💡 Це допомагає підтримувати базу даних чистою та видаляти неактуальне сміття.</i>"
    )
    
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🧹 Очистити прострочені", callback_data="dev_tokens_cleanup")],
        [InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_main_other")],
    ])
    
    await _edit_or_answer(callback.message, text, reply_markup=kb)
    await callback.answer()


async def developer_tokens_cleanup(callback: CallbackQuery):
    """Очищення прострочених токенів."""
    if not await _ensure_developer(callback):
        return
    
    count = await cleanup_expired_tokens()
    await callback.answer(f"Оброблено {count} токенів", show_alert=True)

    # Оновлюємо меню зі свіжою статистикою
    stats = await get_token_stats()
    text = (
        "🎟 <b>УПРАВЛІННЯ ТОКЕНАМИ</b>\n"
        "───────────────────\n"
        f"📊 <b>Статистика:</b>\n"
        f"• Активні (діючі): <b>{stats.get('active', 0)}</b>\n"
        f"• Використані: <b>{stats.get('used', 0)}</b>\n"
        f"• Прострочені: <b>{stats.get('expired', 0)}</b>\n"
        f"• Всього в базі: <b>{stats.get('total', 0)}</b>\n\n"
        f"⚙️ <b>Функції:</b>\n"
        f"🧹 <b>Очищення прострочених:</b> видаляє з бази даних усі посилання-запрошення, термін дії яких (24 години) вже закінчився і які так і не були використані.\n"
        "───────────────────\n"
        "<i>💡 Це допомагає підтримувати базу даних чистою та видаляти неактуальне сміття.</i>"
    )
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🧹 Очистити прострочені", callback_data="dev_tokens_cleanup")],
        [InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_main_other")],
    ])
    await _edit_or_answer(callback.message, text, reply_markup=kb)


# ==================== Health Status ====================

async def developer_health_status(callback: CallbackQuery):
    """Показує статус здоров'я бота з покращеним інтерфейсом."""
    if not await _ensure_developer(callback):
        return
    
    status = get_health_status()
    
    health_icon = "🟢" if status["status"] == "healthy" else "🔴"
    scheduler_icon = "✅" if status["scheduler_healthy"] else "❌"
    reminder_icon = "✅" if status["reminder_healthy"] else "❌"
    
    from bot.services.health import health_check
    
    now = datetime.now(pytz.timezone(TIMEZONE))
    
    def format_delta(last_time):
        if not last_time: return "немає даних"
        delta = now - last_time
        minutes = int(delta.total_seconds() / 60)
        if minutes == 0: return "щойно"
        if minutes < 60: return f"{minutes} хв тому"
        return f"{minutes // 60} год {minutes % 60} хв тому"

    scheduler_delta = format_delta(health_check._last_scheduler_heartbeat)
    reminder_delta = format_delta(health_check._last_reminder_heartbeat)
    
    # Форматуємо дату очищення токенів
    token_cleanup = status.get('last_token_cleanup')
    if token_cleanup:
        try:
            token_cleanup = datetime.fromisoformat(str(token_cleanup).replace('Z', '+00:00')).astimezone(pytz.timezone(TIMEZONE)).strftime('%H:%M:%S (%d.%m)')
        except Exception:
            pass

    text = (
        f"🛡 <b>МОНІТОРИНГ СИСТЕМИ</b>\n"
        f"───────────────────\n"
        f"{health_icon} Стан системи: <b>{status['status'].upper()}</b>\n"
        f"⏱ Час роботи: <b>{status['uptime_human']}</b>\n"
        f"❌ Помилок за сесію: <b>{status['errors_count']}</b>\n\n"
        
        f"⚙️ <b>Життєздатність компонентів:</b>\n"
        f"{scheduler_icon} <b>Планувальник (Scheduler):</b>\n"
        f"  └ Активність: {scheduler_delta}\n"
        f"{reminder_icon} <b>Цикл нагадувань:</b>\n"
        f"  └ Активність: {reminder_delta}\n\n"
        
        f"📊 <b>Додаткові дані:</b>\n"
        f"🕒 Сервер: <code>{now.strftime('%H:%M:%S')}</code>\n"
        f"🧹 Очищення токенів: <code>{token_cleanup or 'немає'}</code>\n"
        f"───────────────────\n"
        f"<i>💡 Кнопка «Оновити» перевіряє актуальні дані датчиків активності без перезавантаження сторінки.</i>"
    )
    
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🔄 Оновити статус", callback_data="dev_health_status")],
        [InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_main_other")],
    ])
    
    await _edit_or_answer(callback.message, text, reply_markup=kb)
    await callback.answer()


# ---------- New User Management Handlers (from user profile) ----------

async def _rebuild_user_profile_view(callback: CallbackQuery, user_id: int, state: FSMContext):
    """Helper to refresh the user profile view after an action."""
    try:
        report = await get_user_days_report(user_id)
    except Exception:
        await _edit_or_answer(
            callback.message,
            "⚠️ Користувача не знайдено або дані видалено.",
            reply_markup=await _get_users_menu_kb(callback.from_user.id)
        )
        await state.clear()
        return

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(text="📅 Навчальні дні", callback_data=f"dev_days_manage:{user_id}"),
            InlineKeyboardButton(text="✍️ Змінити", callback_data=f"dev_user_modify:{user_id}")
        ],
        [
            InlineKeyboardButton(text="❌ Видалити", callback_data=f"dev_user_delete_confirm:{user_id}"),
            InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_users_menu")
        ]
    ])
    await _edit_or_answer(callback.message, report, reply_markup=kb)
    await state.clear()

# --- Learning Days Management ---

async def developer_days_manage_menu(callback: CallbackQuery, state: FSMContext, user_id: Optional[int] = None):
    if user_id is None:
        try:
            user_id = int(callback.data.split(":")[1])
        except (ValueError, IndexError):
            await callback.answer("Помилка ID користувача.", show_alert=True)
            return

    user = await get_user_details(user_id)
    if not user:
        await callback.answer("Користувача не знайдено.", show_alert=True)
        return

    overview = await get_days_overview(user_id)
    
    buttons = []
    for day, status in overview:
        action = "close" if status == DayStatus.OPEN else "open"
        icon = "🟢" if status == DayStatus.OPEN else "✅" if status == DayStatus.COMPLETED else "🔒"
        buttons.append([InlineKeyboardButton(
            text=f"{icon} День {day}",
            callback_data=f"dev_toggle_day:{user_id}:{day}:{action}"
        )])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад до профілю", callback_data=f"dev_user_profile_back:{user_id}")])
    
    await _edit_or_answer(
        callback.message,
        f"<b>Керування днями для {user.get('full_name', 'Користувача')}</b>\nНатисніть на день, щоб змінити його статус (відкрити/закрити).",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    await callback.answer()

async def developer_toggle_day_status(callback: CallbackQuery, state: FSMContext):
    try:
        parts = callback.data.split(":")
        user_id = int(parts[1])
        day = int(parts[2])
        action = parts[3]
    except (ValueError, IndexError):
        await callback.answer("Помилка даних.", show_alert=True)
        return

    if action == "open":
        await open_days_for_user(user_id, [day])
    else: # close
        await close_days_for_user(user_id, [day])
    
    await callback.answer(f"День {day}: {'відкрито' if action == 'open' else 'закрито'}")
    # Refresh the menu
    await developer_days_manage_menu(callback, state, user_id=user_id)

async def developer_user_profile_back(callback: CallbackQuery, state: FSMContext):
    try:
        user_id = int(callback.data.split(":")[1])
    except (ValueError, IndexError):
        await callback.answer("Помилка ID користувача.", show_alert=True)
        return

    await state.clear()
    await _rebuild_user_profile_view(callback, user_id, state)
    await callback.answer()

# --- Modify User ---

async def developer_user_modify_menu(callback: CallbackQuery, state: FSMContext):
    try:
        user_id = int(callback.data.split(":")[1])
    except (ValueError, IndexError):
        await callback.answer("Помилка ID користувача.", show_alert=True)
        return

    await state.clear()
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="Змінити керівника", callback_data=f"dev_change_manager:{user_id}")],
        [InlineKeyboardButton(text="Змінити посаду", callback_data=f"dev_change_role:{user_id}")],
        [InlineKeyboardButton(text="Змінити місто", callback_data=f"dev_change_city:{user_id}")],
        [InlineKeyboardButton(text="⬅️ Назад до профілю", callback_data=f"dev_user_profile_back:{user_id}")]
    ])
    await _edit_or_answer(callback.message, "✍️ <b>Що саме ви хочете змінити?</b>", reply_markup=kb)
    await callback.answer()

async def developer_change_manager_start(callback: CallbackQuery, state: FSMContext):
    try:
        user_id = int(callback.data.split(":")[1])
    except (ValueError, IndexError):
        await callback.answer("Помилка ID користувача.", show_alert=True)
        return

    await state.set_state(DeveloperStates.waiting_change_manager_id)
    await state.update_data(user_id_to_modify=user_id)
    await _edit_or_answer(callback.message, "Введіть новий ID керівника:", reply_markup=_cancel_keyboard(f"dev_user_modify:{user_id}"))
    await callback.answer()

async def developer_process_change_manager(message: Message, state: FSMContext):
    data = await state.get_data()
    user_id = data.get("user_id_to_modify")
    if not user_id:
        await message.answer("⚠️ Помилка стану: користувача не знайдено.")
        await state.clear()
        return

    try:
        new_manager_id = int(message.text.strip())
    except ValueError:
        await message.answer("❌ ID має бути числом. Спробуйте ще раз:")
        return

    # Check if manager exists
    manager = await get_manager_by_uid(new_manager_id)
    if not manager:
        await message.answer("⚠️ Керівника з таким ID не знайдено в базі керівників. Спробуйте ще раз:")
        return

    user_details = await get_user_details(user_id)
    if not user_details:
        await message.answer("⚠️ Користувача не знайдено.")
        await state.clear()
        return

    # Update manager_id directly without altering shop or logging added events
    async with aiosqlite.connect(DB_PATH) as db:
        await db.execute(
            "UPDATE users SET manager_id = ? WHERE user_id = ?",
            (new_manager_id, user_id)
        )
        await db.commit()

    await state.clear()
    await message.answer(f"✅ Керівника успішно змінено на <b>{manager.get('full_name', new_manager_id)}</b>!")
    
    try:
        report = await get_user_days_report(user_id)
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [
                InlineKeyboardButton(text="📅 Навчальні дні", callback_data=f"dev_days_manage:{user_id}"),
                InlineKeyboardButton(text="✍️ Змінити", callback_data=f"dev_user_modify:{user_id}")
            ],
            [
                InlineKeyboardButton(text="❌ Видалити", callback_data=f"dev_user_delete_confirm:{user_id}"),
                InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_users_menu")
            ]
        ])
        await message.answer(report, reply_markup=kb)
    except Exception as e:
        logger = get_logger()
        logger.error(f"Failed to display updated user profile: {e}")

async def developer_change_role_start(callback: CallbackQuery, state: FSMContext):
    try:
        user_id = int(callback.data.split(":")[1])
    except (ValueError, IndexError):
        await callback.answer("Помилка ID користувача.", show_alert=True)
        return

    positions = await get_all_positions()
    role_names = [p["name"] for p in positions] if positions else AVAILABLE_ROLES

    buttons = []
    for idx, role in enumerate(role_names):
        short_name = role[:25] + "..." if len(role) > 28 else role
        buttons.append([InlineKeyboardButton(text=short_name, callback_data=f"dev_set_role:{user_id}:{idx}")])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data=f"dev_user_modify:{user_id}")])
    await _edit_or_answer(callback.message, "Оберіть нову посаду:", reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons))
    await callback.answer()

async def developer_process_change_role(callback: CallbackQuery, state: FSMContext):
    try:
        parts = callback.data.split(":")
        user_id = int(parts[1])
        role_idx = int(parts[2])
    except (ValueError, IndexError):
        await callback.answer("Помилка даних.", show_alert=True)
        return

    positions = await get_all_positions()
    role_names = [p["name"] for p in positions] if positions else AVAILABLE_ROLES

    if not (0 <= role_idx < len(role_names)):
        await callback.answer("Невідома посада.", show_alert=True)
        return

    new_role = role_names[role_idx]
    user_details = await get_user_details(user_id)
    if not user_details:
        await callback.answer("Користувача не знайдено.", show_alert=True)
        return

    # Update role directly without erasing shop
    async with aiosqlite.connect(DB_PATH) as db:
        await db.execute(
            "UPDATE users SET role = ? WHERE user_id = ?",
            (new_role, user_id)
        )
        await db.commit()

    await callback.answer(f"Посаду змінено на {new_role}!", show_alert=True)
    await _rebuild_user_profile_view(callback, user_id, state)

async def developer_change_city_start(callback: CallbackQuery, state: FSMContext):
    try:
        user_id = int(callback.data.split(":")[1])
    except (ValueError, IndexError):
        await callback.answer("Помилка ID користувача.", show_alert=True)
        return

    cities_data = await get_all_cities()
    city_names = [c["name"] for c in cities_data] if cities_data else AVAILABLE_CITIES

    buttons = []
    for idx, city in enumerate(city_names):
        buttons.append([InlineKeyboardButton(text=city, callback_data=f"dev_set_city:{user_id}:{idx}")])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data=f"dev_user_modify:{user_id}")])
    await _edit_or_answer(callback.message, "Оберіть нове місто:", reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons))
    await callback.answer()

async def developer_process_change_city(callback: CallbackQuery, state: FSMContext):
    try:
        parts = callback.data.split(":")
        user_id = int(parts[1])
        city_idx = int(parts[2])
    except (ValueError, IndexError):
        await callback.answer("Помилка даних.", show_alert=True)
        return

    cities_data = await get_all_cities()
    city_names = [c["name"] for c in cities_data] if cities_data else AVAILABLE_CITIES

    if not (0 <= city_idx < len(city_names)):
        await callback.answer("Невідоме місто.", show_alert=True)
        return

    new_city = city_names[city_idx]
    user_details = await get_user_details(user_id)
    if not user_details:
        await callback.answer("Користувача не знайдено.", show_alert=True)
        return

    # Update city directly without erasing shop
    async with aiosqlite.connect(DB_PATH) as db:
        await db.execute(
            "UPDATE users SET city = ? WHERE user_id = ?",
            (new_city, user_id)
        )
        await db.commit()

    await callback.answer(f"Місто змінено на {new_city}!", show_alert=True)
    await _rebuild_user_profile_view(callback, user_id, state)

# --- Delete User ---

async def developer_user_delete_confirm(callback: CallbackQuery, state: FSMContext):
    try:
        user_id = int(callback.data.split(":")[1])
    except (ValueError, IndexError):
        await callback.answer("Помилка ID користувача.", show_alert=True)
        return

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(text="✅ Так, видалити", callback_data=f"dev_user_delete_perform:{user_id}"),
            InlineKeyboardButton(text="❌ Ні, назад", callback_data=f"dev_user_profile_back:{user_id}")
        ]
    ])
    await _edit_or_answer(callback.message, f"Ви впевнені, що хочете видалити користувача <code>{user_id}</code>?", reply_markup=kb)
    await callback.answer()

async def developer_user_delete_perform(callback: CallbackQuery, state: FSMContext):
    try:
        user_id = int(callback.data.split(":")[1])
    except (ValueError, IndexError):
        await callback.answer("Помилка ID користувача.", show_alert=True)
        return

    await delete_user(user_id)
    await callback.answer(f"Користувача {user_id} видалено.", show_alert=True)
    await _edit_or_answer(callback.message, f"✅ Користувача <code>{user_id}</code> видалено.", reply_markup=await _get_users_menu_kb(callback.from_user.id))


# ==================== XLSX Reports ====================

def _subtract_months(dt: datetime, months: int) -> datetime:
    """Повертає дату, зсунуту на `months` місяців назад."""
    year = dt.year
    month = dt.month - months
    while month <= 0:
        month += 12
        year -= 1
    day = min(dt.day, calendar.monthrange(year, month)[1])
    return dt.replace(year=year, month=month, day=day)


def _report_range_start(now: datetime, months: int) -> datetime:
    # 1 місяць = поточний місяць з 1-го числа
    # 2+ місяці = з 1-го числа місяця, що був N-1 місяців тому
    anchor = _subtract_months(now, max(months - 1, 0))
    return anchor.replace(day=1, hour=0, minute=0, second=0, microsecond=0)


def _period_label(months: int) -> str:
    mapping = {
        1: "Останній місяць",
        2: "Останні 2 місяці",
        3: "Останні 3 місяці",
        6: "Останні 6 місяців",
        12: "Останній рік",
    }
    return mapping.get(months, f"Останні {months} місяців")

async def developer_xlsx_menu(callback: CallbackQuery):
    """Меню формування XLSX звіту."""
    logger = get_logger()
    logger.debug(f"XLSX menu requested by user {callback.from_user.id}")
    
    try:
        await callback.answer()
    except Exception as e:
        logger.error(f"Failed to answer callback in developer_xlsx_menu: {e}", exc_info=True)
        
    if not await _ensure_developer(callback):
        logger.warning(f"User {callback.from_user.id} FAILED developer check for XLSX menu")
        return
    
    now = datetime.now(pytz.timezone(TIMEZONE))
    default_months = 1
    start_date = _report_range_start(now, default_months)
    text = (
        "📊 <b>Звітність у форматі XLSX</b>\n"
        "───────────────────\n"
        f"📅 Період за замовчуванням: <b>{_period_label(default_months)}</b>\n"
        f"   <b>{start_date.strftime('%d.%m.%Y')} — {now.strftime('%d.%m.%Y')}</b>\n\n"
        "Звіт містить наступні метрики по містах та магазинах:\n"
        "• Кількість нових стажерів\n"
        "• Кількість активних стажерів\n"
        "• Кількість відсіву (не завершили)\n"
        "• Частка запитів до керівників\n"
        "• Детальні списки ПІБ: додались / відсіялись / стали працівниками / відхилені\n"
        "───────────────────\n"
        "<i>💡 Оберіть період для формування звіту.</i>"
    )
    
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📥 За останній місяць", callback_data="dev_xlsx_send_period:1")],
        [InlineKeyboardButton(text="📥 За 2 місяці", callback_data="dev_xlsx_send_period:2")],
        [InlineKeyboardButton(text="📥 За 3 місяці", callback_data="dev_xlsx_send_period:3")],
        [InlineKeyboardButton(text="📥 За пів року", callback_data="dev_xlsx_send_period:6")],
        [InlineKeyboardButton(text="📥 За рік", callback_data="dev_xlsx_send_period:12")],
        [InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_main_analyt")],
    ])
    
    try:
        await _edit_or_answer(callback.message, text, reply_markup=kb)
    except Exception as e:
        logger.error(f"Failed to display XLSX menu: {e}", exc_info=True)
        try:
            await callback.message.answer(text, reply_markup=kb)
        except Exception as e2:
            logger.error(f"Fallback also failed: {e2}", exc_info=True)

async def developer_send_current_report(callback: CallbackQuery):
    """Генерує та надсилає поточний XLSX звіт."""
    await _developer_send_xlsx_report(callback, months=1)


async def developer_send_period_report(callback: CallbackQuery):
    """Генерує та надсилає XLSX звіт за обраний період."""
    try:
        months = int((callback.data or "").split(":")[1])
    except (IndexError, ValueError):
        months = 1
    await _developer_send_xlsx_report(callback, months=months)


async def _developer_send_xlsx_report(callback: CallbackQuery, months: int):
    """Спільна логіка генерації XLSX за обраний період."""
    logger = get_logger()
    logger.debug(f"XLSX report generation requested by user {callback.from_user.id}")
    
    if not await _ensure_developer(callback):
        return
    
    try:
        await callback.answer("⏳ Генерую звіт...", show_alert=False)
    except Exception as e:
        logger.error(f"Failed to answer callback in developer_send_current_report: {e}", exc_info=True)
    
    now = datetime.now(pytz.timezone(TIMEZONE))
    start_date = _report_range_start(now, months)
    
    try:
        logger.debug(f"Getting report data from {start_date} to {now}")
        data = await get_report_data(start_date, now)
        details = await get_report_details(start_date, now)
        has_details = any(details.get(k) for k in ("added", "dropped", "promoted", "rejected"))
        if not data and not has_details:
            await callback.message.answer("❌ За цей період ще немає даних для звіту.")
            logger.debug(f"No data for report for user {callback.from_user.id}")
            return
        
        logger.debug(f"Generating XLSX report with {len(data)} records")
        xlsx_file = generate_xlsx_report(data, start_date, now, details)
        
        filename = f"Bulka_Report_{months}m_{start_date.strftime('%Y%m%d')}_{now.strftime('%Y%m%d')}.xlsx"
        from aiogram.types import BufferedInputFile
        
        await callback.message.answer_document(
            document=BufferedInputFile(xlsx_file.getvalue(), filename=filename),
            caption=(
                f"📊 <b>{_period_label(months)}</b>\n"
                f"{start_date.strftime('%d.%m.%Y')} - {now.strftime('%d.%m.%Y')}"
            ),
        )
        logger.debug(f"Report sent to user {callback.from_user.id}")
    except Exception as e:
        logger.error(f"Помилка генерації або надсилання XLSX звіту: {e}", exc_info=True)
        try:
            await callback.message.answer("❌ Помилка при генерації звіту.")
        except Exception:
            pass

# ==================== Reminder History ====================

async def developer_reminder_history_menu(callback: CallbackQuery):
    """Displays the reminder history."""
    if not await _ensure_developer(callback):
        return
    
    # Показуємо 15 останніх записів, щоб не перевищити ліміт символів
    logs = await get_reminder_history(limit=15)
    
    if not logs:
        await _edit_or_answer(
            callback.message,
            "📭 Історія нагадувань порожня.",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_main_analyt")]
            ])
        )
        await callback.answer()
        return

    lines = [f"📜 <b>Історія нагадувань (останні {len(logs)})</b>", ""]
    
    for log in logs:
        sent_at = log['sent_at']
        intern_name = log['intern_name'] or f"ID {log['intern_id']}"
        source = log['source']
        sender_name = log['sender_name'] or (f"ID {log['sender_id']}" if log['sender_id'] else None)
        
        icon = "🤖"
        if source == "manager":
            icon = "👨‍🏫"
        elif source == "hr":
            icon = "👔"
            
        by_whom = ""
        if source == "auto":
            by_whom = "Автоматично"
        elif sender_name:
            by_whom = f"{sender_name}"
        else:
            by_whom = source.capitalize()
            
        # Екрануємо дані з БД
        e_intern_name = html.escape(intern_name)
        e_by_whom = html.escape(by_whom)
        
        lines.append(f"{icon} <b>{sent_at}</b> → {e_intern_name}")
        lines.append(f"   <i>Від: {e_by_whom}</i>")
        lines.append("───────────────")

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🔄 Оновити", callback_data="dev_reminder_history")],
        [InlineKeyboardButton(text="⬅️ Назад", callback_data="developer_menu")]
    ])
    
    await _edit_or_answer(callback.message, "\n".join(lines), reply_markup=kb)
    await callback.answer()


# ==================== Analytics ====================

async def developer_analytics_menu(callback: CallbackQuery):
    """Меню вибору аналітичних звітів."""
    if not await _ensure_developer(callback):
        return
    
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📅 Щоденний дайджест", callback_data="dev_analytics_digest")],
        [InlineKeyboardButton(text="📉 Воронка відсіву", callback_data="dev_analytics_funnel")],
        [InlineKeyboardButton(text="🎓 Навчальний процес", callback_data="dev_training_menu")],
        [InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_main_analyt")]
    ])
    
    await _edit_or_answer(
        callback.message,
        "📊 <b>Аналітика</b>\nОберіть тип звіту:",
        reply_markup=kb
    )
    await callback.answer()

async def developer_daily_digest(callback: CallbackQuery):
    """Показує статистику за останні 24 години."""
    if not await _ensure_developer(callback):
        return
        
    stats = await get_daily_stats()
    
    text = (
        "📅 <b>Щоденний дайджест (24 год)</b>\n\n"
        f"🆕 Нових стажерів: <b>{stats['new_users']}</b>\n"
        f"✅ Пройдено блоків: <b>{stats['completions']}</b>\n"
        f"🏃‍♂️ Активних користувачів: <b>{stats['active_users']}</b>\n"
    )
    
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_analytics_menu")]
    ])
    
    await _edit_or_answer(callback.message, text, reply_markup=kb)
    await callback.answer()

async def developer_dropout_report(callback: CallbackQuery):
    """Показує воронку відсіву користувачів."""
    if not await _ensure_developer(callback):
        return
        
    data = await get_dropout_funnel(active_days=3)
    
    # 1. Секція для АКТИВНИХ
    active_total = data['active_users']
    active_dist = data['active_distribution']
    
    lines = [
        "📉 <b>Воронка відсіву (АКТИВНІ до 3 дн)</b>",
        f"Активних стажерів: <b>{active_total}</b>",
        ""
    ]
    
    if active_total > 0:
        for day in range(1, DAYS_TOTAL + 1):
            count = active_dist.get(day, 0)
            percent = (count / active_total * 100)
            bar = "█" * int(percent / 5)
            lines.append(f"День {day}: <b>{count}</b> ({percent:.1f}%)")
            lines.append(f"<code>{bar}</code>")
    else:
        lines.append("<i>Немає активних стажерів за останні 3 дні.</i>")

    lines.append("\n" + "─" * 15 + "\n")

    # 2. Секція для ВСІХ
    total_count = data['total_users']
    total_dist = data['total_distribution']
    
    lines.append("📊 <b>Загальна статистика (стажери в процесі)</b>")
    lines.append(f"Всього у процесі: <b>{total_count}</b>\n")
    
    if total_count > 0:
        for day in range(1, DAYS_TOTAL + 1):
            count = total_dist.get(day, 0)
            percent = (count / total_count * 100)
            bar = "░" * int(percent / 5)
            lines.append(f"День {day}: <b>{count}</b> ({percent:.1f}%)")
            lines.append(f"<code>{bar}</code>")
    
    footer = "* Секція 'Активні' не враховує тих, хто не заходив у бот ≥ 3 дні."
    lines.append(f"\n<i>{html.escape(footer)}</i>")
    
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🔄 Оновити", callback_data="dev_analytics_funnel")],
        [InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_analytics_menu")]
    ])
    
    await _edit_or_answer(callback.message, "\n".join(lines), reply_markup=kb)
    await callback.answer()


def _range_days_from_token(token: str) -> Optional[int]:
    if token == "7":
        return 7
    if token == "30":
        return 30
    return None


def _safe_name(row: dict) -> str:
    return row.get("full_name") or (f"@{row.get('username')}" if row.get("username") else f"ID {row.get('user_id')}")

async def _resolve_training_added_city(city_token: str) -> Optional[str]:
    if city_token == "all":
        return None
    try:
        city_idx = int(city_token)
    except ValueError:
        return None
    cities = await get_training_added_cities()
    if 0 <= city_idx < len(cities):
        return cities[city_idx]
    return None


async def developer_training_menu(callback: CallbackQuery):
    if not await _ensure_developer(callback):
        return
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🆕 Додані стажери", callback_data="dev_training_added:all")],
        [InlineKeyboardButton(text="😴 Залишили навчання", callback_data="dev_training_left")],
        [InlineKeyboardButton(text="👷 Стали працівниками", callback_data="dev_training_promoted:all")],
        [InlineKeyboardButton(text="❌ Відхилені стажери", callback_data="dev_training_rejected:all")],
        [InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_analytics_menu")],
    ])
    await _edit_or_answer(callback.message, "🎓 <b>Навчальний процес</b>\nОберіть розділ:", reply_markup=kb)
    await callback.answer()


async def _render_training_added(
    callback: CallbackQuery,
    range_token: str,
    city: Optional[str] = None,
    city_idx: Optional[int] = None,
):
    range_days = _range_days_from_token(range_token)
    rows = await get_training_added_interns(range_days=range_days, city=city)
    filter_label = {"7": "останні 7 дн.", "30": "останні 30 дн.", "all": "весь час"}.get(range_token, "весь час")
    city_label = f" | Місто: <b>{html.escape(city)}</b>" if city else ""

    lines = [
        f"🆕 <b>Додані стажери</b> ({filter_label}){city_label}",
        f"Всього: <b>{len(rows)}</b>",
        "",
    ]
    for idx, row in enumerate(rows[:200], start=1):
        name = html.escape(_safe_name(row))
        city_v = html.escape(row.get("city") or "—")
        role_v = html.escape(row.get("role") or "—")
        lines.append(f"{idx}. {name} — {city_v} | {role_v} | {row.get('event_at')}")

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(text="7 дн", callback_data="dev_training_added:7"),
            InlineKeyboardButton(text="30 дн", callback_data="dev_training_added:30"),
            InlineKeyboardButton(text="Весь час", callback_data="dev_training_added:all"),
        ],
        [InlineKeyboardButton(text="🏙️ За містом", callback_data=f"dev_training_added_city_menu:{range_token}")],
        [
            InlineKeyboardButton(
                text="📥 Вивантажити",
                callback_data=f"dev_training_added_export:{range_token}:{city_idx if city_idx is not None else 'all'}",
            )
        ],
        [
            InlineKeyboardButton(
                text="🧹 Очистити",
                callback_data=f"dev_training_added_clear_ask:{range_token}:{city_idx if city_idx is not None else 'all'}",
            )
        ],
        [InlineKeyboardButton(text="⬅️ Навчальний процес", callback_data="dev_training_menu")],
    ])
    await _edit_or_answer(callback.message, "\n".join(lines), reply_markup=kb)


async def developer_training_added(callback: CallbackQuery):
    if not await _ensure_developer(callback):
        return
    parts = (callback.data or "").split(":")
    token = parts[1] if len(parts) > 1 else "all"
    await _render_training_added(callback, token)
    await callback.answer()


async def developer_training_added_city_menu(callback: CallbackQuery):
    if not await _ensure_developer(callback):
        return
    parts = (callback.data or "").split(":")
    token = parts[1] if len(parts) > 1 else "all"
    cities = await get_training_added_cities()
    if not cities:
        await callback.answer("У подіях немає міст для фільтра.", show_alert=True)
        return

    rows = [[InlineKeyboardButton(text=city, callback_data=f"dev_training_added_city:{token}:{idx}")] for idx, city in enumerate(cities)]
    rows.append([InlineKeyboardButton(text="⬅️ Назад", callback_data=f"dev_training_added:{token}")])
    await _edit_or_answer(
        callback.message,
        "🏙️ <b>Фільтр доданих стажерів за містом</b>\nОберіть місто:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=rows),
    )
    await callback.answer()


async def developer_training_added_city(callback: CallbackQuery):
    if not await _ensure_developer(callback):
        return
    parts = (callback.data or "").split(":")
    if len(parts) < 4:
        await callback.answer("Некоректний фільтр.", show_alert=True)
        return
    token = parts[2]
    try:
        city_idx = int(parts[3])
    except ValueError:
        await callback.answer("Некоректне місто.", show_alert=True)
        return

    cities = await get_training_added_cities()
    if city_idx < 0 or city_idx >= len(cities):
        await callback.answer("Місто не знайдено.", show_alert=True)
        return

    await _render_training_added(callback, token, city=cities[city_idx], city_idx=city_idx)
    await callback.answer()


async def developer_training_added_export(callback: CallbackQuery):
    if not await _ensure_developer(callback):
        return
    parts = (callback.data or "").split(":")
    token = parts[1] if len(parts) > 1 else "all"
    city_token = parts[2] if len(parts) > 2 else "all"

    city: Optional[str] = None
    if city_token != "all":
        try:
            city_idx = int(city_token)
            cities = await get_training_added_cities()
            if 0 <= city_idx < len(cities):
                city = cities[city_idx]
        except ValueError:
            city = None

    rows = await get_training_added_interns(range_days=_range_days_from_token(token), city=city)
    await _export_training_events_xlsx(callback, rows, "Додані_стажери.xlsx", "Додані стажери")
    await callback.answer()


async def developer_training_added_clear_ask(callback: CallbackQuery):
    if not await _ensure_developer(callback):
        return
    parts = (callback.data or "").split(":")
    token = parts[1] if len(parts) > 1 else "all"
    city_token = parts[2] if len(parts) > 2 else "all"
    city = await _resolve_training_added_city(city_token)
    filter_label = {"7": "останні 7 дн.", "30": "останні 30 дн.", "all": "весь час"}.get(token, "весь час")
    city_label = f"\n🏙️ Місто: <b>{html.escape(city)}</b>" if city else ""

    text = (
        "🧹 <b>Підтвердження очищення</b>\n\n"
        "Ви дійсно хочете очистити записи в розділі <b>Додані стажери</b> за цим фільтром?\n"
        f"📅 Період: <b>{filter_label}</b>{city_label}\n\n"
        "Після очищення ці записи не будуть відображатися у звіті."
    )
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(
                text="✅ Так, очистити",
                callback_data=f"dev_training_added_clear_confirm:{token}:{city_token}",
            ),
            InlineKeyboardButton(
                text="❌ Скасувати",
                callback_data=f"dev_training_added_clear_cancel:{token}:{city_token}",
            ),
        ]
    ])
    await _edit_or_answer(callback.message, text, reply_markup=kb)
    await callback.answer()


async def developer_training_added_clear_cancel(callback: CallbackQuery):
    if not await _ensure_developer(callback):
        return
    parts = (callback.data or "").split(":")
    token = parts[1] if len(parts) > 1 else "all"
    city_token = parts[2] if len(parts) > 2 else "all"
    city = await _resolve_training_added_city(city_token)
    city_idx = int(city_token) if city_token != "all" and city_token.isdigit() else None
    await _render_training_added(callback, token, city=city, city_idx=city_idx)
    await callback.answer("Скасовано")


async def developer_training_added_clear_confirm(callback: CallbackQuery):
    if not await _ensure_developer(callback):
        return
    parts = (callback.data or "").split(":")
    token = parts[1] if len(parts) > 1 else "all"
    city_token = parts[2] if len(parts) > 2 else "all"
    city = await _resolve_training_added_city(city_token)

    deleted_count = await clear_training_added_interns(
        range_days=_range_days_from_token(token),
        city=city,
    )
    city_idx = int(city_token) if city_token != "all" and city_token.isdigit() else None
    await _render_training_added(callback, token, city=city, city_idx=city_idx)
    await callback.answer(f"Очищено: {deleted_count}", show_alert=True)


async def developer_training_left(callback: CallbackQuery):
    if not await _ensure_developer(callback):
        return
    rows = await get_training_left_inactive(days=3)
    lines = [
        "😴 <b>Залишили навчання</b> (неактивні ≥ 3 днів)",
        f"Всього: <b>{len(rows)}</b>",
        "",
    ]
    for idx, row in enumerate(rows[:200], start=1):
        name = html.escape(_safe_name(row))
        shop = html.escape((row.get("shop") or "—").split(" ")[0])
        if row.get("deleted"):
            lines.append(
                f"{idx}. {name} — неактивний з <b>{row.get('inactive_since') or '—'}</b> | "
                f"магазин: <b>{shop}</b> | статус: <b>Видалено</b> ({row.get('deleted_at') or '—'})"
            )
        else:
            lines.append(
                f"{idx}. {name} — неактивний з <b>{row.get('inactive_since') or '—'}</b> | "
                f"магазин: <b>{shop}</b> | статус: <b>Неактивний</b>"
            )

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🔄 Оновити", callback_data="dev_training_left")],
        [InlineKeyboardButton(text="⬅️ Навчальний процес", callback_data="dev_training_menu")],
    ])
    await _edit_or_answer(callback.message, "\n".join(lines), reply_markup=kb)
    await callback.answer()


async def _render_training_promoted(callback: CallbackQuery, token: str, page: int = 0):
    rows = await get_training_promoted(range_days=_range_days_from_token(token))
    filter_label = {"7": "останні 7 дн.", "30": "останні 30 дн.", "all": "весь час"}.get(token, "весь час")
    items_per_page = 20
    pages_total = max((len(rows) + items_per_page - 1) // items_per_page, 1)
    curr_page = max(0, min(page, pages_total - 1))
    start = curr_page * items_per_page
    end = start + items_per_page
    page_rows = rows[start:end]

    lines = [
        f"👷 <b>Стали працівниками</b> ({filter_label})",
        f"Всього: <b>{len(rows)}</b> | Стор. <b>{curr_page + 1}/{pages_total}</b>",
        "",
    ]
    for idx, row in enumerate(page_rows, start=start + 1):
        name = html.escape(_safe_name(row))
        role_v = html.escape(row.get("role") or "—")
        lines.append(f"{idx}. {name} — {role_v} | {row.get('event_at')}")

    nav_row = []
    if curr_page > 0:
        nav_row.append(InlineKeyboardButton(text="⬅️ Попередня", callback_data=f"dev_training_promoted_page:{token}:{curr_page - 1}"))
    if curr_page < pages_total - 1:
        nav_row.append(InlineKeyboardButton(text="Наступна ➡️", callback_data=f"dev_training_promoted_page:{token}:{curr_page + 1}"))

    kb_rows = []
    if nav_row:
        kb_rows.append(nav_row)
    kb = InlineKeyboardMarkup(inline_keyboard=[
        *kb_rows,
        [
            InlineKeyboardButton(text="7 дн", callback_data="dev_training_promoted:7"),
            InlineKeyboardButton(text="30 дн", callback_data="dev_training_promoted:30"),
            InlineKeyboardButton(text="Весь час", callback_data="dev_training_promoted:all"),
        ],
        [InlineKeyboardButton(text="📥 Вивантажити", callback_data=f"dev_training_promoted_export:{token}")],
        [InlineKeyboardButton(text="⬅️ Навчальний процес", callback_data="dev_training_menu")],
    ])
    await _edit_or_answer(callback.message, "\n".join(lines), reply_markup=kb)


async def developer_training_promoted(callback: CallbackQuery):
    if not await _ensure_developer(callback):
        return
    parts = (callback.data or "").split(":")
    token = parts[1] if len(parts) > 1 else "all"
    await _render_training_promoted(callback, token)
    await callback.answer()

async def developer_training_promoted_page(callback: CallbackQuery):
    if not await _ensure_developer(callback):
        return
    parts = (callback.data or "").split(":")
    if len(parts) < 3:
        await callback.answer("Некоректна сторінка.", show_alert=True)
        return
    token = parts[1]
    try:
        page = int(parts[2])
    except ValueError:
        await callback.answer("Некоректна сторінка.", show_alert=True)
        return
    await _render_training_promoted(callback, token, page=page)
    await callback.answer()


async def _render_training_rejected(callback: CallbackQuery, token: str):
    rows = await get_training_rejected(range_days=_range_days_from_token(token))
    filter_label = {"7": "останні 7 дн.", "30": "останні 30 дн.", "all": "весь час"}.get(token, "весь час")
    lines = [
        f"❌ <b>Відхилені стажери</b> ({filter_label})",
        f"Всього: <b>{len(rows)}</b>",
        "",
    ]
    for idx, row in enumerate(rows[:200], start=1):
        name = html.escape(_safe_name(row))
        lines.append(f"{idx}. {name} — дата відхилення: <b>{row.get('event_at') or '—'}</b>")

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(text="7 дн", callback_data="dev_training_rejected:7"),
            InlineKeyboardButton(text="30 дн", callback_data="dev_training_rejected:30"),
            InlineKeyboardButton(text="Весь час", callback_data="dev_training_rejected:all"),
        ],
        [InlineKeyboardButton(text="📥 Вивантажити", callback_data=f"dev_training_rejected_export:{token}")],
        [InlineKeyboardButton(text="⬅️ Навчальний процес", callback_data="dev_training_menu")],
    ])
    await _edit_or_answer(callback.message, "\n".join(lines), reply_markup=kb)


async def developer_training_rejected(callback: CallbackQuery):
    if not await _ensure_developer(callback):
        return
    parts = (callback.data or "").split(":")
    token = parts[1] if len(parts) > 1 else "all"
    await _render_training_rejected(callback, token)
    await callback.answer()


async def _export_training_events_xlsx(callback: CallbackQuery, rows: list[dict], filename: str, title: str):
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from aiogram.types import BufferedInputFile

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title

    headers = ["Дата", "ПІБ", "Username", "Місто", "Посада", "Керівник ID", "Telegram ID"]
    for idx, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=idx, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill(fill_type="solid", fgColor="305496")
        c.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(rows, start=2):
        ws.cell(row=row_idx, column=1, value=row.get("event_at"))
        ws.cell(row=row_idx, column=2, value=row.get("full_name") or "—")
        ws.cell(row=row_idx, column=3, value=row.get("username") or "—")
        ws.cell(row=row_idx, column=4, value=row.get("city") or "—")
        ws.cell(row=row_idx, column=5, value=row.get("role") or "—")
        ws.cell(row=row_idx, column=6, value=row.get("manager_id") or "—")
        ws.cell(row=row_idx, column=7, value=row.get("user_id") or "—")

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 35
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 15
    ws.column_dimensions["E"].width = 25
    ws.column_dimensions["F"].width = 14
    ws.column_dimensions["G"].width = 14

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    await callback.message.answer_document(
        document=BufferedInputFile(buf.getvalue(), filename=filename),
        caption=f"📥 <b>Вивантаження</b> — {len(rows)} записів",
        parse_mode="HTML",
    )


async def developer_training_promoted_export(callback: CallbackQuery):
    if not await _ensure_developer(callback):
        return
    parts = (callback.data or "").split(":")
    token = parts[1] if len(parts) > 1 else "all"
    rows = await get_training_promoted(range_days=_range_days_from_token(token))
    await _export_training_events_xlsx(callback, rows, "Стали_працівниками.xlsx", "Стали працівниками")
    await callback.answer("Файл готовий ✅")


async def developer_training_rejected_export(callback: CallbackQuery):
    if not await _ensure_developer(callback):
        return
    parts = (callback.data or "").split(":")
    token = parts[1] if len(parts) > 1 else "all"
    rows = await get_training_rejected(range_days=_range_days_from_token(token))
    await _export_training_events_xlsx(callback, rows, "Відхилені_стажери.xlsx", "Відхилені стажери")
    await callback.answer("Файл готовий ✅")



# ==============================================================================
# УПРАВЛІННЯ ПОСАДАМИ
# ==============================================================================

def _positions_keyboard(positions: list, page: int = 1) -> InlineKeyboardMarkup:
    items_per_page = 6
    total_pages = (len(positions) - 1) // items_per_page + 1
    start_idx = (page - 1) * items_per_page
    end_idx = start_idx + items_per_page
    current_positions = positions[start_idx:end_idx]

    buttons = []
    for pos in current_positions:
        buttons.append([InlineKeyboardButton(text=f"👔 {pos['name']}", callback_data=f"dev_pos_view:{pos['id']}")])
    
    # Пагінація
    nav_row = []
    if page > 1:
        nav_row.append(InlineKeyboardButton(text="⬅️", callback_data=f"dev_positions_page:{page-1}"))
    nav_row.append(InlineKeyboardButton(text=f"📄 {page}/{total_pages or 1}", callback_data="ignore"))
    if page < total_pages:
        nav_row.append(InlineKeyboardButton(text="➡️", callback_data=f"dev_positions_page:{page+1}"))
    if nav_row:
        buttons.append(nav_row)

    buttons.append([InlineKeyboardButton(text="➕ Додати посаду", callback_data="dev_pos_add")])
    buttons.append([InlineKeyboardButton(text="🔙 Назад", callback_data="dev_main_other")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)


async def dev_positions_menu(callback: CallbackQuery):
    positions = await get_all_positions()
    await _edit_or_answer(
        callback.message,
        "👔 <b>Управління посадами</b>\nОберіть посаду для редагування або додайте нову:",
        reply_markup=_positions_keyboard(positions, 1)
    )
    await callback.answer()


async def dev_positions_page(callback: CallbackQuery):
    page = int(callback.data.split(":")[1])
    positions = await get_all_positions()
    await _edit_or_answer(
        callback.message,
        "👔 <b>Управління посадами</b>\nОберіть посаду для редагування або додайте нову:",
        reply_markup=_positions_keyboard(positions, page)
    )
    await callback.answer()


async def dev_pos_add_start(callback: CallbackQuery, state: FSMContext):
    await _edit_or_answer(
        callback.message,
        "Введіть назву нової посади:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="Скасувати", callback_data="dev_positions_menu")]])
    )
    await state.set_state(DeveloperStates.waiting_add_position_name)
    await callback.answer()


async def dev_pos_add_name(message: Message, state: FSMContext):
    name = message.text.strip()
    await state.update_data(pos_name=name)
    await message.answer(
        f"Посада: <b>{name}</b>\nВведіть кількість днів навчання (за замовчуванням 5):",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="Залишити 5", callback_data="dev_pos_add_days:5")],
            [InlineKeyboardButton(text="Скасувати", callback_data="dev_positions_menu")]
        ])
    )
    await state.set_state(DeveloperStates.waiting_add_position_days)


async def _dev_pos_prompt_type(message_or_callback, state: FSMContext, pos_name: str, days: int):
    await state.update_data(pos_days=days)
    buttons = [
        [InlineKeyboardButton(text="🏪 Торговий зал (ТЗ)", callback_data="dev_pos_add_type:ТЗ")],
        [InlineKeyboardButton(text="🍞 Власне виробництво (ВВ)", callback_data="dev_pos_add_type:ВВ")],
        [InlineKeyboardButton(text="❌ Скасувати", callback_data="dev_positions_menu")]
    ]
    text = (
        f"Посада: <b>{pos_name}</b> ({days} днів)\n\n"
        f"Оберіть напрямок посади:\n"
        f"• <b>ТЗ (Торговий зал)</b> — касири, продавці тощо\n"
        f"• <b>ВВ (Власне виробництво)</b> — пекарі, кухарі, піцайоло тощо"
    )
    if isinstance(message_or_callback, CallbackQuery):
        await _edit_or_answer(message_or_callback.message, text, reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons))
    else:
        await message_or_callback.answer(text, reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons))
    await state.set_state(DeveloperStates.waiting_add_position_type)


async def dev_pos_add_days_text(message: Message, state: FSMContext):
    data = await state.get_data()
    pos_name = data.get('pos_name')
    if not pos_name:
        await state.clear()
        return
        
    try:
        days = int(message.text.strip())
        if days < 1:
            await message.answer("Кількість днів повинна бути не менше 1.")
            return
        await _dev_pos_prompt_type(message, state, pos_name, days)
    except ValueError:
        await message.answer("Будь ласка, введіть числове значення кількості днів.")


async def dev_pos_add_days_callback(callback: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    pos_name = data.get('pos_name')
    if not pos_name:
        await state.clear()
        return
        
    days = int(callback.data.split(":")[1])
    await _dev_pos_prompt_type(callback, state, pos_name, days)
    await callback.answer()


async def dev_pos_add_type_callback(callback: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    pos_name = data.get('pos_name')
    days = data.get('pos_days', 5)
    if not pos_name:
        await state.clear()
        await callback.answer("Дані сесії застаріли.", show_alert=True)
        return

    t_type = callback.data.split(":")[1]  # 'ТЗ' або 'ВВ'
    type_label = "🏪 Торговий зал (ТЗ)" if t_type == "ТЗ" else "🍞 Власне виробництво (ВВ)"
    
    success = await add_position(pos_name, days, territorial_type=t_type)
    if success:
        await _edit_or_answer(
            callback.message,
            f"✅ Посаду <b>{pos_name}</b> ({days} днів, {type_label}) успішно додано!"
        )
    else:
        await _edit_or_answer(
            callback.message,
            f"❌ Помилка: посада <b>{pos_name}</b> вже існує або виникла інша помилка."
        )
    await state.clear()
    
    positions = await get_all_positions()
    await callback.message.answer(
        "👔 <b>Управління посадами</b>",
        reply_markup=_positions_keyboard(positions, 1)
    )
    await callback.answer()


async def dev_pos_view(callback: CallbackQuery):
    pos_id = int(callback.data.split(":")[1])
    pos = await get_position_by_id(pos_id)
    if not pos:
        await callback.answer("Посаду не знайдено!", show_alert=True)
        return
        
    stats = await get_position_stats(pos['name'])
    t_type = pos.get('territorial_type', 'ТЗ')
    type_label = "🏪 Торговий зал (ТЗ)" if t_type == "ТЗ" else "🍞 Власне виробництво (ВВ)"
    
    text = (
        f"👔 <b>Посада:</b> {pos['name']}\n"
        f"🏷 <b>Тип:</b> {type_label}\n"
        f"📅 <b>Днів навчання:</b> {pos['days_count']}\n\n"
        f"📊 <b>Статистика:</b>\n"
        f"• Працівників: {stats['workers_count']}\n"
        f"• Стажерів: {stats['interns_count']}\n"
        f"• Всього: {stats['total']}"
    )
    
    other_type = "ВВ" if t_type == "ТЗ" else "ТЗ"
    other_label = "🍞 Змінити на ВВ" if t_type == "ТЗ" else "🏪 Змінити на ТЗ"
    
    buttons = [
        [InlineKeyboardButton(text="✏️ Змінити назву", callback_data=f"dev_pos_edit_name:{pos_id}")],
        [InlineKeyboardButton(text="✏️ Змінити кількість днів", callback_data=f"dev_pos_edit_days:{pos_id}")],
        [InlineKeyboardButton(text=other_label, callback_data=f"dev_pos_toggle_type:{pos_id}:{other_type}")],
        [InlineKeyboardButton(text="🔙 До списку посад", callback_data="dev_positions_menu")]
    ]
    
    await _edit_or_answer(callback.message, text, reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons))
    await callback.answer()


async def dev_pos_toggle_type(callback: CallbackQuery):
    """Перемикає тип посади між ТЗ та ВВ."""
    parts = callback.data.split(":")
    pos_id = int(parts[1])
    new_type = parts[2]  # 'ТЗ' або 'ВВ'
    
    success = await update_position_type(pos_id, new_type)
    if not success:
        await callback.answer("❌ Помилка зміни типу", show_alert=True)
        return
    
    pos = await get_position_by_id(pos_id)
    stats = await get_position_stats(pos['name'])
    t_type = pos.get('territorial_type', 'ТЗ')
    type_label = "🏪 Торговий зал (ТЗ)" if t_type == "ТЗ" else "🍞 Власне виробництво (ВВ)"
    
    text = (
        f"👔 <b>Посада:</b> {pos['name']}\n"
        f"🏷 <b>Тип:</b> {type_label}\n"
        f"📅 <b>Днів навчання:</b> {pos['days_count']}\n\n"
        f"📊 <b>Статистика:</b>\n"
        f"• Працівників: {stats['workers_count']}\n"
        f"• Стажерів: {stats['interns_count']}\n"
        f"• Всього: {stats['total']}"
    )
    
    other_type = "ВВ" if t_type == "ТЗ" else "ТЗ"
    other_label = "🍞 Змінити на ВВ" if t_type == "ТЗ" else "🏪 Змінити на ТЗ"
    
    buttons = [
        [InlineKeyboardButton(text="✏️ Змінити назву", callback_data=f"dev_pos_edit_name:{pos_id}")],
        [InlineKeyboardButton(text="✏️ Змінити кількість днів", callback_data=f"dev_pos_edit_days:{pos_id}")],
        [InlineKeyboardButton(text=other_label, callback_data=f"dev_pos_toggle_type:{pos_id}:{other_type}")],
        [InlineKeyboardButton(text="🔙 До списку посад", callback_data="dev_positions_menu")]
    ]
    
    await _edit_or_answer(callback.message, text, reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons))
    await callback.answer(f"✅ Тип змінено на {t_type}")


async def dev_pos_edit_name_start(callback: CallbackQuery, state: FSMContext):
    pos_id = int(callback.data.split(":")[1])
    pos = await get_position_by_id(pos_id)
    
    await _edit_or_answer(
        callback.message,
        f"Введіть нову назву для посади <b>{pos['name']}</b>:\n\n"
        f"⚠️ <i>Увага: зміна назви автоматично оновить всіх користувачів, матеріали та історію!</i>",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="Скасувати", callback_data=f"dev_pos_view:{pos_id}")]])
    )
    await state.update_data(edit_pos_id=pos_id)
    await state.set_state(DeveloperStates.waiting_edit_position_name)
    await callback.answer()


async def dev_pos_edit_name_process(message: Message, state: FSMContext):
    data = await state.get_data()
    pos_id = data.get('edit_pos_id')
    new_name = message.text.strip()
    
    success = await update_position_name(pos_id, new_name)
    if success:
        await message.answer(f"✅ Назву посади змінено на <b>{new_name}</b>, зв'язані сутності оновлено.")
    else:
        await message.answer("❌ Помилка: можливо посада з такою назвою вже існує.")
        
    await state.clear()
    pos = await get_position_by_id(pos_id)
    
    # Повертаємось до перегляду
    stats = await get_position_stats(pos['name'])
    text = (
        f"👔 <b>Посада:</b> {pos['name']}\n"
        f"📅 <b>Днів навчання:</b> {pos['days_count']}\n\n"
        f"📊 <b>Статистика:</b>\n"
        f"• Працівників: {stats['workers_count']}\n"
        f"• Стажерів: {stats['interns_count']}\n"
        f"• Всього: {stats['total']}"
    )
    buttons = [
        [InlineKeyboardButton(text="✏️ Змінити назву", callback_data=f"dev_pos_edit_name:{pos_id}")],
        [InlineKeyboardButton(text="✏️ Змінити кількість днів", callback_data=f"dev_pos_edit_days:{pos_id}")],
        [InlineKeyboardButton(text="🔙 До списку посад", callback_data="dev_positions_menu")]
    ]
    await message.answer(text, reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons))


async def dev_pos_edit_days_start(callback: CallbackQuery, state: FSMContext):
    pos_id = int(callback.data.split(":")[1])
    pos = await get_position_by_id(pos_id)
    
    await _edit_or_answer(
        callback.message,
        f"Введіть нову кількість днів навчання для посади <b>{pos['name']}</b> (зараз {pos['days_count']}):",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="Скасувати", callback_data=f"dev_pos_view:{pos_id}")]])
    )
    await state.update_data(edit_pos_id=pos_id)
    await state.set_state(DeveloperStates.waiting_edit_position_days)
    await callback.answer()


async def dev_pos_edit_days_process(message: Message, state: FSMContext):
    data = await state.get_data()
    pos_id = data.get('edit_pos_id')
    
    try:
        new_days = int(message.text.strip())
        success = await update_position_days(pos_id, new_days)
        if success:
            await message.answer(f"✅ Кількість днів змінено на {new_days}.")
        else:
            await message.answer("❌ Помилка під час оновлення.")
    except ValueError:
        await message.answer("Будь ласка, введіть число.")
        return
        
    await state.clear()
    pos = await get_position_by_id(pos_id)
    
    # Повертаємось до перегляду
    stats = await get_position_stats(pos['name'])
    text = (
        f"👔 <b>Посада:</b> {pos['name']}\n"
        f"📅 <b>Днів навчання:</b> {pos['days_count']}\n\n"
        f"📊 <b>Статистика:</b>\n"
        f"• Працівників: {stats['workers_count']}\n"
        f"• Стажерів: {stats['interns_count']}\n"
        f"• Всього: {stats['total']}"
    )
    buttons = [
        [InlineKeyboardButton(text="✏️ Змінити назву", callback_data=f"dev_pos_edit_name:{pos_id}")],
        [InlineKeyboardButton(text="✏️ Змінити кількість днів", callback_data=f"dev_pos_edit_days:{pos_id}")],
        [InlineKeyboardButton(text="🔙 До списку посад", callback_data="dev_positions_menu")]
    ]
    await message.answer(text, reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons))


# ==============================================================================
# УПРАВЛІННЯ МІСТАМИ
# ==============================================================================

def _cities_keyboard(cities: list, page: int = 1) -> InlineKeyboardMarkup:
    items_per_page = 10
    total_pages = (len(cities) - 1) // items_per_page + 1
    start_idx = (page - 1) * items_per_page
    end_idx = start_idx + items_per_page
    current_cities = cities[start_idx:end_idx]

    buttons = []
    for city in current_cities:
        buttons.append([InlineKeyboardButton(text=f"🏙 {city['name']}", callback_data=f"dev_city_view:{city['id']}")])
    
    # Пагінація
    nav_row = []
    if page > 1:
        nav_row.append(InlineKeyboardButton(text="⬅️", callback_data=f"dev_cities_page:{page-1}"))
    nav_row.append(InlineKeyboardButton(text=f"📄 {page}/{total_pages or 1}", callback_data="ignore"))
    if page < total_pages:
        nav_row.append(InlineKeyboardButton(text="➡️", callback_data=f"dev_cities_page:{page+1}"))
    if nav_row:
        buttons.append(nav_row)

    buttons.append([InlineKeyboardButton(text="➕ Додати місто", callback_data="dev_city_add")])
    buttons.append([InlineKeyboardButton(text="🔙 Назад", callback_data="dev_main_other")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)


async def dev_cities_menu(callback: CallbackQuery):
    cities = await get_all_cities()
    await _edit_or_answer(
        callback.message,
        "🏙 <b>Управління містами</b>\nОберіть місто для редагування або додайте нове:",
        reply_markup=_cities_keyboard(cities, 1)
    )
    await callback.answer()


async def dev_cities_page(callback: CallbackQuery):
    page = int(callback.data.split(":")[1])
    cities = await get_all_cities()
    await _edit_or_answer(
        callback.message,
        "🏙 <b>Управління містами</b>\nОберіть місто для редагування або додайте нове:",
        reply_markup=_cities_keyboard(cities, page)
    )
    await callback.answer()


async def dev_city_add_start(callback: CallbackQuery, state: FSMContext):
    await _edit_or_answer(
        callback.message,
        "Введіть назву нового міста:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="Скасувати", callback_data="dev_cities_menu")]])
    )
    await state.set_state(DeveloperStates.waiting_add_city_name)
    await callback.answer()


async def dev_city_add_name(message: Message, state: FSMContext):
    name = message.text.strip()
    success = await add_city(name)
    if success:
        await message.answer(f"✅ Місто <b>{name}</b> успішно додано!")
    else:
        await message.answer(f"❌ Помилка: місто <b>{name}</b> вже існує.")
    await state.clear()
    
    cities = await get_all_cities()
    await message.answer(
        "🏙 <b>Управління містами</b>",
        reply_markup=_cities_keyboard(cities, 1)
    )


async def dev_city_view(callback: CallbackQuery):
    city_id = int(callback.data.split(":")[1])
    city = await get_city_by_id(city_id)
    if not city:
        await callback.answer("Місто не знайдено!", show_alert=True)
        return
        
    text = f"🏙 <b>Місто:</b> {city['name']}"
    
    buttons = [
        [InlineKeyboardButton(text="✏️ Змінити назву", callback_data=f"dev_city_edit_name:{city_id}")],
        [InlineKeyboardButton(text="🗑 Видалити місто", callback_data=f"dev_city_delete:{city_id}")],
        [InlineKeyboardButton(text="🔙 До списку міст", callback_data="dev_cities_menu")]
    ]
    
    await _edit_or_answer(callback.message, text, reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons))
    await callback.answer()


async def dev_city_edit_name_start(callback: CallbackQuery, state: FSMContext):
    city_id = int(callback.data.split(":")[1])
    city = await get_city_by_id(city_id)
    
    await _edit_or_answer(
        callback.message,
        f"Введіть нову назву для міста <b>{city['name']}</b>:\n\n"
        f"⚠️ <i>Увага: зміна назви автоматично оновить всіх користувачів!</i>",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="Скасувати", callback_data=f"dev_city_view:{city_id}")]])
    )
    await state.update_data(edit_city_id=city_id)
    await state.set_state(DeveloperStates.waiting_edit_city_name)
    await callback.answer()


async def dev_city_edit_name_process(message: Message, state: FSMContext):
    data = await state.get_data()
    city_id = data.get('edit_city_id')
    new_name = message.text.strip()
    
    success = await update_city_name(city_id, new_name)
    if success:
        await message.answer(f"✅ Назву міста змінено на <b>{new_name}</b>, зв'язані сутності оновлено.")
    else:
        await message.answer("❌ Помилка під час оновлення назви міста.")
        
    await state.clear()
    
    city = await get_city_by_id(city_id)
    text = f"🏙 <b>Місто:</b> {city['name']}"
    buttons = [
        [InlineKeyboardButton(text="✏️ Змінити назву", callback_data=f"dev_city_edit_name:{city_id}")],
        [InlineKeyboardButton(text="🗑 Видалити місто", callback_data=f"dev_city_delete:{city_id}")],
        [InlineKeyboardButton(text="🔙 До списку міст", callback_data="dev_cities_menu")]
    ]
    await message.answer(text, reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons))


async def dev_city_delete(callback: CallbackQuery):
    city_id = int(callback.data.split(":")[1])
    city = await get_city_by_id(city_id)
    
    if not city:
        await callback.answer("Місто не знайдено!", show_alert=True)
        return
        
    await _edit_or_answer(
        callback.message,
        f"❗️ Ви впевнені, що хочете видалити місто <b>{city['name']}</b>?\n\n"
        f"<i>Користувачі з цим містом збережуть його назву, але воно зникне зі списку вибору.</i>",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [
                InlineKeyboardButton(text="❌ Ні, скасувати", callback_data=f"dev_city_view:{city_id}"),
                InlineKeyboardButton(text="✅ Так, видалити", callback_data=f"dev_city_delete_confirm:{city_id}")
            ]
        ])
    )
    await callback.answer()


async def dev_city_delete_confirm(callback: CallbackQuery):
    city_id = int(callback.data.split(":")[1])
    success = await delete_city(city_id)
    
    if success:
        await callback.answer("Місто видалено!", show_alert=True)
    else:
        await callback.answer("Помилка при видаленні!", show_alert=True)
        
    cities = await get_all_cities()
    await _edit_or_answer(
        callback.message,
        "🏙 <b>Управління містами</b>",
        reply_markup=_cities_keyboard(cities, 1)
    )



# =========================================================
# ТЕРИТОРІАЛИ (Territorials)
# =========================================================
from database.managers import get_all_territorials, update_manager_responsible, reassign_city_managers_to_territorial, get_territorials_by_city

TERRITORIAL_TYPES = {"ТЗ": "🏪 Торговий зал", "ВВ": "🍞 Власне виробництво"}

async def _build_territorials_team_view(bot: Optional[Bot] = None) -> tuple[str, InlineKeyboardMarkup]:
    """Builds the view for the Territorials management panel with inline cards and Telegram sync."""
    territorials = await get_all_territorials()
    
    city_abbr = {
        "Хмельницький": "ХМ",
        "Камʼянець-Подільський": "КП"
    }
    
    lines = [
        f"🗺 <b>Команда Територіалів (всього: {len(territorials)})</b>",
        "Натисніть на територіала для перегляду картки:",
        ""
    ]
    buttons = []
    if not territorials:
        lines.append("  У системі поки немає територіалів.")
    else:
        for t in territorials:
            name, username = await _format_identity(
                t["uid"], t.get("full_name"), t.get("username"), bot=bot
            )
            city = t.get("city", "")
            abbr = city_abbr.get(city, city or "??")
            t_type = t.get("territorial_type") or "ТЗ"
            
            btn_text = f"🗺 {name} | {username} | {abbr} · {t_type}"
            buttons.append([InlineKeyboardButton(
                text=btn_text,
                callback_data=f"dev_territorial_view:{t['uid']}"
            )])
            
    buttons.append([InlineKeyboardButton(text="➕ Додати територіала", callback_data="dev_territorial_add")])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_main_team")])
    
    return "\n".join(lines), InlineKeyboardMarkup(inline_keyboard=buttons)


async def developer_territorials_menu(callback: CallbackQuery):
    if not await _ensure_developer(callback):
        return
    text, kb = await _build_territorials_team_view(bot=callback.bot)
    await _show_admin_photo_menu(
        callback,
        "admin_spus.jpg",
        text,
        kb,
    )


async def developer_list_territorials(callback: CallbackQuery):
    await developer_territorials_menu(callback)


async def developer_territorial_view(callback: CallbackQuery):
    if not await _ensure_developer(callback):
        return
    t_uid = int(callback.data.split(":")[1])
    manager = await get_manager_by_uid(t_uid)
    if not manager:
        await callback.answer("Територіала не знайдено!", show_alert=True)
        return

    bot = callback.bot
    t_type = manager.get('territorial_type') or 'ТЗ'
    type_label = TERRITORIAL_TYPES.get(t_type, t_type)
    city = manager.get('city', '')

    name, username = await _format_identity(t_uid, manager.get("full_name"), manager.get("username"), bot=bot)

    # Підлеглі керівники
    subordinate_managers = await get_managers_by_responsible(t_uid)
    managers_count = len(subordinate_managers)

    # Збираємо унікальні магазини серед підлеглих керівників
    shops_set = set()
    for m in subordinate_managers:
        m_shops = m.get('shops') or []
        if isinstance(m_shops, str):
            try:
                m_shops = json.loads(m_shops)
            except Exception:
                m_shops = []
        for s in m_shops:
            if s:
                shops_set.add(s)
    shops_count = len(shops_set)

    # Рахуємо стажерів та працівників по місту з урахуванням типу посади (ВВ або ТЗ)
    from database.users import get_users_by_city
    from bot.services.positions import get_all_positions
    
    positions = await get_all_positions()
    pos_type_map = {p["name"]: p.get("territorial_type", "ТЗ") for p in positions}
    
    city_users = await get_users_by_city(city) if city else []
    workers_count = 0
    trainees_count = 0
    for u in city_users:
        u_pos = u.get("role") or u.get("position") or ""
        u_type = pos_type_map.get(u_pos, "ТЗ")
        if u_type == t_type:
            if u.get("status") == "Працівник":
                workers_count += 1
            else:
                trainees_count += 1

    total_subordinates = managers_count + workers_count + trainees_count

    text = (
        f"🗺 <b>Територіал:</b> {html.escape(name)}\n"
        f"👤 <b>Username:</b> {html.escape(username)}\n"
        f"🆔 <b>Telegram ID:</b> <code>{manager.get('uid')}</code>\n"
        f"🏙 <b>Місто:</b> {html.escape(city or 'не вказано')}\n"
        f"🏷 <b>Тип:</b> {type_label} ({t_type})\n\n"
        f"📊 <b>Статистика по напрямку ({t_type} — {html.escape(city)}):</b>\n"
        f"• 👔 Підлеглих керівників: {managers_count}\n"
        f"• 🏪 Закріплених магазинів: {shops_count}\n"
        f"• 🎓 Стажерів ({t_type}): {trainees_count}\n"
        f"• 👷 Працівників ({t_type}): {workers_count}\n"
        f"• 👥 Всього людей у структурі: {total_subordinates}"
    )

    buttons = [
        [InlineKeyboardButton(text="👔 Керівники Територіала", callback_data=f"dev_territorial_managers:{t_uid}:0")],
        [InlineKeyboardButton(text="🗑 Видалити", callback_data=f"dev_territorial_del:{t_uid}")],
        [InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_territorials_menu")]
    ]

    photo_input = await get_user_avatar_input(bot, t_uid)
    await _send_or_edit_card_photo(
        callback,
        photo_input,
        text,
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )

async def developer_territorial_managers(callback: CallbackQuery):
    """Список керівників конкретного Територіала з пагінацією."""
    if not await _ensure_developer(callback):
        return
    parts = callback.data.split(":")
    t_uid = int(parts[1])
    page = int(parts[2]) if len(parts) > 2 else 0

    manager_info = await get_manager_by_uid(t_uid)
    t_name = manager_info.get('full_name', f'UID {t_uid}') if manager_info else f'UID {t_uid}'

    subordinates = await get_managers_by_responsible(t_uid)

    if not subordinates:
        await _edit_or_answer(
            callback.message,
            f"👔 <b>Керівники Територіала {t_name}</b>\n\nПоки що немає жодного керівника.",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="⬅️ Назад", callback_data=f"dev_territorial_view:{t_uid}")]
            ])
        )
        await callback.answer()
        return

    per_page = 10
    total = len(subordinates)
    pages = (total + per_page - 1) // per_page
    page = max(0, min(page, pages - 1))
    chunk = subordinates[page * per_page:(page + 1) * per_page]

    lines = []
    for m in chunk:
        name = m.get('full_name', 'Без імені')
        shops = m.get('shops') or []
        if isinstance(shops, str):
            import json
            try:
                shops = json.loads(shops)
            except Exception:
                shops = []
        shop_str = ', '.join(shops) if shops else '—'
        lines.append(f"• <b>{name}</b> | {shop_str}")

    text = (
        f"👔 <b>Керівники Територіала {t_name}</b> ({page+1}/{pages}):\n\n"
        + "\n".join(lines)
    )

    nav_buttons = []
    if page > 0:
        nav_buttons.append(InlineKeyboardButton(text="◀️", callback_data=f"dev_territorial_managers:{t_uid}:{page-1}"))
    if page < pages - 1:
        nav_buttons.append(InlineKeyboardButton(text="▶️", callback_data=f"dev_territorial_managers:{t_uid}:{page+1}"))

    buttons = []
    if nav_buttons:
        buttons.append(nav_buttons)
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data=f"dev_territorial_view:{t_uid}")])

    await _edit_or_answer(
        callback.message,
        text,
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    await callback.answer()

async def developer_territorial_del(callback: CallbackQuery):
    if not await _ensure_developer(callback):
        return
    t_uid = int(callback.data.split(":")[1])
    manager = await get_manager_by_uid(t_uid)
    if not manager:
        await callback.answer("Територіала не знайдено!", show_alert=True)
        return

    name = manager.get('full_name', 'Без імені')
    city = manager.get('city', '—')
    t_type = manager.get('territorial_type') or '—'

    # Підтвердження видалення
    buttons = [
        [InlineKeyboardButton(text="✅ Так, видалити", callback_data=f"dev_territorial_del_confirm:{t_uid}")],
        [InlineKeyboardButton(text="❌ Скасувати", callback_data=f"dev_territorial_view:{t_uid}")]
    ]
    await _edit_or_answer(
        callback.message,
        f"⚠️ Видалити Територіала <b>{name}</b> ({city} · {t_type})?\n\n"
        f"Всі підлеглі керівники повернуться до Адміністратора (без сповіщень).",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    await callback.answer()

async def developer_territorial_del_confirm(callback: CallbackQuery):
    if not await _ensure_developer(callback):
        return
    t_uid = int(callback.data.split(":")[1])
    await delete_manager_by_uid(t_uid)
    await callback.answer("Територіала видалено! Підлеглих передано Адміністратору.", show_alert=True)
    await developer_list_territorials(callback)

# ADDING TERRITORIAL — FSM: uid → city → type → transfer
async def developer_add_territorial_start(callback: CallbackQuery, state: FSMContext):
    if not await _ensure_developer(callback):
        return
    await _edit_or_answer(
        callback.message,
        "Введіть Telegram ID (UID) нового територіала:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="Скасувати", callback_data="dev_territorials_menu")]])
    )
    await state.set_state(DeveloperStates.waiting_add_territorial_uid)
    await callback.answer()

async def developer_add_territorial_uid(message: Message, state: FSMContext):
    try:
        uid = int(message.text.strip())
    except ValueError:
        await message.answer("Будь ласка, введіть коректний числовий UID.")
        return

    await state.update_data(t_uid=uid)
    cities = await get_all_cities()

    if not cities:
        await message.answer("У базі немає міст. Додайте міста спочатку.")
        await state.clear()
        return

    buttons = []
    for c in cities:
        buttons.append([InlineKeyboardButton(text=c['name'], callback_data=f"dev_t_city:{c['id']}")])

    await message.answer(
        "Оберіть місто для цього територіала:",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    await state.set_state(DeveloperStates.waiting_add_territorial_city)

async def developer_add_territorial_city(callback: CallbackQuery, state: FSMContext):
    city_id = int(callback.data.split(":")[1])
    city_data = await get_city_by_id(city_id)
    city_name = city_data['name']

    await state.update_data(t_city=city_name)

    # Перевіряємо, які типи вже є в цьому місті
    existing = await get_territorials_by_city(city_name)
    existing_types = {t.get('territorial_type') for t in existing}

    buttons = []
    for code, label in TERRITORIAL_TYPES.items():
        if code in existing_types:
            buttons.append([InlineKeyboardButton(
                text=f"{label} ({code}) — вже існує ⚠️",
                callback_data=f"dev_t_type:{code}"
            )])
        else:
            buttons.append([InlineKeyboardButton(
                text=f"{label} ({code})",
                callback_data=f"dev_t_type:{code}"
            )])

    warning = ""
    if existing_types:
        warning = f"\n\n⚠️ У місті <b>{city_name}</b> вже є: {', '.join(existing_types)}."

    await callback.message.edit_text(
        f"Оберіть тип Територіала для міста <b>{city_name}</b>:{warning}",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    await state.set_state(DeveloperStates.waiting_add_territorial_type)
    await callback.answer()

async def developer_add_territorial_type(callback: CallbackQuery, state: FSMContext):
    t_type = callback.data.split(":")[1]  # "ТЗ" або "ВВ"
    await state.update_data(t_type=t_type)
    data = await state.get_data()
    city_name = data.get("t_city")

    buttons = [
        [InlineKeyboardButton(text="✅ Так, перенести", callback_data="dev_t_transfer:yes")],
        [InlineKeyboardButton(text="❌ Ні, залишити як є", callback_data="dev_t_transfer:no")]
    ]

    type_label = TERRITORIAL_TYPES.get(t_type, t_type)
    await callback.message.edit_text(
        f"Перенести <b>всіх керівників</b> міста <b>{city_name}</b> під управління цього Територіала ({type_label})?\n\n"
        f"💡 Після перенесення можна перепривʼязати окремих керівників вручну.",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons)
    )
    await state.set_state(DeveloperStates.waiting_add_territorial_transfer)
    await callback.answer()

async def developer_add_territorial_finish(callback: CallbackQuery, state: FSMContext):
    transfer = callback.data.split(":")[1] == "yes"
    data = await state.get_data()
    uid = data.get("t_uid")
    city = data.get("t_city")
    t_type = data.get("t_type")

    await add_manager(uid, "Територіал", city=city, territorial_type=t_type)

    if transfer:
        await reassign_city_managers_to_territorial(city, uid)

    type_label = TERRITORIAL_TYPES.get(t_type, t_type)
    await callback.message.edit_text(
        f"✅ Територіала <b>{type_label} ({t_type})</b> для міста <b>{city}</b> успішно додано!" +
        ("\n👔 Всіх керівників міста перенесено." if transfer else ""),
        reply_markup=_territorials_menu_keyboard()
    )
    await state.clear()
    await callback.answer()




# ==============================================================================
# УПРАВЛІННЯ НАГЛЯДАЧАМИ (ПУНКТ 6)
# ==============================================================================

async def _build_observers_menu_view(page: int = 0, bot: Optional[Bot] = None) -> tuple[str, InlineKeyboardMarkup]:
    """Будує інтерфейс списку наглядачів для адмін-панелі."""
    observers = await get_all_observers()
    
    PER_PAGE = 6
    total_count = len(observers)
    pages_total = max(1, (total_count + PER_PAGE - 1) // PER_PAGE)
    page = max(0, min(page, pages_total - 1))
    
    start = page * PER_PAGE
    paginated_obs = observers[start:start + PER_PAGE]
    
    lines = [
        f"👁 <b>Команда Наглядачів (всього: {total_count})</b>",
        "Наглядачі мають доступ лише до перегляду матеріалів та аналітики.",
        ""
    ]
    
    if not observers:
        lines.append("  <i>У системі поки немає призначених наглядачів.</i>")
    
    buttons = []
    for obs in paginated_obs:
        name, username = await _format_identity(obs["uid"], obs.get("full_name"), obs.get("username"), bot=bot)
        btn_text = f"👁 {name} | {username}"
        buttons.append([InlineKeyboardButton(text=btn_text, callback_data=f"dev_obs_view:{obs['uid']}")])
    
    # Кнопки пагінації
    nav_row = []
    if page > 0:
        nav_row.append(InlineKeyboardButton(text="⬅️", callback_data=f"dev_obs_page:{page - 1}"))
    if pages_total > 1:
        nav_row.append(InlineKeyboardButton(text=f"📄 {page + 1}/{pages_total}", callback_data="ignore"))
    if page < pages_total - 1:
        nav_row.append(InlineKeyboardButton(text="➡️", callback_data=f"dev_obs_page:{page + 1}"))
    
    if nav_row:
        buttons.append(nav_row)
        
    buttons.append([InlineKeyboardButton(text="➕ Створити запрошення", callback_data="dev_obs_create_invite")])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="dev_main_team")])
    
    return "\n".join(lines), InlineKeyboardMarkup(inline_keyboard=buttons)


async def developer_observers_menu(callback: CallbackQuery):
    """Показує список наглядачів."""
    if not await _ensure_developer(callback):
        return
    page = 0
    if callback.data and callback.data.startswith("dev_obs_page:"):
        try:
            page = int(callback.data.split(":")[1])
        except (ValueError, IndexError):
            page = 0
            
    text, kb = await _build_observers_menu_view(page=page, bot=callback.bot)
    await _show_admin_photo_menu(callback, "admin_spus.jpg", text, kb)


async def developer_observer_create_invite(callback: CallbackQuery):
    """Генерує токен та посилання для нового наглядача."""
    if not await _ensure_developer(callback):
        return
    
    admin_uid = callback.from_user.id
    token = await generate_token(manager_id=admin_uid, role="Наглядач", expires_in_hours=24)
    bot_info = await callback.bot.get_me()
    bot_username = bot_info.username
    invite_link = f"https://t.me/{bot_username}?start={admin_uid}-{token}"
    
    text = (
        "👁 <b>Запрошення для Наглядача</b>\n\n"
        "🔗 <b>Одноразове посилання:</b>\n"
        f"<code>{invite_link}</code>\n\n"
        "⏳ <b>Термін дії:</b> 24 години.\n"
        "👤 Надішліть це посилання людині, яку призначаєте Наглядачем.\n"
        "Після переходу за посиланням користувач введе своє ПІБ та автоматично отримає роль <b>Наглядач</b>."
    )
    
    buttons = [
        [InlineKeyboardButton(text="⬅️ До списку наглядачів", callback_data="dev_observers_menu")]
    ]
    
    await _show_admin_photo_menu(callback, "admin_spus.jpg", text, InlineKeyboardMarkup(inline_keyboard=buttons))


async def developer_observer_view(callback: CallbackQuery):
    """Показує картку наглядача з фото."""
    if not await _ensure_developer(callback):
        return
    try:
        obs_uid = int(callback.data.split(":")[1])
    except (ValueError, IndexError):
        await callback.answer("Помилка ідентифікатора.", show_alert=True)
        return
        
    obs = await get_observer_by_uid(obs_uid)
    if not obs:
        await callback.answer("Наглядача не знайдено або він вже видалений.", show_alert=True)
        return
        
    name, username = await _format_identity(obs["uid"], obs.get("full_name"), obs.get("username"), bot=callback.bot)
    
    resp_name = "Адміністратор"
    if obs.get("responsible_uid"):
        resp_mgr = await get_manager_by_uid(obs["responsible_uid"])
        if resp_mgr and resp_mgr.get("full_name"):
            resp_name = resp_mgr["full_name"]
            
    text = (
        f"👁 <b>Картка Наглядача</b>\n\n"
        f"👤 <b>ПІБ:</b> {html.escape(name)}\n"
        f"🆔 <b>Telegram ID:</b> <code>{obs['uid']}</code>\n"
        f"📱 <b>Username:</b> {html.escape(username)}\n"
        f"👑 <b>Призначив:</b> {html.escape(resp_name)}\n"
        f"🟢 <b>Статус:</b> Активний (лише перегляд)"
    )
    
    buttons = [
        [InlineKeyboardButton(text="🗑 Видалити наглядача", callback_data=f"dev_obs_del_confirm:{obs_uid}")],
        [InlineKeyboardButton(text="⬅️ До списку наглядачів", callback_data="dev_observers_menu")]
    ]
    
    photo_input = await get_user_avatar_input(callback.bot, obs_uid)
    await _send_or_edit_card_photo(callback, photo_input, text, reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons))


async def developer_observer_delete_confirm(callback: CallbackQuery):
    """Підтвердження видалення наглядача."""
    if not await _ensure_developer(callback):
        return
    try:
        obs_uid = int(callback.data.split(":")[1])
    except (ValueError, IndexError):
        await callback.answer("Помилка ідентифікатора.", show_alert=True)
        return
        
    obs = await get_observer_by_uid(obs_uid)
    if not obs:
        await callback.answer("Наглядача не знайдено.", show_alert=True)
        return
        
    name, _ = await _format_identity(obs["uid"], obs.get("full_name"))
    
    text = (
        f"⚠️ <b>Видалення Наглядача</b>\n\n"
        f"Ви впевнені, що хочете видалити наглядача <b>{html.escape(name)}</b>?\n"
        f"Користувач втратить доступ до панелі наглядача."
    )
    
    buttons = [
        [InlineKeyboardButton(text="🗑 Так, видалити", callback_data=f"dev_obs_del_do:{obs_uid}")],
        [InlineKeyboardButton(text="❌ Скасувати", callback_data=f"dev_obs_view:{obs_uid}")]
    ]
    
    await _edit_or_answer(callback.message, text, reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons))
    await callback.answer()


async def developer_observer_delete_do(callback: CallbackQuery):
    """Виконує видалення наглядача."""
    if not await _ensure_developer(callback):
        return
    try:
        obs_uid = int(callback.data.split(":")[1])
    except (ValueError, IndexError):
        await callback.answer("Помилка ідентифікатора.", show_alert=True)
        return
        
    await delete_observer_by_uid(obs_uid)
    await callback.answer("✅ Наглядача успішно видалено.", show_alert=True)
    text, kb = await _build_observers_menu_view(page=0)
    await _edit_or_answer(callback.message, text, reply_markup=kb)


async def developer_promote_manager_handler(callback: CallbackQuery):
    """Обробляє рішення територіала перевести керівника-стажера в статус 'Керівник'."""
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access or (not is_admin and not is_territorial):
        await callback.answer("⛔️ Доступ заборонено.", show_alert=True)
        return
        
    try:
        manager_uid = int(callback.data.split(":")[1])
    except (ValueError, IndexError):
        await callback.answer("Помилка ідентифікатора.", show_alert=True)
        return
        
    mgr = await get_manager_by_uid(manager_uid)
    if not mgr:
        await callback.answer("Керівника не знайдено.", show_alert=True)
        return
        
    await promote_manager_trainee(manager_uid)
    name, _ = await _format_identity(mgr["uid"], mgr.get("full_name"), mgr.get("username"))
    
    updated_text = (
        f"✅ <b>Керівника успішно переведено!</b>\n\n"
        f"👤 <b>Керівник:</b> {html.escape(name)}\n"
        f"🟢 <b>Новий статус:</b> Керівник\n"
        f"🕒 <b>Дата переведення:</b> {datetime.now(pytz.timezone(TIMEZONE)).strftime('%Y-%m-%d %H:%M')}"
    )
    
    try:
        if callback.message.photo:
            await callback.message.edit_caption(caption=updated_text, reply_markup=None)
        else:
            await callback.message.edit_text(updated_text, reply_markup=None)
    except Exception:
        pass
        
    await callback.answer("✅ Керівника переведено в статус «Керівник»!", show_alert=True)
    
    # Сповіщення самому керівнику
    try:
        congrats_text = (
            "🎉 <b>Вітаємо!</b>\n\n"
            "Ваш територіальний керівник підтвердив успішне проходження навчального курсу.\n"
            "Вам офіційно присвоєно статус <b>Керівник</b>! 👔✨\n\n"
            "Всі навчальні матеріали залишаються доступними для перегляду в <b>Панелі керівника</b>."
        )
        await callback.bot.send_message(manager_uid, congrats_text, parse_mode="HTML")
    except Exception:
        pass


async def developer_keep_manager_trainee_handler(callback: CallbackQuery):
    """Обробляє рішення територіала залишити керівника стажером."""
    has_access, is_admin, is_territorial = await _check_access(callback)
    if not has_access or (not is_admin and not is_territorial):
        await callback.answer("⛔️ Доступ заборонено.", show_alert=True)
        return
        
    try:
        manager_uid = int(callback.data.split(":")[1])
    except (ValueError, IndexError):
        await callback.answer("Помилка ідентифікатора.", show_alert=True)
        return
        
    mgr = await get_manager_by_uid(manager_uid)
    name = mgr.get("full_name") or f"ID {manager_uid}" if mgr else f"ID {manager_uid}"
    
    updated_text = (
        f"⏳ <b>Рішення зафіксовано</b>\n\n"
        f"👤 <b>Керівник:</b> {html.escape(name)}\n"
        f"🟡 <b>Поточний статус:</b> Керівник Стажер (навчання продовжується)\n"
        f"Ви зможете перевести його у статус «Керівник» пізніше в меню управління керівниками."
    )
    
    try:
        if callback.message.photo:
            await callback.message.edit_caption(caption=updated_text, reply_markup=None)
        else:
            await callback.message.edit_text(updated_text, reply_markup=None)
    except Exception:
        pass
        
    await callback.answer("Статус керівника залишено без змін.", show_alert=True)


def register_developer_menu_handlers(dp: Dispatcher):

    # Manager Trainee Promotion (Пункт 6)
    dp.callback_query.register(developer_promote_manager_handler, lambda c: c.data and c.data.startswith("dev_promote_mgr:"))
    dp.callback_query.register(developer_keep_manager_trainee_handler, lambda c: c.data and c.data.startswith("dev_keep_mgr_trainee:"))

    # Observers (Пункт 6)
    dp.callback_query.register(developer_observers_menu, lambda c: c.data == "dev_observers_menu" or (c.data and c.data.startswith("dev_obs_page:")))
    dp.callback_query.register(developer_observer_create_invite, lambda c: c.data == "dev_obs_create_invite")
    dp.callback_query.register(developer_observer_view, lambda c: c.data and c.data.startswith("dev_obs_view:"))
    dp.callback_query.register(developer_observer_delete_confirm, lambda c: c.data and c.data.startswith("dev_obs_del_confirm:"))
    dp.callback_query.register(developer_observer_delete_do, lambda c: c.data and c.data.startswith("dev_obs_del_do:"))

    # Territorials
    dp.callback_query.register(developer_territorials_menu, lambda c: c.data == "dev_territorials_menu")
    dp.callback_query.register(developer_list_territorials, lambda c: c.data == "dev_territorials_list")
    dp.callback_query.register(developer_territorial_view, lambda c: c.data and c.data.startswith("dev_territorial_view:"))
    dp.callback_query.register(developer_territorial_managers, lambda c: c.data and c.data.startswith("dev_territorial_managers:"))
    dp.callback_query.register(developer_territorial_del, lambda c: c.data and c.data.startswith("dev_territorial_del:") and not c.data.startswith("dev_territorial_del_confirm:"))
    dp.callback_query.register(developer_territorial_del_confirm, lambda c: c.data and c.data.startswith("dev_territorial_del_confirm:"))

    dp.callback_query.register(developer_add_territorial_start, lambda c: c.data == "dev_territorial_add")
    dp.message.register(developer_add_territorial_uid, DeveloperStates.waiting_add_territorial_uid)
    dp.callback_query.register(developer_add_territorial_city, DeveloperStates.waiting_add_territorial_city, lambda c: c.data and c.data.startswith("dev_t_city:"))
    dp.callback_query.register(developer_add_territorial_type, DeveloperStates.waiting_add_territorial_type, lambda c: c.data and c.data.startswith("dev_t_type:"))
    dp.callback_query.register(developer_add_territorial_finish, DeveloperStates.waiting_add_territorial_transfer, lambda c: c.data and c.data.startswith("dev_t_transfer:"))

    # Positions
    dp.callback_query.register(dev_positions_menu, lambda c: c.data == "dev_positions_menu")
    dp.callback_query.register(dev_positions_page, lambda c: c.data and c.data.startswith("dev_positions_page:"))
    dp.callback_query.register(dev_pos_add_start, lambda c: c.data == "dev_pos_add")
    dp.message.register(dev_pos_add_name, DeveloperStates.waiting_add_position_name)
    dp.message.register(dev_pos_add_days_text, DeveloperStates.waiting_add_position_days)
    dp.callback_query.register(dev_pos_add_days_callback, DeveloperStates.waiting_add_position_days, lambda c: c.data and c.data.startswith("dev_pos_add_days:"))
    dp.callback_query.register(dev_pos_add_type_callback, DeveloperStates.waiting_add_position_type, lambda c: c.data and c.data.startswith("dev_pos_add_type:"))
    dp.callback_query.register(dev_pos_view, lambda c: c.data and c.data.startswith("dev_pos_view:"))
    dp.callback_query.register(dev_pos_edit_name_start, lambda c: c.data and c.data.startswith("dev_pos_edit_name:"))
    dp.message.register(dev_pos_edit_name_process, DeveloperStates.waiting_edit_position_name)
    dp.callback_query.register(dev_pos_edit_days_start, lambda c: c.data and c.data.startswith("dev_pos_edit_days:"))
    dp.message.register(dev_pos_edit_days_process, DeveloperStates.waiting_edit_position_days)
    dp.callback_query.register(dev_pos_toggle_type, lambda c: c.data and c.data.startswith("dev_pos_toggle_type:"))

    # Cities
    dp.callback_query.register(dev_cities_menu, lambda c: c.data == "dev_cities_menu")
    dp.callback_query.register(dev_cities_page, lambda c: c.data and c.data.startswith("dev_cities_page:"))
    dp.callback_query.register(dev_city_add_start, lambda c: c.data == "dev_city_add")
    dp.message.register(dev_city_add_name, DeveloperStates.waiting_add_city_name)
    dp.callback_query.register(dev_city_view, lambda c: c.data and c.data.startswith("dev_city_view:"))
    dp.callback_query.register(dev_city_edit_name_start, lambda c: c.data and c.data.startswith("dev_city_edit_name:"))
    dp.message.register(dev_city_edit_name_process, DeveloperStates.waiting_edit_city_name)
    dp.callback_query.register(dev_city_delete, lambda c: c.data and c.data.startswith("dev_city_delete:"))
    dp.callback_query.register(dev_city_delete_confirm, lambda c: c.data and c.data.startswith("dev_city_delete_confirm:"))

    # XLSX Reports (Priority)
    dp.callback_query.register(developer_xlsx_menu, lambda c: c.data == "dev_xlsx_menu")
    dp.callback_query.register(developer_send_current_report, lambda c: c.data == "dev_xlsx_send_current")
    dp.callback_query.register(developer_send_period_report, lambda c: c.data and c.data.startswith("dev_xlsx_send_period:"))

    dp.callback_query.register(developer_menu_callback, lambda c: c.data == "developer_menu")
    dp.callback_query.register(dev_main_study_handler, lambda c: c.data == "dev_main_study")
    dp.callback_query.register(dev_main_analyt_handler, lambda c: c.data == "dev_main_analyt")
    dp.callback_query.register(dev_main_team_handler, lambda c: c.data == "dev_main_team")
    dp.callback_query.register(dev_main_other_handler, lambda c: c.data == "dev_main_other")

    # Users
    dp.callback_query.register(developer_users_menu, lambda c: c.data == "dev_users_menu")
    dp.callback_query.register(developer_list_users, lambda c: c.data == "dev_users_list")
    dp.callback_query.register(developer_list_interns, lambda c: c.data == "dev_users_interns")
    dp.callback_query.register(developer_list_workers, lambda c: c.data == "dev_users_workers")
    dp.callback_query.register(developer_interns_export_xlsx, lambda c: c.data == "dev_interns_export_xlsx")
    dp.callback_query.register(developer_workers_export_xlsx, lambda c: c.data == "dev_workers_export_xlsx")
    dp.callback_query.register(developer_bulk_promote_completed_interns, lambda c: c.data == "dev_users_bulk_promote")
    
    # Active/Inactive users
    dp.callback_query.register(developer_active_users, lambda c: c.data == "dev_users_active")
    dp.callback_query.register(developer_inactive_users, lambda c: c.data == "dev_users_inactive")
    
    # City filter
    dp.callback_query.register(developer_users_by_city_menu, lambda c: c.data == "dev_users_by_city")
    dp.callback_query.register(developer_users_filter_city, lambda c: c.data and c.data.startswith("dev_users_filter_city:"))
    
    # Updated pagination
    dp.callback_query.register(developer_users_pagination_handler, lambda c: c.data and (c.data.startswith("dev_users_pag:") or c.data.startswith("dev_users_page:")))
    
    dp.callback_query.register(developer_request_user_search, lambda c: c.data == "dev_users_search")
    dp.callback_query.register(developer_request_delete_user, lambda c: c.data == "dev_users_delete")
    dp.message.register(developer_process_user_search, DeveloperStates.waiting_user_search)
    dp.message.register(developer_process_delete_user, DeveloperStates.waiting_delete_user)

    # Managers Team Management
    dp.callback_query.register(developer_manage_managers_menu, lambda c: c.data == "dev_manage_managers")
    dp.callback_query.register(developer_managers_pagination_handler, lambda c: c.data and c.data.startswith("dev_mgr_page:"))
    dp.callback_query.register(developer_list_managers, lambda c: c.data == "dev_managers_list")
    dp.callback_query.register(developer_request_add_manager, lambda c: c.data == "dev_add_manager")
    dp.message.register(developer_process_add_manager_id, DeveloperStates.waiting_add_manager_id)
    dp.message.register(developer_process_add_manager_name, DeveloperStates.waiting_add_manager_name) # NEW
    dp.callback_query.register(developer_process_add_manager_city, DeveloperStates.waiting_add_manager_city, lambda c: c.data.startswith("dev_mgr_city_select:"))
    dp.callback_query.register(developer_process_shop_selection, DeveloperStates.waiting_add_manager_shops, lambda c: c.data.startswith("dev_mgr_shop_toggle:"))
    dp.callback_query.register(developer_finish_shop_selection, DeveloperStates.waiting_add_manager_shops, lambda c: c.data == "dev_mgr_shop_done")
    dp.callback_query.register(developer_remove_manager_menu, lambda c: c.data == "dev_remove_manager_menu")
    dp.callback_query.register(developer_remove_manager, lambda c: c.data and c.data.startswith("mgr_remove:"))
    dp.callback_query.register(manager_cancel_add, lambda c: c.data == "mgr_cancel_add")
    dp.callback_query.register(manager_team_back, lambda c: c.data == "mgr_team_back")

    # Manager Search & Card Management Handlers (Point 5)
    dp.callback_query.register(dev_mgr_search_start, lambda c: c.data == "dev_mgr_search_start")
    dp.message.register(dev_mgr_search_process, DeveloperStates.waiting_search_manager)
    dp.callback_query.register(dev_mgr_view, lambda c: c.data and c.data.startswith("dev_mgr_view:"))
    dp.callback_query.register(dev_mgr_edit_menu, lambda c: c.data and c.data.startswith("dev_mgr_edit_menu:"))
    dp.callback_query.register(dev_mgr_change_name_start, lambda c: c.data and c.data.startswith("dev_mgr_ch_name:"))
    dp.message.register(dev_mgr_change_name_process, DeveloperStates.waiting_edit_manager_name)
    dp.callback_query.register(dev_mgr_change_shops_start, lambda c: c.data and c.data.startswith("dev_mgr_ch_shops:"))
    dp.callback_query.register(dev_mgr_edit_shop_toggle, DeveloperStates.waiting_edit_manager_shops, lambda c: c.data and c.data.startswith("dev_mgr_eshop_toggle:"))
    dp.callback_query.register(dev_mgr_edit_shop_done, DeveloperStates.waiting_edit_manager_shops, lambda c: c.data and c.data.startswith("dev_mgr_eshop_done:"))
    dp.callback_query.register(dev_mgr_transfer_shops_yes, lambda c: c.data and c.data.startswith("dev_mgr_tr_yes:"))
    dp.callback_query.register(dev_mgr_fire_ask, lambda c: c.data and c.data.startswith("dev_mgr_fire_ask:"))
    dp.callback_query.register(dev_mgr_fire_confirm, lambda c: c.data and c.data.startswith("dev_mgr_fire_confirm:"))
    dp.callback_query.register(dev_mgr_reassign_menu, lambda c: c.data and c.data.startswith("dev_mgr_reassign_menu:"))
    dp.callback_query.register(dev_mgr_reassign_set, lambda c: c.data and c.data.startswith("dev_mgr_reassign_set:"))
    
    # HR Team Management (moved up or handled here)
    dp.callback_query.register(developer_request_add_hr, lambda c: c.data == "dev_add_hr")
    dp.message.register(developer_process_add_hr, DeveloperStates.waiting_add_hr)
    dp.callback_query.register(hr_cancel_add_role, lambda c: c.data == "hr_cancel_add_role")
    dp.callback_query.register(hr_team_member_handler, lambda c: c.data and c.data.startswith("hr_team_member_"))
    dp.callback_query.register(hr_remove_member, lambda c: c.data and c.data.startswith("hr_remove_"))
    dp.callback_query.register(hr_team_back, lambda c: c.data == "hr_team_back")

    # Tokens & Health
    dp.callback_query.register(developer_tokens_menu, lambda c: c.data == "dev_tokens_menu")
    dp.callback_query.register(developer_tokens_cleanup, lambda c: c.data == "dev_tokens_cleanup")
    dp.callback_query.register(developer_health_status, lambda c: c.data == "dev_health_status")

    # Days picker pagination (Materials, Videos, Photos, Tests)
    dp.callback_query.register(developer_days_page_callback, lambda c: c.data and any(c.data.startswith(p) for p in ["dev_mat_day_pg|", "dev_test_day_pg|", "dev_video_day_pg|", "dev_photo_day_pg|"]))

    # Material Acknowledgment Analytics (Point 8)
    dp.callback_query.register(developer_ack_analytics_menu, lambda c: c.data == "dev_an_ack_menu")
    dp.callback_query.register(developer_ack_category_events_view, lambda c: c.data and (c.data.startswith("dev_ack_cat:") or c.data.startswith("dev_ack_cat_pg:")))
    dp.callback_query.register(developer_ack_event_card_view, lambda c: c.data and c.data.startswith("dev_ack_ev:"))
    dp.callback_query.register(developer_ack_shops_list_view, lambda c: c.data and c.data.startswith("dev_ack_shops:"))
    dp.callback_query.register(developer_ack_shop_users_view, lambda c: c.data and c.data.startswith("dev_ack_shop_u:"))

    # Materials Editor
    dp.callback_query.register(developer_materials_menu, lambda c: c.data == "dev_materials_menu")
    dp.callback_query.register(developer_materials_select_day, lambda c: c.data and c.data.startswith("dev_mat_role|"))
    dp.callback_query.register(developer_materials_select_type, lambda c: c.data and c.data.startswith("dev_mat_day|"))
    dp.callback_query.register(developer_materials_view, lambda c: c.data and c.data.startswith("dev_mat_type|"))
    dp.callback_query.register(developer_materials_full_view, lambda c: c.data and c.data.startswith("dev_mat_full_view|"))
    dp.callback_query.register(developer_pagination_handler, lambda c: c.data and c.data.startswith("dev_pag:")) # NEW
    dp.callback_query.register(developer_material_fix_start, lambda c: c.data and c.data.startswith("dev_mat_fix:"))
    dp.callback_query.register(developer_material_delete_request, lambda c: c.data and c.data.startswith("dev_mat_del_req:"))
    dp.callback_query.register(developer_material_delete_confirm, lambda c: c.data and c.data.startswith("dev_mat_del_conf:"))
    dp.message.register(developer_process_fix_page, DeveloperStates.waiting_fix_page)
    
    # Complement handlers
    dp.callback_query.register(developer_material_complement_menu, lambda c: c.data and c.data.startswith("dev_mat_complement|"))
    dp.callback_query.register(developer_material_complement_start, lambda c: c.data and c.data.startswith("dev_comp_type:"))
    dp.message.register(developer_process_insert_start_idx, DeveloperStates.waiting_insert_start_idx)
    dp.message.register(developer_process_insert_end_idx, DeveloperStates.waiting_insert_end_idx)
    dp.message.register(developer_process_complement_content, DeveloperStates.waiting_complement_content)
    dp.callback_query.register(developer_material_menu_back, lambda c: c.data == "dev_mat_menu_back")
    
    # Array insert and range replace handlers
    dp.callback_query.register(developer_process_insert_mode, lambda c: c.data and c.data.startswith("dev_ins_mode:"))
    dp.message.register(developer_process_insert_array_content, DeveloperStates.waiting_insert_array_content)
    dp.callback_query.register(developer_process_insert_array_done, lambda c: c.data == "dev_ins_array_done")
    
    dp.callback_query.register(developer_material_replace_range_start, lambda c: c.data and c.data.startswith("dev_mat_replace|"))
    dp.message.register(developer_process_replace_range, DeveloperStates.waiting_replace_range)
    dp.message.register(developer_process_replace_array_content, DeveloperStates.waiting_replace_array_content)
    dp.callback_query.register(developer_process_replace_array_done, lambda c: c.data == "dev_rep_array_done")
    dp.callback_query.register(developer_process_notify_decision, lambda c: c.data and c.data.startswith("dev_notify:"))
    
    dp.callback_query.register(developer_materials_edit_start, lambda c: c.data and c.data.startswith("dev_mat_edit|"))
    dp.message.register(developer_process_material_parts, DeveloperStates.waiting_material_parts)
    dp.message.register(developer_process_material_video, DeveloperStates.waiting_material_video)


    # Video Editor
    dp.callback_query.register(developer_videos_menu, lambda c: c.data == "dev_videos_menu")
    dp.callback_query.register(developer_videos_select_day, lambda c: c.data and c.data.startswith("dev_video_role|"))
    dp.callback_query.register(developer_videos_view, lambda c: c.data and c.data.startswith("dev_video_day|"))
    dp.callback_query.register(developer_video_edit_start, lambda c: c.data == "dev_video_edit")
    dp.message.register(developer_process_video_uploads, DeveloperStates.waiting_video_uploads)

    # Photo Editor
    dp.callback_query.register(developer_photos_menu, lambda c: c.data == "dev_photos_menu")
    dp.callback_query.register(developer_photos_select_day, lambda c: c.data and c.data.startswith("dev_photo_role|"))
    dp.callback_query.register(developer_photos_view, lambda c: c.data and c.data.startswith("dev_photo_day|"))
    dp.callback_query.register(developer_photo_edit_start, lambda c: c.data == "dev_photo_edit")
    dp.message.register(developer_process_photo_uploads, DeveloperStates.waiting_photo_uploads)

    # Syllabus Editor
    dp.callback_query.register(developer_syllabus_menu, lambda c: c.data == "dev_syllabus_menu")
    dp.callback_query.register(developer_syllabus_view, lambda c: c.data and c.data.startswith("dev_syl_role|"))
    dp.callback_query.register(developer_syllabus_toggle, lambda c: c.data and c.data.startswith("dev_syl_toggle:"))
    dp.callback_query.register(developer_syllabus_edit, lambda c: c.data == "dev_syl_edit")
    dp.message.register(developer_process_syllabus, DeveloperStates.waiting_syllabus_text)

    # Tests Editor
    dp.callback_query.register(developer_tests_menu, lambda c: c.data == "dev_tests_menu")
    dp.callback_query.register(developer_tests_select_day, lambda c: c.data and c.data.startswith("dev_test_role|"))
    dp.callback_query.register(developer_tests_view, lambda c: c.data and c.data.startswith("dev_test_day|"))
    dp.callback_query.register(developer_test_toggle, lambda c: c.data and c.data.startswith("dev_test_toggle:"))
    dp.callback_query.register(developer_tests_edit_start, lambda c: c.data == "dev_test_edit")
    dp.message.register(developer_tests_process_input, DeveloperStates.waiting_test_input)

    # Notification handlers
    dp.callback_query.register(notify_decision_yes, lambda c: c.data == "notify_yes")
    dp.callback_query.register(notify_decision_no, lambda c: c.data == "notify_no")
    dp.callback_query.register(notify_cancel, lambda c: c.data == "notify_cancel")
    dp.message.register(process_notify_message, DeveloperStates.waiting_notify_message)

    # Days
    dp.callback_query.register(developer_days_menu, lambda c: c.data == "dev_days_menu")
    dp.callback_query.register(developer_request_days_status, lambda c: c.data == "dev_days_status")
    dp.callback_query.register(developer_request_days_open, lambda c: c.data == "dev_days_open")
    dp.callback_query.register(developer_request_days_close, lambda c: c.data == "dev_days_close")
    dp.callback_query.register(developer_request_days_reset, lambda c: c.data == "dev_days_reset")
    dp.message.register(developer_process_days_status, DeveloperStates.waiting_days_status)
    dp.message.register(developer_process_days_open, DeveloperStates.waiting_days_open)
    dp.message.register(developer_process_days_close, DeveloperStates.waiting_days_close)
    dp.message.register(developer_process_days_reset, DeveloperStates.waiting_days_reset)

    # Dev Team Management
    dp.callback_query.register(developer_dev_team_menu, lambda c: c.data == "dev_team_menu")
    dp.callback_query.register(developer_admin_view, lambda c: c.data and c.data.startswith("dev_admin_view:"))
    dp.callback_query.register(developer_request_add_dev, lambda c: c.data == "dev_add_dev")
    dp.callback_query.register(developer_remove_dev_menu, lambda c: c.data == "dev_remove_dev_menu")
    dp.callback_query.register(developer_remove_dev, lambda c: c.data and c.data.startswith("dev_remove_confirm:"))
    dp.callback_query.register(developer_cancel_add_role, lambda c: c.data == "dev_cancel_add_role")
    dp.message.register(developer_process_add_dev, DeveloperStates.waiting_add_developer)
    dp.message.register(developer_process_add_dev_name, DeveloperStates.waiting_add_developer_name)

    # New User Profile Actions
    dp.callback_query.register(developer_days_manage_menu, lambda c: c.data and c.data.startswith("dev_days_manage:"))
    dp.callback_query.register(developer_toggle_day_status, lambda c: c.data and c.data.startswith("dev_toggle_day:"))
    dp.callback_query.register(developer_user_profile_back, lambda c: c.data and c.data.startswith("dev_user_profile_back:"))
    dp.callback_query.register(developer_user_modify_menu, lambda c: c.data and c.data.startswith("dev_user_modify:"))
    dp.callback_query.register(developer_change_manager_start, lambda c: c.data and c.data.startswith("dev_change_manager:"))
    dp.message.register(developer_process_change_manager, DeveloperStates.waiting_change_manager_id)
    dp.callback_query.register(developer_change_role_start, lambda c: c.data and c.data.startswith("dev_change_role:"))
    dp.callback_query.register(developer_process_change_role, lambda c: c.data and (c.data.startswith("dev_process_change_role:") or c.data.startswith("dev_set_role:")))
    dp.callback_query.register(developer_change_city_start, lambda c: c.data and c.data.startswith("dev_change_city:"))
    dp.callback_query.register(developer_process_change_city, lambda c: c.data and (c.data.startswith("dev_process_change_city:") or c.data.startswith("dev_set_city:")))
    dp.callback_query.register(developer_user_delete_confirm, lambda c: c.data and c.data.startswith("dev_user_delete_confirm:"))
    dp.callback_query.register(developer_user_delete_perform, lambda c: c.data and c.data.startswith("dev_user_delete_perform:"))
    
    # HR Team Management
    dp.callback_query.register(developer_manage_hrs_menu, lambda c: c.data == "dev_hr_team_menu")
    
    # Legacy/Unified handlers (keep for MAIN_DEVELOPER_ID or specific access)
    # Removed duplicate manager handlers as they are now handled in the 'Managers' section above

    # The following handlers were moved or replaced:
    # dp.callback_query.register(developer_manage_managers_menu, lambda c: c.data == "dev_manage_managers")
    # dp.callback_query.register(developer_request_add_manager, lambda c: c.data == "dev_add_manager")
    # dp.message.register(developer_process_add_manager_id, DeveloperStates.waiting_add_manager_id)
    # dp.callback_query.register(developer_process_shop_selection, DeveloperStates.waiting_add_manager_shops, lambda c: c.data.startswith("dev_mgr_shop_toggle:"))
    # dp.callback_query.register(developer_finish_shop_selection, DeveloperStates.waiting_add_manager_shops, lambda c: c.data == "dev_mgr_shop_done")
    dp.callback_query.register(developer_remove_manager_menu, lambda c: c.data == "dev_remove_manager_menu")
    dp.callback_query.register(developer_remove_manager, lambda c: c.data and c.data.startswith("mgr_remove:"))
    dp.callback_query.register(manager_cancel_add, lambda c: c.data == "mgr_cancel_add")
    dp.callback_query.register(manager_team_back, lambda c: c.data == "mgr_team_back")

    # Analytics
    dp.callback_query.register(developer_analytics_menu, lambda c: c.data == "dev_analytics_menu")
    dp.callback_query.register(developer_daily_digest, lambda c: c.data == "dev_analytics_digest")
    dp.callback_query.register(developer_dropout_report, lambda c: c.data == "dev_analytics_funnel")
    dp.callback_query.register(developer_training_menu, lambda c: c.data == "dev_training_menu")
    dp.callback_query.register(developer_training_added, lambda c: c.data and c.data.startswith("dev_training_added:"))
    dp.callback_query.register(developer_training_added_city_menu, lambda c: c.data and c.data.startswith("dev_training_added_city_menu:"))
    dp.callback_query.register(developer_training_added_city, lambda c: c.data and c.data.startswith("dev_training_added_city:"))
    dp.callback_query.register(developer_training_added_export, lambda c: c.data and c.data.startswith("dev_training_added_export:"))
    dp.callback_query.register(developer_training_added_clear_ask, lambda c: c.data and c.data.startswith("dev_training_added_clear_ask:"))
    dp.callback_query.register(developer_training_added_clear_cancel, lambda c: c.data and c.data.startswith("dev_training_added_clear_cancel:"))
    dp.callback_query.register(developer_training_added_clear_confirm, lambda c: c.data and c.data.startswith("dev_training_added_clear_confirm:"))
    dp.callback_query.register(developer_training_left, lambda c: c.data == "dev_training_left")
    dp.callback_query.register(developer_training_promoted, lambda c: c.data and c.data.startswith("dev_training_promoted:"))
    dp.callback_query.register(developer_training_promoted_page, lambda c: c.data and c.data.startswith("dev_training_promoted_page:"))
    dp.callback_query.register(developer_training_rejected, lambda c: c.data and c.data.startswith("dev_training_rejected:"))
    dp.callback_query.register(developer_training_promoted_export, lambda c: c.data and c.data.startswith("dev_training_promoted_export:"))
    dp.callback_query.register(developer_training_rejected_export, lambda c: c.data and c.data.startswith("dev_training_rejected_export:"))
    dp.callback_query.register(developer_reminder_history_menu, lambda c: c.data == "dev_reminder_history")

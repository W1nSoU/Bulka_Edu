
from __future__ import annotations
import aiosqlite
import json
import os
from datetime import datetime, timedelta
import pytz
from typing import List, Dict, Any
from io import BytesIO
import openpyxl
from aiogram import Bot
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from database import DB_PATH
from bot.config import TIMEZONE, DAYS_TOTAL

async def get_report_data(start_date: datetime, end_date: datetime) -> List[Dict[str, Any]]:
    """
    Агрегує дані для звіту за вказаний період.
    Повертає список словників з метриками по містах та магазинах.
    """
    from database.positions import get_days_count_for_role
    from database.users import MANAGEMENT_ROLES, _get_privileged_user_ids

    start_str = start_date.strftime("%Y-%m-%d %H:%M:%S")
    end_str = end_date.strftime("%Y-%m-%d %H:%M:%S")
    
    # 3 дні неактивності для метрики "Активні"
    active_cutoff = (datetime.now(pytz.timezone(TIMEZONE)) - timedelta(days=3)).strftime("%Y-%m-%d %H:%M:%S")
    privileged_ids = await _get_privileged_user_ids()

    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        
        # Отримуємо унікальні пари місто-магазин
        cursor = await db.execute("SELECT DISTINCT city, shop FROM users WHERE city IS NOT NULL AND shop IS NOT NULL")
        locations = await cursor.fetchall()
        
        report_rows = []
        
        for loc in locations:
            city = loc['city']
            shop = loc['shop']
            
            # 1. Нові стажери (зареєстровані в періоді, не працівники, не менеджмент)
            cursor = await db.execute(
                """
                SELECT user_id, role FROM users 
                WHERE city = ? AND shop = ? 
                  AND (status IS NULL OR status != 'Працівник')
                  AND first_seen BETWEEN ? AND ?
                """,
                (city, shop, start_str, end_str)
            )
            new_candidates = await cursor.fetchall()
            new_count = sum(1 for r in new_candidates if r[0] not in privileged_ids and r[1] not in MANAGEMENT_ROLES)
            
            # 2. Активні стажери (в процесі ТА активність < 3 дні)
            cursor = await db.execute(
                """
                SELECT u.user_id, u.role FROM users u
                WHERE u.city = ? AND u.shop = ?
                  AND (u.status IS NULL OR u.status != 'Працівник')
                  AND u.last_activity >= ?
                """,
                (city, shop, active_cutoff)
            )
            active_candidates = await cursor.fetchall()
            active_count = 0
            for row in active_candidates:
                uid, role = row[0], row[1]
                if uid in privileged_ids or role in MANAGEMENT_ROLES:
                    continue
                req_days = await get_days_count_for_role(role or "")
                cursor = await db.execute("SELECT COUNT(*) FROM progress WHERE user_id = ? AND completed = 1", (uid,))
                completed_days = (await cursor.fetchone())[0]
                if completed_days < req_days:
                    active_count += 1
            
            # 3. Відсів (dropout)
            # Визначаємо як: не завершили навчання ТА не заходили більше 3 днів
            cursor = await db.execute(
                """
                SELECT u.user_id, u.role FROM users u
                WHERE u.city = ? AND u.shop = ?
                  AND (u.status IS NULL OR u.status != 'Працівник')
                  AND u.last_activity < ?
                """,
                (city, shop, active_cutoff)
            )
            potential_dropouts = await cursor.fetchall()
            dropout_count = 0
            for row in potential_dropouts:
                uid, role = row[0], row[1]
                if uid in privileged_ids or role in MANAGEMENT_ROLES:
                    continue
                req_days = await get_days_count_for_role(role or "")
                cursor = await db.execute("SELECT COUNT(*) FROM progress WHERE user_id = ? AND completed = 1", (uid,))
                completed_days = (await cursor.fetchone())[0]
                if completed_days < req_days:
                    dropout_count += 1
            
            # 4. Частота запитів до керівників
            cursor = await db.execute(
                """
                SELECT COUNT(*) FROM support_logs sl
                JOIN users u ON sl.user_id = u.user_id
                WHERE u.city = ? AND u.shop = ? AND sl.created_at BETWEEN ? AND ?
                """,
                (city, shop, start_str, end_str)
            )
            support_requests = (await cursor.fetchone())[0]
            
            if new_count > 0 or active_count > 0 or dropout_count > 0 or support_requests > 0:
                report_rows.append({
                    "city": city,
                    "shop": shop,
                    "new": new_count,
                    "active": active_count,
                    "dropout": dropout_count,
                    "support": support_requests
                })
                
        return report_rows


async def get_report_details(start_date: datetime, end_date: datetime) -> Dict[str, List[Dict[str, Any]]]:
    """
    Повертає деталізацію для XLSX:
    - хто додався
    - хто відсіявся (автовидалені або звільнені)
    - хто став працівником
    - хто відхилений керівником
    """
    start_str = start_date.strftime("%Y-%m-%d %H:%M:%S")
    end_str = end_date.strftime("%Y-%m-%d %H:%M:%S")

    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row

        async def _fetch(event_types: list[str]) -> List[Dict[str, Any]]:
            placeholders = ",".join("?" for _ in event_types)
            cur = await db.execute(
                f"""
                SELECT user_id, full_name, username, city, role, manager_id, event_at
                FROM training_events
                WHERE event_type IN ({placeholders}) AND event_at BETWEEN ? AND ?
                ORDER BY event_at DESC
                """,
                (*event_types, start_str, end_str),
            )
            return [dict(r) for r in await cur.fetchall()]

        return {
            "added": await _fetch(["added"]),
            "dropped": await _fetch(["left_deleted", "fired"]),
            "promoted": await _fetch(["promoted"]),
            "rejected": await _fetch(["rejected"]),
        }


def _write_detail_sheet(
    wb: openpyxl.Workbook,
    title: str,
    rows: List[Dict[str, Any]],
    *,
    color: str,
) -> None:
    ws = wb.create_sheet(title=title)
    headers = ["Дата", "ПІБ", "Username", "Місто", "Посада", "Керівник ID", "Telegram ID"]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )

    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = border

    for row in rows:
        ws.append([
            row.get("event_at"),
            row.get("full_name") or "—",
            row.get("username") or "—",
            row.get("city") or "—",
            row.get("role") or "—",
            row.get("manager_id") or "—",
            row.get("user_id") or "—",
        ])
        for cell in ws[ws.max_row]:
            cell.border = border

    if not rows:
        ws.append(["Немає даних за обраний період", "", "", "", "", "", ""])

    widths = {"A": 20, "B": 35, "C": 20, "D": 15, "E": 25, "F": 14, "G": 14}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def generate_xlsx_report(
    data: List[Dict[str, Any]],
    start_date: datetime,
    end_date: datetime,
    details: Dict[str, List[Dict[str, Any]]] | None = None,
) -> BytesIO:
    """Генерує XLSX файл у пам'яті."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Зведення"
    
    # Стилі
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'), 
        right=Side(style='thin'), 
        top=Side(style='thin'), 
        bottom=Side(style='thin')
    )
    
    # Заголовок звіту
    ws.merge_cells('A1:F1')
    ws['A1'] = f"Звіт з {start_date.strftime('%d.%m.%Y')} по {end_date.strftime('%d.%m.%Y')}"
    ws['A1'].font = Font(size=14, bold=True)
    ws['A1'].alignment = center_align
    
    # Заголовки колонок
    headers = ["Місто", "Магазин", "Нові стажери", "Активні", "Відсів", "Запити до керівників"]
    ws.append(headers)
    
    for cell in ws[2]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = border
        
    # Дані
    for row_data in data:
        row = [
            row_data['city'],
            row_data['shop'],
            row_data['new'],
            row_data['active'],
            row_data['dropout'],
            row_data['support']
        ]
        ws.append(row)
        for cell in ws[ws.max_row]:
            cell.border = border
            cell.alignment = Alignment(vertical="center")
            if isinstance(cell.value, int):
                cell.alignment = center_align

    # Автоматична ширина колонок
    for col in ws.columns:
        max_length = 0
        column = None
        for cell in col:
            # Пропускаємо MergedCell
            if hasattr(cell, 'column_letter'):
                if column is None:
                    column = cell.column_letter
                try:
                    if cell.value and len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
        if column:
            ws.column_dimensions[column].width = max(max_length + 2, 10)

    details = details or {}
    _write_detail_sheet(wb, "Додались (ПІБ)", details.get("added", []), color="2F5597")
    _write_detail_sheet(wb, "Відсіялись (ПІБ)", details.get("dropped", []), color="C00000")
    _write_detail_sheet(wb, "Стали працівниками", details.get("promoted", []), color="217346")
    _write_detail_sheet(wb, "Відхилені", details.get("rejected", []), color="7F6000")

    # Зберігаємо в буфер
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output

async def auto_monthly_report_sender(bot: Bot):
    """
    Автоматично розсилає місячний звіт усім розробникам.
    """
    from database.managers import get_all_developers
    from aiogram.types import BufferedInputFile
    from bot.services.logger import get_logger
    
    logger = get_logger()
    logger.info("🔵 Початок автоматичної розсилки місячного звіту...")
    
    now = datetime.now(pytz.timezone(TIMEZONE))
    # Звіт за поточний місяць (оскільки запускається в кінці місяця)
    start_date = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    
    try:
        data = await get_report_data(start_date, now)
        details = await get_report_details(start_date, now)
        has_details = any(details.get(k) for k in ("added", "dropped", "promoted", "rejected"))
        if not data and not has_details:
            logger.info("ℹ️ Немає даних для автоматичного звіту.")
            return
            
        xlsx_file = generate_xlsx_report(data, start_date, now, details)
        filename = f"Bulka_Monthly_Report_{start_date.strftime('%Y-%m')}.xlsx"
        
        devs = await get_all_developers()
        
        sent_count = 0
        for dev in devs:
            try:
                await bot.send_document(
                    chat_id=dev['uid'],
                    document=BufferedInputFile(xlsx_file.getvalue(), filename=filename),
                    caption=f"📊 <b>Підсумковий місячний звіт</b> ({start_date.strftime('%d.%m')} - {now.strftime('%d.%m')})"
                )
                sent_count += 1
            except Exception:
                pass
        
        logger.info(f"✅ Місячний звіт надіслано {sent_count} розробникам.")
        
    except Exception as exc:
        logger.error(f"Error in auto_monthly_report_sender: {exc}", exc_info=True)

from __future__ import annotations
import io
import re
from datetime import datetime
from typing import Dict, Any, List

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from database.material_notifications import (
    get_event_by_id,
    get_event_shops_summary,
    get_event_all_shop_users,
    get_event_overall_stats
)


def _sanitize_sheet_title(title: str, existing_titles: set[str]) -> str:
    """
    Очищає назву вкладки Excel від заборонених символів,
    обмежує довжину до 31 символу та гарантує унікальність.
    """
    # Видаляємо неприпустимі символи: \ / ? * : [ ]
    cleaned = re.sub(r'[\\/*?:\[\]]', '_', title).strip()
    if not cleaned:
        cleaned = "Магазин"
    cleaned = cleaned[:28]  # залишаємо запас для суфіксу

    final_title = cleaned
    counter = 1
    while final_title.lower() in existing_titles:
        final_title = f"{cleaned}_{counter}"[:31]
        counter += 1

    existing_titles.add(final_title.lower())
    return final_title


async def generate_material_ack_xlsx(event_id: int) -> io.BytesIO:
    """
    Генерує Excel-звіт про ознайомлення з матеріалами для події:
    - Аркуш 1: «Всі магазини» (Зведений звіт з усіма магазинами та співробітниками)
    - Аркуші 2..N: Окремий аркуш на кожен магазин
    """
    event = await get_event_by_id(event_id)
    if not event:
        raise ValueError(f"Подію з id={event_id} не знайдено.")

    role = event.get('role', 'Невідома посада')
    day = event.get('day', 1)
    desc = event.get('description', '')
    created_str = event.get('created_at', '')
    try:
        dt = datetime.fromisoformat(created_str.replace("Z", "+00:00"))
        dt_formatted = dt.strftime("%d.%m.%Y %H:%M")
    except Exception:
        dt_formatted = created_str[:16]

    overall_stats = await get_event_overall_stats(event_id)
    shops_summary = await get_event_shops_summary(event_id)

    wb = openpyxl.Workbook()
    existing_sheet_titles: set[str] = set()

    # Стилі
    font_title = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    font_subtitle = Font(name="Calibri", size=11, italic=True, color="595959")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_regular = Font(name="Calibri", size=11)
    font_bold = Font(name="Calibri", size=11, bold=True)

    fill_header = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    fill_shop_group = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    fill_ack = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")  # світло-зелений
    fill_noack = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")  # світло-червоний

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    border_thin = Border(
        left=Side(style='thin', color="D9D9D9"),
        right=Side(style='thin', color="D9D9D9"),
        top=Side(style='thin', color="D9D9D9"),
        bottom=Side(style='thin', color="D9D9D9")
    )
    border_header = Border(
        left=Side(style='thin', color="1F4E79"),
        right=Side(style='thin', color="1F4E79"),
        top=Side(style='thin', color="1F4E79"),
        bottom=Side(style='medium', color="1F4E79")
    )

    # -------------------------------------------------------------
    # 1. Аркуш «Всі магазини» (Зведений звіт)
    # -------------------------------------------------------------
    ws_all = wb.active
    ws_all.title = _sanitize_sheet_title("Всі магазини", existing_sheet_titles)

    # Шапка звіту
    ws_all.merge_cells("A1:G1")
    ws_all["A1"] = "ЗВІТ ПРО ОЗНАЙОМЛЕННЯ З НАВЧАЛЬНИМИ МАТЕРІАЛАМИ"
    ws_all["A1"].font = font_title
    ws_all["A1"].alignment = align_left

    ws_all.merge_cells("A2:G2")
    ws_all["A2"] = f"Посада: {role} | Навчальний день: {day} | Дата оновлення: {dt_formatted}"
    ws_all["A2"].font = font_subtitle
    ws_all["A2"].alignment = align_left

    ws_all.merge_cells("A3:G3")
    ws_all["A3"] = f"Опис змін: {desc}" if desc else "Опис змін: не вказано"
    ws_all["A3"].font = font_subtitle
    ws_all["A3"].alignment = align_left

    ws_all.merge_cells("A4:G4")
    ws_all["A4"] = (
        f"Всього отримувачів: {overall_stats['total']} | "
        f"Ознайомились: {overall_stats['acknowledged_count']} ({overall_stats['ack_percent']}%) | "
        f"Не ознайомились: {overall_stats['unacknowledged_count']} ({overall_stats['unack_percent']}%)"
    )
    ws_all["A4"].font = font_bold
    ws_all["A4"].alignment = align_left

    # Заголовки колонок
    headers_all = [
        "Магазин", "Категорія", "ПІБ співробітника", "Username",
        "Статус ознайомлення", "Дата ознайомлення", "Вчасно / Запізнення"
    ]
    ws_all.append([])  # Рядок 5 порожній
    ws_all.append(headers_all)  # Рядок 6
    header_row_idx = 6

    for col_idx in range(1, len(headers_all) + 1):
        cell = ws_all.cell(row=header_row_idx, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_header

    # Заповнення даними по магазинах
    shops_data_cache: Dict[str, List[Dict[str, Any]]] = {}

    for s in shops_summary:
        shop_name = s['shop']
        users = await get_event_all_shop_users(event_id, shop_name)
        shops_data_cache[shop_name] = users

        for u in users:
            cat_title = u['role_type'].capitalize()
            is_ack = bool(u.get('acknowledged_at'))
            status_text = "✅ Ознайомлений" if is_ack else "❌ Не ознайомлений"

            ack_date_str = "-"
            timeliness_str = "-"
            if is_ack:
                raw_ack = u['acknowledged_at']
                try:
                    ack_dt = datetime.fromisoformat(raw_ack.replace("Z", "+00:00"))
                    ack_date_str = ack_dt.strftime("%d.%m.%Y %H:%M")
                except Exception:
                    ack_date_str = raw_ack[:16]

                timeliness_str = "Після 72 год" if u.get('is_late') else "Вчасно"

            row_values = [
                shop_name,
                cat_title,
                u.get('full_name', ''),
                f"@{u['username']}" if u.get('username') else "-",
                status_text,
                ack_date_str,
                timeliness_str
            ]
            ws_all.append(row_values)
            curr_row = ws_all.max_row

            # Оформлення клітинок
            for col_idx in range(1, len(row_values) + 1):
                cell = ws_all.cell(row=curr_row, column=col_idx)
                cell.font = font_regular
                cell.border = border_thin
                if col_idx in (1, 2, 4, 5, 6, 7):
                    cell.alignment = align_center

            # Підсвітка статусу
            status_cell = ws_all.cell(row=curr_row, column=5)
            status_cell.fill = fill_ack if is_ack else fill_noack
            status_cell.font = font_bold

    # -------------------------------------------------------------
    # 2. Окремі аркуші для кожного магазину
    # -------------------------------------------------------------
    headers_shop = [
        "Категорія", "ПІБ співробітника", "Username",
        "Статус ознайомлення", "Дата ознайомлення", "Вчасно / Запізнення"
    ]

    for s in shops_summary:
        shop_name = s['shop']
        users = shops_data_cache.get(shop_name, [])

        sheet_title = _sanitize_sheet_title(shop_name, existing_sheet_titles)
        ws_shop = wb.create_sheet(title=sheet_title)

        # Заголовок аркуша магазину
        ws_shop.merge_cells("A1:F1")
        ws_shop["A1"] = f"МАГАЗИН: {shop_name}"
        ws_shop["A1"].font = font_title
        ws_shop["A1"].alignment = align_left

        total_sh = s.get('total_count', len(users))
        ack_sh = s.get('ack_count', sum(1 for u in users if u.get('acknowledged_at')))
        pct_sh = round(ack_sh / total_sh * 100, 1) if total_sh > 0 else 0.0

        ws_shop.merge_cells("A2:F2")
        ws_shop["A2"] = f"Всього співробітників: {total_sh} | Ознайомились: {ack_sh} ({pct_sh}%) | Не ознайомились: {total_sh - ack_sh}"
        ws_shop["A2"].font = font_bold
        ws_shop["A2"].alignment = align_left

        ws_shop.append([])  # Рядок 3
        ws_shop.append(headers_shop)  # Рядок 4
        sh_header_idx = 4

        for col_idx in range(1, len(headers_shop) + 1):
            cell = ws_shop.cell(row=sh_header_idx, column=col_idx)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = border_header

        for u in users:
            cat_title = u['role_type'].capitalize()
            is_ack = bool(u.get('acknowledged_at'))
            status_text = "✅ Ознайомлений" if is_ack else "❌ Не ознайомлений"

            ack_date_str = "-"
            timeliness_str = "-"
            if is_ack:
                raw_ack = u['acknowledged_at']
                try:
                    ack_dt = datetime.fromisoformat(raw_ack.replace("Z", "+00:00"))
                    ack_date_str = ack_dt.strftime("%d.%m.%Y %H:%M")
                except Exception:
                    ack_date_str = raw_ack[:16]
                timeliness_str = "Після 72 год" if u.get('is_late') else "Вчасно"

            row_values = [
                cat_title,
                u.get('full_name', ''),
                f"@{u['username']}" if u.get('username') else "-",
                status_text,
                ack_date_str,
                timeliness_str
            ]
            ws_shop.append(row_values)
            curr_row = ws_shop.max_row

            for col_idx in range(1, len(row_values) + 1):
                cell = ws_shop.cell(row=curr_row, column=col_idx)
                cell.font = font_regular
                cell.border = border_thin
                if col_idx in (1, 3, 4, 5, 6):
                    cell.alignment = align_center

            status_cell = ws_shop.cell(row=curr_row, column=4)
            status_cell.fill = fill_ack if is_ack else fill_noack
            status_cell.font = font_bold

    # -------------------------------------------------------------
    # 3. Авто-підбір ширини колонок для всіх аркушів
    # -------------------------------------------------------------
    for ws in wb.worksheets:
        ws.views.sheetView[0].showGridLines = True
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # Ігноруємо об'єднані заголовки
                if cell.row < 6 and ws.title == "Всі магазини":
                    continue
                if cell.row < 4 and ws.title != "Всі магазини":
                    continue
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

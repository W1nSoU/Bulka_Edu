from __future__ import annotations
import io
import re
from typing import Dict, Any, List, Set

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from database import DB_PATH
from database.news import (
    get_news_by_id,
    get_news_reactions_summary,
    get_news_failed_deliveries_detailed,
    get_news_reactions_detailed
)


def _sanitize_sheet_title(title: str, existing_titles: Set[str]) -> str:
    cleaned = re.sub(r'[\\/*?:\[\]]', '_', title).strip()
    if not cleaned:
        cleaned = "Sheet"
    cleaned = cleaned[:28]

    final_title = cleaned
    counter = 1
    while final_title.lower() in existing_titles:
        final_title = f"{cleaned}_{counter}"[:31]
        counter += 1

    existing_titles.add(final_title.lower())
    return final_title


async def generate_news_report_xlsx(news_id: int, db_path: str = DB_PATH) -> io.BytesIO:
    """
    Генерує брендований Excel-звіт за результатами розсилки новини:
    - Аркуш 1: «Загальний» (метадані, KPI розсилки, аналітика реакцій).
    - Аркуш 2: «Не доставлено» (список співробітників, яким не вдалося доставити, із класифікованими причинами).
    - Аркуш 3: «Реакції» (список співробітників, які залишили реакцію: ПІБ, магазин, посада, емодзі, час).
    """
    news = await get_news_by_id(news_id, db_path=db_path)
    if not news:
        raise ValueError(f"Новину з id={news_id} не знайдено.")

    reactions_summary = await get_news_reactions_summary(news_id, db_path=db_path)
    failed_deliveries = await get_news_failed_deliveries_detailed(news_id, db_path=db_path)
    reactions_detailed = await get_news_reactions_detailed(news_id, db_path=db_path)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Видаляємо стандартний порожній аркуш

    # Стилі
    font_title = Font(name="Arial", size=13, bold=True, color="1F2937")
    font_sub = Font(name="Arial", size=9, italic=True, color="4B5563")
    font_section = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Arial", size=10, color="1F2937")
    font_bold = Font(name="Arial", size=10, bold=True, color="1F2937")

    fill_primary = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")     # Темно-синій
    fill_header = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")      # Синій
    fill_kpi_header = PatternFill(start_color="D97706", end_color="D97706", fill_type="solid")  # Янтарний
    fill_zebra = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")       # Світло-сірий
    fill_error_header = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")# Червоний
    fill_success_header = PatternFill(start_color="059669", end_color="059669", fill_type="solid") # Зелений

    thin_border_side = Side(style='thin', color="D1D5DB")
    border_data = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)

    used_sheet_titles: Set[str] = set()

    # =========================================================================
    # АРКУШ 1: Загальний
    # =========================================================================
    ws_summary = wb.create_sheet(title=_sanitize_sheet_title("Загальний", used_sheet_titles))
    ws_summary.views.sheetView[0].showGridLines = True

    # Заголовок
    ws_summary.merge_cells("A1:F1")
    ws_summary["A1"] = f"📰 ЗВІТ РОЗСИЛКИ НОВИНИ #{news_id}"
    ws_summary["A1"].font = font_title
    ws_summary["A1"].alignment = align_left
    ws_summary.row_dimensions[1].height = 25

    created_at = news.get("created_at") or "—"
    target_city = news.get("selected_city") or "all"
    city_str = "Усі міста" if target_city == "all" else target_city
    roles_list = news.get("selected_roles") or []
    roles_str = ", ".join(roles_list) if roles_list else "Всі категорії"

    ws_summary["A2"] = f"📅 Дата розсилки: {created_at} | 🏙 Місто: {city_str} | 👥 Аудиторія: {roles_str}"
    ws_summary["A2"].font = font_sub
    ws_summary["A2"].alignment = align_left

    # Текст новини
    ws_summary.merge_cells("A4:F4")
    ws_summary["A4"] = "Текст публікації:"
    ws_summary["A4"].font = font_bold

    news_text = news.get("text") or "—"
    ws_summary.merge_cells("A5:F5")
    ws_summary["A5"] = news_text
    ws_summary["A5"].font = font_data
    ws_summary["A5"].alignment = align_wrap
    ws_summary.row_dimensions[5].height = min(120, max(45, len(news_text) // 2))

    # Секція: Показники розсилки
    ws_summary.merge_cells("A7:D7")
    cell_sec1 = ws_summary["A7"]
    cell_sec1.value = "📊 ПОКАЗНИКИ ДОСТАВКИ ТА ОХОПЛЕННЯ"
    cell_sec1.font = font_section
    cell_sec1.fill = fill_primary
    cell_sec1.alignment = align_left
    ws_summary.row_dimensions[7].height = 22

    total_recipients = news.get("total_recipients") or 0
    sent_count = news.get("sent_count") or 0
    failed_count = news.get("failed_count") or 0
    total_waves = news.get("total_waves") or 0

    sent_pct = round((sent_count / total_recipients * 100), 1) if total_recipients > 0 else 0
    failed_pct = round((failed_count / total_recipients * 100), 1) if total_recipients > 0 else 0
    total_reactions = reactions_summary.get("total", 0)
    reaction_pct = round((total_reactions / sent_count * 100), 1) if sent_count > 0 else 0

    kpi_data = [
        ("Всього отримувачів у вибірці", total_recipients, "100%"),
        ("✅ Успішно доставлено", sent_count, f"{sent_pct}%"),
        ("❌ Не вдалося доставити", failed_count, f"{failed_pct}%"),
        ("🌊 Кількість хвиль розсилки", total_waves, "хвиль по 50 осіб"),
        ("💬 Всього отримано реакцій", total_reactions, f"{reaction_pct}% від доставлених")
    ]

    r_idx = 8
    for label, val, note in kpi_data:
        ws_summary.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=2)
        cell_lbl = ws_summary.cell(row=r_idx, column=1, value=label)
        cell_lbl.font = font_bold
        cell_lbl.border = border_data
        ws_summary.cell(row=r_idx, column=2).border = border_data

        cell_val = ws_summary.cell(row=r_idx, column=3, value=val)
        cell_val.font = font_bold
        cell_val.alignment = align_center
        cell_val.border = border_data

        cell_note = ws_summary.cell(row=r_idx, column=4, value=note)
        cell_note.font = font_sub
        cell_note.alignment = align_left
        cell_note.border = border_data
        r_idx += 1

    # Секція: Розподіл реакцій
    r_idx += 1
    ws_summary.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=4)
    cell_sec2 = ws_summary.cell(row=r_idx, column=1)
    cell_sec2.value = "💬 АНАЛІТИКА РЕАКЦІЙ СПІВРОБІТНИКІВ"
    cell_sec2.font = font_section
    cell_sec2.fill = fill_kpi_header
    cell_sec2.alignment = align_left
    ws_summary.row_dimensions[r_idx].height = 22
    r_idx += 1

    reaction_headers = ["Емодзі", "Значення", "Кількість", "Частка"]
    for col_i, h in enumerate(reaction_headers, start=1):
        c = ws_summary.cell(row=r_idx, column=col_i, value=h)
        c.font = font_header
        c.fill = fill_header
        c.alignment = align_center
        c.border = border_data
    r_idx += 1

    reaction_labels = {
        "👎": "Не подобається",
        "🤔": "Замислився",
        "❤️": "Чудово / Любов",
        "🔥": "Супер / Вогонь"
    }

    breakdown = reactions_summary.get("breakdown", {})
    for emoji, label in reaction_labels.items():
        cnt = breakdown.get(emoji, 0)
        pct = round((cnt / total_reactions * 100), 1) if total_reactions > 0 else 0

        c_emoji = ws_summary.cell(row=r_idx, column=1, value=emoji)
        c_emoji.font = Font(name="Arial", size=14)
        c_emoji.alignment = align_center
        c_emoji.border = border_data

        c_lbl = ws_summary.cell(row=r_idx, column=2, value=label)
        c_lbl.font = font_data
        c_lbl.alignment = align_left
        c_lbl.border = border_data

        c_cnt = ws_summary.cell(row=r_idx, column=3, value=cnt)
        c_cnt.font = font_bold
        c_cnt.alignment = align_center
        c_cnt.border = border_data

        c_pct = ws_summary.cell(row=r_idx, column=4, value=f"{pct}%")
        c_pct.font = font_bold
        c_pct.alignment = align_center
        c_pct.border = border_data

        r_idx += 1

    # =========================================================================
    # АРКУШ 2: Не доставлено
    # =========================================================================
    ws_failed = wb.create_sheet(title=_sanitize_sheet_title("Не доставлено", used_sheet_titles))
    ws_failed.views.sheetView[0].showGridLines = True

    ws_failed.merge_cells("A1:H1")
    ws_failed["A1"] = f"❌ СПИСОК НЕ ДОСТАВЛЕНИХ ПОВІДОМЛЕНЬ ({len(failed_deliveries)} осіб)"
    ws_failed["A1"].font = font_title
    ws_failed["A1"].alignment = align_left
    ws_failed.row_dimensions[1].height = 25

    failed_headers = ["№", "ПІБ", "Посада", "Місто", "Магазин", "Telegram ID", "Причина", "Технічна помилка"]
    for col_i, h in enumerate(failed_headers, start=1):
        c = ws_failed.cell(row=3, column=col_i, value=h)
        c.font = font_header
        c.fill = fill_error_header
        c.alignment = align_center
        c.border = border_data
    ws_failed.row_dimensions[3].height = 22

    for row_idx, item in enumerate(failed_deliveries, start=4):
        fill = fill_zebra if row_idx % 2 == 0 else PatternFill(fill_type=None)

        vals = [
            row_idx - 3,
            item.get("full_name") or "Не вказано",
            item.get("role") or "—",
            item.get("city") or "—",
            item.get("shop") or "—",
            item.get("user_id"),
            item.get("error_reason") or "Невідома помилка",
            item.get("raw_error") or "—"
        ]

        for col_i, val in enumerate(vals, start=1):
            c = ws_failed.cell(row=row_idx, column=col_i, value=val)
            c.font = font_data
            c.border = border_data
            if fill.fill_type:
                c.fill = fill
            if col_i in (1, 6):
                c.alignment = align_center
            elif col_i == 7:
                c.alignment = align_left
                c.font = font_bold
            else:
                c.alignment = align_left

    # =========================================================================
    # АРКУШ 3: Реакції
    # =========================================================================
    ws_react = wb.create_sheet(title=_sanitize_sheet_title("Реакції", used_sheet_titles))
    ws_react.views.sheetView[0].showGridLines = True

    ws_react.merge_cells("A1:G1")
    ws_react["A1"] = f"💬 РЕАКЦІЇ СПІВРОБІТНИКІВ ({len(reactions_detailed)} осіб)"
    ws_react["A1"].font = font_title
    ws_react["A1"].alignment = align_left
    ws_react.row_dimensions[1].height = 25

    react_headers = ["№", "ПІБ", "Магазин", "Посада", "Місто", "Реакція", "Час реакції"]
    for col_i, h in enumerate(react_headers, start=1):
        c = ws_react.cell(row=3, column=col_i, value=h)
        c.font = font_header
        c.fill = fill_success_header
        c.alignment = align_center
        c.border = border_data
    ws_react.row_dimensions[3].height = 22

    for row_idx, item in enumerate(reactions_detailed, start=4):
        fill = fill_zebra if row_idx % 2 == 0 else PatternFill(fill_type=None)

        vals = [
            row_idx - 3,
            item.get("full_name") or "Не вказано",
            item.get("shop") or "—",
            item.get("role") or "—",
            item.get("city") or "—",
            item.get("reaction") or "—",
            item.get("updated_at") or item.get("created_at") or "—"
        ]

        for col_i, val in enumerate(vals, start=1):
            c = ws_react.cell(row=row_idx, column=col_i, value=val)
            c.font = font_data
            c.border = border_data
            if fill.fill_type:
                c.fill = fill
            if col_i in (1, 7):
                c.alignment = align_center
            elif col_i == 6:
                c.alignment = align_center
                c.font = Font(name="Arial", size=14)
            else:
                c.alignment = align_left

    # Автопідбір ширини колонок для всіх аркушів
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # Не враховуємо об'єднані довгі рядки
                if cell.row in (1, 2, 4, 5, 7):
                    continue
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(12, min(max_len + 4, 45))

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
